from __future__ import annotations
import logging
from dataclasses import dataclass
from pathlib import Path
from datetime import datetime, timedelta, date
from typing import Tuple, Dict
import numpy as np
from scipy.optimize import linprog
from scipy import sparse
from statsmodels.tsa.ar_model import AutoReg
try:
    from statsmodels.tsa.seasonal import MSTL
    HAS_MSTL = True
except Exception:
    HAS_MSTL = False
try:
    from openpyxl import load_workbook
except ImportError as e:
    raise ImportError(
        "缺少 openpyxl。请在当前 Python 环境执行：pip install openpyxl"
    ) from e
# 绘图
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
plt.rcParams.update({
    "font.family": "sans-serif",
    "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Arial"],
    "axes.unicode_minus": False, "font.size": 12, "axes.spines.top": False,
    "axes.linewidth": 1.5, "xtick.major.width": 1.5, "ytick.major.width": 1.5,
    "xtick.major.size": 6, "ytick.major.size": 6,
})
#设置参数
class Config:
    b_dir = str(Path(__file__).resolve().parent)
    att1 = str(Path(b_dir) / "附件1.xlsx")
    att2 = str(Path(b_dir) / "附件2.xlsx")
    _tpl_candidates = [
        Path(b_dir) / "附件5" / "result2.xlsx",
        Path(b_dir) / "result2.xlsx",
    ]
    template = str(next(p for p in _tpl_candidates if p.exists()))
    out_dir = str(Path(b_dir) / "Q2结果")
    out_xlsx = str(Path(out_dir) / "result2_filled.xlsx")
    fig_dir = str(Path(out_dir) / "第二问图片输出")
    # 时间离散与储能
    slots_per_day = 144; dt = 1.0 / 6.0; horizon_slots = 144
    forecast_slots = 288; eta_c = eta_d = 0.90
    e_min, e_max = 1200.0, 10800.0
    p_max, e0 = 5000.0, 6000.0
    #预测
    l_days = 42; pv_days = 60; l_min_days = 21
    mstl_days = 7       # MSTL每7天完整重估，中间仅轻量更新
    ar_lags_candi = (1, 2, 3, 6, 12, 18, 36)
    pv_candi, kappa_candi = (1,2,3,6,12), (1,2,3,5,7)
    clear_quantile, pv_halfwin = 0.93, 2
    #动态风险权重
    rk_hisdays = 28; rk_decay = 0.90; rk_minsamp = 5
    alpha_cold, alpha_min, alpha_max = 4.0, 1.0, 8.0   # 上界防长期顶格
    beta_fixed, c_eps = 1.0, 1e-4
    tau_min, tau_max = 0.55, 0.88     # 上界避免过度保守
    reserve_scale = 0.75              # 兼顾紧急购电与富余电量
    #优化与运行控制
    eps_cyc, eps_sur, lp_met = 1e-5, 1e-8, "highs"
    max_days = 365   # 测试可改40，正式用365
CFG = Config()
def setup_logger(outdir: Path) -> logging.Logger:
    outdir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("Q2")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    return logger
def excel_serial_to_date(x) -> datetime:
    return datetime(1899, 12, 30) + timedelta(days=float(x))
def cell_to_date(x):
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    try:
        from datetime import date as _date
        if isinstance(x, _date):
            return x
    except Exception:
        pass
    if isinstance(x, (int, float, np.integer, np.floating)):
        return excel_serial_to_date(x).date()
    if isinstance(x, str):
        txt=x.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(txt, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"无法识别日期单元格：{x!r}")
def safe_float_array(x) -> np.ndarray:
    return np.asarray([[float(v) if v is not None else np.nan for v in row] for row in x], dtype=float)
def weighted_quantile(values: np.ndarray, weights: np.ndarray, q: float) -> float:
    values = np.asarray(values, dtype=float)
    weights = np.asarray(weights, dtype=float)
    m = np.isfinite(values) & np.isfinite(weights) & (weights > 0)
    if m.sum() == 0:
        return 0.0
    v, w = values[m], weights[m]
    idx = np.argsort(v)
    v, w = v[idx], w[idx]
    cw = np.cumsum(w) / np.sum(w)
    return float(np.interp(q, cw, v))
def moving_average_circular(x: np.ndarray, half: int) -> np.ndarray:
    if half <= 0:
        return x.copy()
    n = len(x)
    out = np.zeros(n)
    for i in range(n):
        lo = max(0, i-half)
        hi = min(n, i+half+1)
        out[i] = np.mean(x[lo:hi])
    return out
def choose_ar_forecast(series: np.ndarray, horizon: int, candidates: Tuple[int, ...], fallback: float = 0.0) -> np.ndarray:
    y = np.asarray(series, dtype=float)
    y = y[np.isfinite(y)]
    if len(y) < 8 or np.std(y) < 1e-8:
        return np.full(horizon, float(np.mean(y) if len(y) else fallback))
    best = None
    for p in candidates:
        if len(y) <= p + 5:
            continue
        try:
            model = AutoReg(y, lags=p, trend="ct", old_names=False).fit()
            if np.isfinite(model.aic) and (best is None or model.aic < best[0]):
                best = (model.aic, model)
        except Exception:
            continue
    if best is None:
        return np.full(horizon, float(np.mean(y[-min(7, len(y)):])) )
    try:
        pred = np.asarray(best[1].predict(start=len(y), end=len(y)+horizon-1, dynamic=False), dtype=float)
        if np.any(~np.isfinite(pred)):
            raise ValueError("AR forecast nonfinite")
        return pred
    except Exception:
        return np.full(horizon, float(np.mean(y[-min(7, len(y)):])) )
#读取数据
def load_inputs(cfg: Config, logger):
    for p in (cfg.att1, cfg.att2, cfg.template):
        if not Path(p).exists():
            raise FileNotFoundError(f"未找到文件：{p}")
    wb1 = load_workbook(cfg.att1, data_only=True, read_only=True)
    s1 = wb1["Sheet1"]
    rows1 = list(s1.iter_rows(min_row=2, max_row=145, min_col=2, max_col=4, values_only=True))
    price = np.array([float(r[0]) for r in rows1], dtype=float)
    load_prior = np.array([float(r[1]) for r in rows1], dtype=float)
    pv_prior = np.array([float(r[2]) for r in rows1], dtype=float)
    wb1.close()
    wb2 = load_workbook(cfg.att2, data_only=True, read_only=True)
    sl = wb2["小区负载"]
    sp = wb2["光伏发电实际功率"]
    load_all = list(sl.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    pv_all = list(sp.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    wb2.close()
    dates = [cell_to_date(r[0]) for r in load_all]
    load_rows = [r[1:] for r in load_all]
    pv_rows = [r[1:] for r in pv_all]
    load = safe_float_array(load_rows)
    pv_actual = safe_float_array(pv_rows)
    assert load.shape == (365, 144), load.shape
    assert pv_actual.shape == (365, 144), pv_actual.shape
    assert len(price) == 144
    return dates, price, load_prior, pv_prior, load, pv_actual
_LOAD_MSTL_CACHE = {}
def load_forecast(day_idx: int, load: np.ndarray, prior: np.ndarray, cfg: Config, logger=None) -> np.ndarray:
    H = cfg.forecast_slots
    if day_idx == 0:
        return np.tile(prior, 2)
    hist_days = min(day_idx, cfg.l_days)
    hist_mat = load[day_idx-hist_days:day_idx]
    # 冷启动：同一时刻近期日加权，附件1先验
    if hist_days < cfg.l_min_days or not HAS_MSTL:
        k = hist_mat.shape[0]
        weights = np.exp(-0.10*np.arange(k-1, -1, -1))
        weights /= weights.sum()
        shape = np.sum(hist_mat * weights[:, None], axis=0)
        if hist_days >= 7:
            shape = 0.55*shape + 0.45*load[day_idx-7]
        w_prior = max(0.0, 1.0 - hist_days / cfg.l_min_days)
        shape = (1-w_prior)*shape + w_prior*prior
        return np.tile(shape, 2)
    need_refit = (
        not _LOAD_MSTL_CACHE
        or day_idx - _LOAD_MSTL_CACHE.get("refit_day", -10**9) >= cfg.mstl_days
    )
    if need_refit:
        y = hist_mat.reshape(-1)
        try:
            mstl = MSTL(y, periods=(144, 1008), stl_kwargs={"robust": False}).fit()
            trend = np.asarray(mstl.trend)
            seasonal = np.asarray(mstl.seasonal)
            if seasonal.ndim == 1:
                s_daily = seasonal; s_weekly = np.zeros_like(seasonal)
            else:
                s_daily = seasonal[:, 0]; s_weekly = seasonal[:, 1]
            resid = np.asarray(mstl.resid)
            daily_tail = s_daily[-min(len(s_daily), 28*144):]
            weekly_tail = s_weekly[-min(len(s_weekly), 6*1008):]
            daily_pattern = np.array([np.nanmean(daily_tail[i::144]) for i in range(144)])
            weekly_pattern = np.array(
                [np.nanmean(weekly_tail[i::1008]) for i in range(1008)])
            _LOAD_MSTL_CACHE.clear()
            _LOAD_MSTL_CACHE.update({
                "refit_day": day_idx,
                "origin_day": day_idx - hist_days,
                "daily_pattern": daily_pattern,
                "weekly_pattern": weekly_pattern,
                "trend_last": float(trend[-1]),
                "trend_slope": float(np.clip(
                    np.polyfit(
                        np.arange(min(432, len(trend)), dtype=float),
                        trend[-min(432, len(trend)):],
                        1
                    )[0], -5.0, 5.0
                )),
                "resid_tail": resid[-min(len(resid), 14*144):].copy(),
            })
        except Exception as e:
            if logger:
                logger.warning("MSTL失败(day=%d)，回退季节加权预测：%s", day_idx, e)
            k = hist_mat.shape[0]
            weights = np.exp(-0.10*np.arange(k-1, -1, -1))
            weights /= weights.sum()
            shape = np.sum(hist_mat * weights[:, None], axis=0)
            if hist_days >= 7:
                shape = 0.5*shape + 0.5*load[day_idx-7]
            return np.tile(shape, 2)
    # 使用最近一次 MSTL 的季节结构，并用最新 3 天数据快速校正趋势和残差
    daily_pattern = _LOAD_MSTL_CACHE["daily_pattern"]
    weekly_pattern = _LOAD_MSTL_CACHE["weekly_pattern"]
    origin_day = _LOAD_MSTL_CACHE["origin_day"]
    recent_days = min(hist_days, 3)
    recent = load[day_idx-recent_days:day_idx].reshape(-1)
    global_start = (day_idx - recent_days) * 144
    gidx = global_start + np.arange(len(recent))
    sd_hist = daily_pattern[gidx % 144]
    sw_hist = weekly_pattern[(gidx - origin_day*144) % 1008]
    deseason = recent - sd_hist - sw_hist
    m = min(432, len(deseason))
    x = np.arange(m, dtype=float)
    coef = np.polyfit(x, deseason[-m:], 1)
    slope = float(np.clip(coef[0], -5.0, 5.0))
    level = float(deseason[-1])
    trend_future = level + slope*np.arange(1, H+1)
    future_gidx = day_idx*144 + np.arange(H)
    sd_future = daily_pattern[future_gidx % 144]
    sw_future = weekly_pattern[(future_gidx - origin_day*144) % 1008]
    # 残差AR仍每天轻量更新，不需要重新做MSTL
    hist_axis = np.arange(len(deseason), dtype=float)
    fitted_trend_hist = level + slope*(hist_axis - (len(deseason)-1))
    resid_recent = deseason - fitted_trend_hist
    ar_future = choose_ar_forecast(resid_recent[-min(len(resid_recent), 3*144):],
        H,cfg.ar_lags_candi,0.0,)
    pred = trend_future + sd_future + sw_future + ar_future
    return np.maximum(pred, 0.0)
# 光伏：晴空包络×日透射率+残差AR
def build_clear_envelope(day_idx: int, pv: np.ndarray, prior: np.ndarray, cfg: Config) -> np.ndarray:
    if day_idx == 0:
        base = prior.copy()
        return np.maximum(base, 0.0)
    w = min(day_idx, cfg.pv_days)
    hist = pv[day_idx-w:day_idx]
    q = np.quantile(hist, cfg.clear_quantile, axis=0)
    q = moving_average_circular(q, cfg.pv_halfwin)
    prior_w = max(0.0, 1.0 - w/21.0)
    base = (1-prior_w)*q + prior_w*prior
    zero_ratio = np.mean(hist <= 1e-8, axis=0)
    base[zero_ratio > 0.90] = 0.0
    return np.maximum(base, 0.0)
def pv_forecast(day_idx: int, pv: np.ndarray, prior: np.ndarray, cfg: Config, logger=None) -> Tuple[np.ndarray, Dict]:
    H = cfg.forecast_slots
    if day_idx == 0:
        return np.tile(prior, 2), {"kappa_d": 1.0, "kappa_next": 1.0, "envelope": prior.copy()}
    B_d = build_clear_envelope(day_idx, pv, prior, cfg)
    # 下一日包络：年周期变化慢，先用当前包络；用近期日能量变化率做轻微尺度外推
    recent_energy = np.sum(pv[max(0,day_idx-14):day_idx], axis=1)*cfg.dt
    scale_next = 1.0
    if len(recent_energy) >= 7:
        x = np.arange(len(recent_energy))
        slope = np.polyfit(x, recent_energy, 1)[0]
        denom = max(np.mean(recent_energy), 1.0)
        scale_next = float(np.clip(1.0 + slope/denom, 0.95, 1.05))
    B_next = B_d * scale_next
    # 用当前滚动包络近似反演近期日透射率；只使用过去日期
    w = min(day_idx, cfg.pv_days)
    hist = pv[day_idx-w:day_idx]
    denom = max(np.sum(B_d)*cfg.dt, 1e-6)
    kappas = np.clip(np.sum(hist, axis=1)*cfg.dt/denom, 0.0, 1.0)
    kpred = choose_ar_forecast(kappas, 2, cfg.kappa_candi, fallback=float(np.mean(kappas)))
    kpred = np.clip(kpred, 0.0, 1.0)
    base_d = B_d*kpred[0]
    base_next = B_next*kpred[1]
    # 残差构造：每个历史日用其日能量对应kappa相对于当前包络重构
    residual_days = hist - kappas[:, None]*B_d[None, :]
    resid_series = residual_days.reshape(-1)
    rpred = choose_ar_forecast(resid_series[-min(len(resid_series), 21*144):], H, cfg.pv_candi, 0.0)
    pred = np.concatenate([base_d, base_next]) + rpred
    pred = np.maximum(pred, 0.0)
    # 夜间稳定无发电时段置零
    zero_mask = B_d <= 1e-8
    pred[zero_mask[np.arange(H) % 144]] = 0.0
    return pred, {"kappa_d": float(kpred[0]), "kappa_next": float(kpred[1]), "envelope": B_d}
# 动态alpha/beta与风险备用量
def compute_historical_marginal_cost_day(j: int,
                                         available_day_exclusive: int,
                                         net_errors: np.ndarray,
                                         price: np.ndarray,
                                         soc_hist: np.ndarray,
                                         planned_hist: np.ndarray,
                                         load: np.ndarray,
                                         pv: np.ndarray,
                                         cfg: Config) -> Tuple[np.ndarray, np.ndarray]:
    nT = cfg.slots_per_day
    cplus = np.full(nT, np.nan, dtype=float)
    cminus = np.full(nT, np.nan, dtype=float)
    if j < 0 or j >= available_day_exclusive:
        return cplus, cminus
    n_avail = available_day_exclusive * nT
    plan_f = planned_hist[:available_day_exclusive].reshape(-1)
    load_f = load[:available_day_exclusive].reshape(-1)
    pv_f = pv[:available_day_exclusive].reshape(-1)
    soc_f = soc_hist[:available_day_exclusive].reshape(-1)
    err_f = net_errors[:available_day_exclusive].reshape(-1)
    slot_price = np.tile(price, available_day_exclusive)
    base = j * nT
    eff_rt = cfg.eta_c * cfg.eta_d
    for t in range(nT):
        e = net_errors[j, t]
        if not np.isfinite(e):
            continue
        lam_t = price[t]
        s0 = soc_hist[j, t] if np.isfinite(soc_hist[j, t]) else cfg.e0
        g0 = base + t
        g1 = g0 + 1
        g2 = min(g0 + cfg.horizon_slots, n_avail - 1)
        if e > 0:
            c_em = 5.0 * lam_t
            c_bat = np.inf
            if s0 > cfg.e_min + 1e-6 and g1 <= g2:
                sl = slice(g1, g2 + 1)
                margin = plan_f[sl] + pv_f[sl] - load_f[sl]
                sf = soc_f[sl]
                pf = slot_price[sl]
                soc_ok = (~np.isfinite(sf)) | (sf < cfg.e_max - 1e-6)
                feasible = soc_ok & ((margin > 0.0) | (pf < lam_t))
                if np.any(feasible):
                    c_bat = float(np.min(pf[feasible]) / eff_rt)
            cplus[t] = min(c_em, c_bat) if np.isfinite(c_bat) else c_em
        elif e < 0:
            recovery = 0.0
            if s0 < cfg.e_max - 1e-6 and g1 <= g2:
                sl = slice(g1, g2 + 1)
                ef = err_f[sl]
                pf = slot_price[sl]
                m = np.isfinite(ef) & (ef > 0.0)
                if np.any(m):
                    recovery = float(np.max(5.0 * pf[m] * eff_rt))
            cminus[t] = max(0.0, lam_t - recovery)
    return cplus, cminus
def update_risk_cost_cache(day_idx: int,
                           cplus_cache: np.ndarray,
                           cminus_cache: np.ndarray,
                           net_errors: np.ndarray,
                           price: np.ndarray,
                           soc_hist: np.ndarray,
                           planned_hist: np.ndarray,
                           load: np.ndarray,
                           pv: np.ndarray,
                           cfg: Config) -> None:
    if day_idx <= 0:
        return
    j = day_idx - 1
    cp, cm = compute_historical_marginal_cost_day(
        j, day_idx, net_errors, price, soc_hist, planned_hist, load, pv, cfg)
    cplus_cache[j] = cp
    cminus_cache[j] = cm
    if day_idx >= 2:
        j = day_idx - 2
        cp, cm = compute_historical_marginal_cost_day(
            j, day_idx, net_errors, price, soc_hist, planned_hist, load, pv, cfg)
        cplus_cache[j] = cp
        cminus_cache[j] = cm
def dynamic_risk_weights(day_idx: int,
                         net_errors: np.ndarray,
                         cplus_cache: np.ndarray,
                         cminus_cache: np.ndarray,
                         cfg: Config) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    nT = cfg.slots_per_day
    alpha = np.full(nT, cfg.alpha_cold, dtype=float)
    beta = np.full(nT, cfg.beta_fixed, dtype=float)
    reserve = np.zeros(nT, dtype=float)
    if day_idx <= 0:
        return alpha, beta, reserve
    j0 = max(0, day_idx - cfg.rk_hisdays)
    hist_days = np.arange(j0, day_idx)
    if hist_days.size == 0:
        return alpha, beta, reserve
    ages = day_idx - hist_days
    base_w = cfg.rk_decay ** (ages - 1)
    errs = net_errors[hist_days]          # (HIST, 144)
    cp_all = cplus_cache[hist_days]
    cm_all = cminus_cache[hist_days]
    for t in range(nT):
        et = errs[:, t]
        # 正向误差历史
        mp = (et > 0) & np.isfinite(et) & np.isfinite(cp_all[:, t])
        # 负向误差历史
        mn = (et < 0) & np.isfinite(et) & np.isfinite(cm_all[:, t])
        if np.count_nonzero(mp) >= cfg.rk_minsamp and np.count_nonzero(mn) >= cfg.rk_minsamp:
            cp = float(np.average(cp_all[mp, t], weights=base_w[mp]))
            cm = float(np.average(cm_all[mn, t], weights=base_w[mn]))
            raw = cp / max(cm, cfg.c_eps)
            alpha[t] = min(cfg.alpha_max, max(cfg.alpha_min, raw))
        tau = alpha[t] / (alpha[t] + beta[t])
        tau = min(cfg.tau_max, max(cfg.tau_min, tau))
        if np.count_nonzero(mp) >= cfg.rk_minsamp:
            reserve[t] = cfg.reserve_scale * weighted_quantile(et[mp], base_w[mp], tau)
    return alpha, beta, reserve
# 7. LP构造
@dataclass
class LPTemplates:
    da_Aeq: sparse.csr_matrix
    da_c: np.ndarray
    da_bounds: list
    mu_terminal: float
def build_lp_templates(price: np.ndarray, cfg: Config) -> LPTemplates:
    n = cfg.slots_per_day
    N = 5 * n
    ib, ic, idis, isu, ie = 0, n, 2*n, 3*n, 4*n
    rows, cols, data = [], [], []
    r = 0
    # 功率平衡：buy + dis - ch - sur = 风险校准后的预测净负荷
    for t in range(n):
        for col, val in ((ib+t,1.0),(idis+t,1.0),(ic+t,-1.0),(isu+t,-1.0)):
            rows.append(r); cols.append(col); data.append(val)
        r += 1
    # SOC递推
    for t in range(n):
        rows.append(r); cols.append(ie+t); data.append(1.0)
        if t > 0:
            rows.append(r); cols.append(ie+t-1); data.append(-1.0)
        rows.append(r); cols.append(ic+t); data.append(-cfg.eta_c*cfg.dt)
        rows.append(r); cols.append(idis+t); data.append(cfg.dt/cfg.eta_d)
        r += 1
    Aeq = sparse.csr_matrix((data,(rows,cols)), shape=(r,N))
    mu = cfg.eta_d * float(np.min(price))
    c = np.zeros(N)
    c[ib:ib+n] = price * cfg.dt
    c[ic:ic+n] = cfg.eps_cyc * cfg.dt
    c[idis:idis+n] = cfg.eps_cyc * cfg.dt
    c[isu:isu+n] = cfg.eps_sur * cfg.dt
    c[ie+n-1] = -mu
    bounds = ([(0,None)]*n + [(0,cfg.p_max)]*n + [(0,cfg.p_max)]*n +
              [(0,None)]*n + [(cfg.e_min,cfg.e_max)]*n)
    return LPTemplates(Aeq, c, bounds, mu)
def solve_day_ahead(price: np.ndarray,
                    load_fc: np.ndarray,
                    pv_fc: np.ndarray,
                    reserve: np.ndarray,
                    soc0: float,
                    cfg: Config,
                    lp: LPTemplates) -> Dict[str,np.ndarray]:
    """每天0:00求解一次全天计划购电及计划储能LP。"""
    n = cfg.slots_per_day
    net_req = load_fc[:n] - pv_fc[:n] + reserve
    b = np.empty(2*n, dtype=float)
    b[:n] = net_req
    b[n:] = 0.0
    b[n] = soc0
    res = linprog(lp.da_c, A_eq=lp.da_Aeq, b_eq=b, bounds=lp.da_bounds, method=cfg.lp_met)
    if not res.success:
        raise RuntimeError(f"日前LP失败: {res.message}")
    x = res.x
    ib, ic, idis, isu, ie = 0,n,2*n,3*n,4*n
    paper_obj = float(res.fun + lp.mu_terminal * soc0)
    return {"buy":x[ib:ib+n], "ch":x[ic:ic+n], "dis":x[idis:idis+n],
            "sur":x[isu:isu+n], "E":x[ie:ie+n], "obj":paper_obj}
def operate_one_slot(soc0: float, buy: float, load_actual: float, pv_actual: float,
                     plan_ch: float, plan_dis: float, cfg: Config) -> Dict[str,float]:
    dt = cfg.dt
    # 先把计划动作投影到当前SOC可行范围
    ch = max(0.0, min(float(plan_ch), cfg.p_max))
    dis = max(0.0, min(float(plan_dis), cfg.p_max))
    max_ch_soc = max(0.0, (cfg.e_max - soc0) / (cfg.eta_c * dt))
    max_dis_soc = max(0.0, (soc0 - cfg.e_min) * cfg.eta_d / dt)
    ch = min(ch, max_ch_soc)
    dis = min(dis, max_dis_soc)
    # 基于真实负荷/光伏检查功率平衡
    gap = float(load_actual - (buy + pv_actual + dis - ch))
    emg = 0.0
    sur = 0.0
    if gap > 1e-10:
        # 1. 先取消充电，相当于立即释放可用供电
        cut_ch = min(ch, gap)
        ch -= cut_ch
        gap -= cut_ch
        # 2. 再增加放电；紧急购电价格为5倍，故在可行范围内优先用储能
        max_dis_total = min(cfg.p_max, max(0.0, (soc0 - cfg.e_min) * cfg.eta_d / dt))
        extra_dis = min(gap, max(0.0, max_dis_total - dis))
        dis += extra_dis
        gap -= extra_dis
        # 3. 储能仍不足时才紧急购电
        emg = max(0.0, gap)
    elif gap < -1e-10:
        excess = -gap
        # 1. 先取消不必要的计划放电
        cut_dis = min(dis, excess)
        dis -= cut_dis
        excess -= cut_dis
        # 2. 再增加充电
        max_ch_total = min(cfg.p_max, max(0.0, (cfg.e_max - soc0) / (cfg.eta_c * dt)))
        extra_ch = min(excess, max(0.0, max_ch_total - ch))
        ch += extra_ch
        excess -= extra_ch
        # 3. 电池无法继续吸收的部分作为富余功率
        sur = max(0.0, excess)
    soc1 = soc0 + (cfg.eta_c*ch - dis/cfg.eta_d)*dt
    soc1 = min(cfg.e_max, max(cfg.e_min, soc1))
    return {"ch":ch, "dis":dis, "emg":emg, "sur":sur, "E":soc1}
# 结果
def fmt_slot(k):
    mins = k*10
    return "24:00" if mins == 1440 else f"{mins//60}:{mins%60:02d}"
def fill_result2(template_path:str, out_path:str, dates, price,
                 plan_buy, ch, dis, soc_start, soc_end, emg, cfg:Config, logger):
    from copy import copy
    wb = load_workbook(template_path)
    sp = wb["计划购电量"]
    sc = wb["充放电量"]
    se = wb["紧急购电量"]
    date_to_idx={d:i for i,d in enumerate(dates)}
    first_submit = date(2025, 2, 1)
    computed_dates=[d for d in dates if d >= first_submit and d in date_to_idx
                    and np.isfinite(plan_buy[date_to_idx[d]]).all()]
    # 计划购电量
    for r in range(2, sp.max_row + 1):
        raw_date = sp.cell(r,1).value
        if raw_date is None:
            continue
        try:
            d=cell_to_date(raw_date)
        except Exception:
            continue
        if d not in date_to_idx:
            continue
        i=date_to_idx[d]
        if not np.isfinite(plan_buy[i]).all():
            continue
        vals=(plan_buy[i]*cfg.dt).tolist()  # kWh/10min
        for c, v in enumerate(vals, start=2):
            sp.cell(r,c).value = None if not np.isfinite(v) else float(v)
        sp.cell(r,146).value=float(np.nansum(plan_buy[i])*cfg.dt)            # EP 全天购电量
        sp.cell(r,147).value=float(np.nansum(plan_buy[i]*price)*cfg.dt)      # EQ 全天购电费
    # 保存原模板样式作为“样板行”
    charge_proto=[{"height": sc.row_dimensions[rr].height,
                   "cells": [sc.cell(rr,cc) for cc in range(1,7)]} for rr in range(2,8)]
    emg_proto=[{"height": se.row_dimensions[rr].height,
                "cells": [se.cell(rr,cc) for cc in range(1,4)]} for rr in range(2,5)]
    def copy_cell_style(src, dst):
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
    # 充放电量
    if sc.max_row >= 2:
        sc.delete_rows(2, sc.max_row-1)
    intervals=[(0,4,"0:00-4:00"),(4,8,"4:00-8:00"),(8,12,"8:00-12:00"),
               (12,16,"12:00-16:00"),(16,20,"16:00-20:00"),(20,24,"20:00-24:00")]
    r=2
    for d in computed_dates:
        i=date_to_idx[d]
        for k,(h0,h1,label) in enumerate(intervals):
            proto=charge_proto[k]
            sc.row_dimensions[r].height = proto["height"]
            for cc in range(1,7):
                copy_cell_style(proto["cells"][cc-1], sc.cell(r,cc))
            sc.cell(r,1).value = d if k==0 else None
            sc.cell(r,2).value = label
            s0=h0*6; s1=h1*6
            sc.cell(r,3).value=float(np.nansum(ch[i,s0:s1])*cfg.dt)
            sc.cell(r,4).value=float(np.nansum(dis[i,s0:s1])*cfg.dt)
            sc.cell(r,5).value = ("0:00" if k==0 else "24:00" if k==1 else None)
            sc.cell(r,6).value = (float(soc_start[i]) if k==0 else
                                  float(soc_end[i]) if k==1 else None)
            r += 1
    # 紧急购电
    if se.max_row >= 2:
        se.delete_rows(2, se.max_row-1)
    r=2
    for d in computed_dates:
        i=date_to_idx[d]
        events=[]; t=0
        while t<144:
            if emg[i,t] > 1e-6:
                st=t; energy=0.0
                while t<144 and emg[i,t] > 1e-6:
                    energy += emg[i,t]*cfg.dt
                    t += 1
                events.append((f"{fmt_slot(st)}-{fmt_slot(t)}", float(energy)))
            else:
                t += 1
        if not events:
            events=[(None,None)]
        for k,(period,energy) in enumerate(events):
            proto=emg_proto[min(k, len(emg_proto)-1)]
            se.row_dimensions[r].height = proto["height"]
            for cc in range(1,4):
                copy_cell_style(proto["cells"][cc-1], se.cell(r,cc))
            se.cell(r,1).value = d if k==0 else None
            se.cell(r,2).value = period
            se.cell(r,3).value = energy
            r += 1
    # 日期与数值格式统一
    sc.column_dimensions['A'].width = max(sc.column_dimensions['A'].width or 10, 12)
    sc.column_dimensions['B'].width = max(sc.column_dimensions['B'].width or 12, 15)
    se.column_dimensions['A'].width = max(se.column_dimensions['A'].width or 10, 12)
    se.column_dimensions['B'].width = max(se.column_dimensions['B'].width or 12, 18)
    for ws in (sc, se):
        for rr in range(2, ws.max_row+1):
            if ws.cell(rr,1).value is not None:
                ws.cell(rr,1).number_format='yyyy/m/d'
        for cc in ((3,4,6) if ws is sc else (3,)):
            for rr in range(2, ws.max_row+1):
                ws.cell(rr,cc).number_format='0.00'
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    wb.close()
def write_specified_date_results(dates, price, plan_buy, ch, dis, emg,
                                 soc_start, soc_end, cfg:Config, logger):
    targets=[date(2025,3,20), date(2025,6,21), date(2025,9,23), date(2025,12,21)]
    date_to_idx={d:i for i,d in enumerate(dates)}
    out_dir=Path(cfg.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    slot_hours=[10,12,14,16,18,20]
    intvs=[(0,4,'0:00-4:00'),(4,8,'4:00-8:00'),(8,12,'8:00-12:00'),
           (12,16,'12:00-16:00'),(16,20,'16:00-20:00'),(20,24,'20:00-24:00')]
    rows1=[]; rows2=[]; rows3=[]
    for d in targets:                     # 一次性计算，供Excel复用
        if d not in date_to_idx: continue
        i=date_to_idx[d]
        if not np.isfinite(plan_buy[i]).all(): continue
        vals=[float(plan_buy[i,h*6]*cfg.dt) for h in slot_hours]
        totE=float(np.nansum(plan_buy[i])*cfg.dt)
        totC=float(np.nansum(plan_buy[i]*price)*cfg.dt)
        rows1.append([str(d),*vals,totE,totC])
        for k,(h0,h1,label) in enumerate(intvs):
            s0,s1=h0*6,h1*6
            rows2.append([str(d) if k==0 else '', label,
                          float(np.nansum(ch[i,s0:s1])*cfg.dt),
                          float(np.nansum(dis[i,s0:s1])*cfg.dt),
                          float(soc_start[i]) if k==0 else '',
                          float(soc_end[i]) if k==0 else ''])
        events=[]; t=0                    # 连续10min紧急购电时段自动合并
        while t<144:
            if emg[i,t] > 1e-6:
                st=t; energy=0.0
                while t<144 and emg[i,t] > 1e-6:
                    energy += emg[i,t]*cfg.dt; t += 1
                events.append((f'{fmt_slot(st)}-{fmt_slot(t)}', float(energy)))
            else:
                t += 1
        for k,(period,energy) in enumerate(events if events else [('无紧急购电',0.0)]):
            rows3.append([str(d) if k==0 else '', period, energy])
    # 表1/表2/表3  Excel
    from openpyxl import Workbook
    wb_out = Workbook()
    sheets = [
        ('表1计划购电',
         ['日期','10:00-10:10/kWh','12:00-12:10/kWh','14:00-14:10/kWh',
          '16:00-16:10/kWh','18:00-18:10/kWh','20:00-20:10/kWh',
          '全天购电量/kWh','全天购电费/元'], rows1),
        ('表2充放电',
         ['日期','时间段','充电量/kWh','放电量/kWh','0:00储电量/kWh',
          '24:00储电量/kWh'], rows2),
        ('表3紧急购电', ['日期','紧急购电时间段','紧急购电量/kWh'], rows3),
    ]
    for k,(title,hdr,rows) in enumerate(sheets):
        ws = wb_out.active if k==0 else wb_out.create_sheet(title)
        ws.title = title
        ws.append(hdr)
        for row in rows:
            ws.append([round(v,4) if isinstance(v,(int,float)) else v for v in row])
    px = out_dir/'Q2_指定日期表格数据.xlsx'
    wb_out.save(px)
# 可视化与结果摘要
C_BUY, C_EM, C_LOAD   = '#2E74B5', '#E74C3C', '#34495E'   # 蓝/红/深灰蓝
C_PV, C_CH, C_DIS     = '#27AE60', '#E74C3C', '#E67E22'   # 绿/红/橙
C_SUR, C_SOC, C_ALPHA = '#95A5A6', '#7F8C8D', '#8E44AD'   # 灰/灰/紫
C_RES, C_MAX, C_MIN   = '#16A085', '#C0392B', '#2980B9'   # 青/深红/蓝
C_NET, C_SUR_GREEN    = '#8E44AD', '#2ECC71'   # 紫 / 富余(亮绿,月度组合图)
XTICKS_DAY = list(range(0, 25, 2))
# 论文图样
NO_TITLE  = True   # 图内不带标题（论文中标题由图注给出）
AXIS_FS   = 15     # 坐标轴标签字号（黑色）
TICK_FS   = 12     # 刻度字号
LEGEND_FS = 17     # 图例字号
SOFT_RED, SOFT_GREEN, SOFT_BLUE = '#E9A6A1', '#8BCF8B', '#7097CA'  # figures4papers 柔和色
BAR_EDGE  = dict(edgecolor='black', linewidth=1.2)
def _style_axes(ax):
    #绘图
    ax.spines['top'].set_visible(False)
    ax.grid(True, axis='y', linestyle='--', linewidth=0.7, alpha=0.3)
    ax.tick_params(labelsize=TICK_FS)
def _limit_annot(ax, yv, color, text):
    ax.axhline(yv, linestyle='--', linewidth=1.2, color=color, alpha=0.55)
    ax.text(1.005, yv, text, transform=ax.get_yaxis_transform(),
            fontsize=8.5, color=color, alpha=0.85, va='center', ha='left')
def _top_legend(ax, ncol=None, **kw):
    h, l = ax.get_legend_handles_labels()
    if not h:
        return
    ax.legend(h, l, loc='lower center', frameon=False,
              ncol=ncol or max(1, len(h)), fontsize=LEGEND_FS,
              bbox_to_anchor=(0.5, 1.02), columnspacing=1.8, **kw)
def _save_line_fig(x, ys, labels, title, xlabel, ylabel, path, date_axis=False,
                   hlines=None, ylim=None, colors=None, line_styles=None,
                   xticks=None, xlim=None, linewidths=None, hline_annot=False,
                   fill_under=None):
    fig, ax = plt.subplots(figsize=(10.8, 5.4), dpi=300)
    for k, (y, lab) in enumerate(zip(ys, labels)):
        c = colors[k] if colors else None
        ls = line_styles[k] if line_styles else '-'
        lw = linewidths[k] if linewidths else 2.2
        ax.plot(x, y, linewidth=lw, label=lab, color=c, linestyle=ls)
    if fill_under is not None:
        yi, base, fc, fa = fill_under
        ax.fill_between(x, ys[yi], base, color=fc, alpha=fa, linewidth=0)
    if hlines:
        for item in hlines:
            yv, lab = item[0], item[1]
            hc = item[2] if len(item) > 2 else '#7F8C8D'
            if hline_annot:
                _limit_annot(ax, yv, hc, lab)
            else:
                ax.axhline(yv, linestyle='--', linewidth=1.0, color=hc, alpha=0.55, label=lab)
    if not NO_TITLE:
        ax.set_title(title, fontsize=13, fontweight='bold', pad=28)
    ax.set_xlabel(xlabel, fontsize=AXIS_FS, color='black')
    ax.set_ylabel(ylabel, fontsize=AXIS_FS, color='black')
    _style_axes(ax)
    _top_legend(ax)
    if date_axis:
        ax.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=5, maxticks=9))
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'))
        fig.autofmt_xdate(rotation=25)
    if xticks is not None:
        ax.set_xticks(xticks)
    if xlim is not None:
        ax.set_xlim(*xlim)
    if ylim is not None:
        ax.set_ylim(*ylim)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches='tight')
    plt.close(fig)
def _save_bar_fig(xlabels, values, title, xlabel, ylabel, path, color=SOFT_RED):
    """柱状图：figures4papers 柔和色 + 黑描边，无标题、大字号。"""
    fig, ax = plt.subplots(figsize=(10.8, 5.4), dpi=300)
    xx = np.arange(len(values))
    ax.bar(xx, values, width=0.62, color=color, alpha=0.95, zorder=3, **BAR_EDGE)
    ax.set_xticks(xx)
    ax.set_xticklabels(xlabels, fontsize=TICK_FS)
    if not NO_TITLE:
        ax.set_title(title, fontsize=13, fontweight='bold', pad=12)
    ax.set_xlabel(xlabel, fontsize=AXIS_FS, color='black')
    ax.set_ylabel(ylabel, fontsize=AXIS_FS, color='black')
    _style_axes(ax)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches='tight')
    plt.close(fig)
def generate_visualizations(dates, price, load, pv, load_fc_store, pv_fc_store,
                            plan_buy, ch, dis, emg, sur, soc_start, soc_end,
                            alpha_store, reserve_store, run_D,
                            cfg:Config, logger):
    #生成论文结果
    fig_dir=Path(cfg.fig_dir); fig_dir.mkdir(parents=True, exist_ok=True)
    submit_start=date(2025,2,1)
    valid_days=[i for i in range(run_D)
                if dates[i] >= submit_start and np.isfinite(plan_buy[i]).all()]
    if not valid_days:
        logger.warning("尚无2月及以后的有效结果，已跳过可视化。")
        return
    xs=[datetime.combine(dates[i], datetime.min.time()) for i in valid_days]
    # 1) 2月起日末SOC轨迹（fig4 风格：蓝粗线+浅填充、右缘限值标注）
    fig, ax = plt.subplots(figsize=(10.8, 5.4), dpi=300)
    ax.fill_between(xs, soc_end[valid_days], cfg.e_min,
                    color=C_BUY, alpha=0.10, linewidth=0)
    ax.plot(xs, soc_end[valid_days], color=C_BUY, linewidth=1.8, label='日末储电量')
    _limit_annot(ax, cfg.e_min, C_MIN, f'下限 {cfg.e_min:.0f}')
    _limit_annot(ax, cfg.e_max, C_MAX, f'上限 {cfg.e_max:.0f}')
    if not NO_TITLE:
        ax.set_title('2月起储能设备日末电量轨迹', fontsize=13, fontweight='bold', pad=28)
    ax.set_xlabel('日期', fontsize=AXIS_FS, color='black')
    ax.set_ylabel('储电量 / kWh', fontsize=AXIS_FS, color='black')
    _style_axes(ax)
    _top_legend(ax)
    ax.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=5, maxticks=9))
    ax.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'))
    fig.autofmt_xdate(rotation=25)
    ax.set_ylim(max(0, cfg.e_min-500), cfg.e_max+500)
    fig.tight_layout()
    fig.savefig(fig_dir/"第二问储能设备日末电量轨迹.png", dpi=300, bbox_inches='tight')
    plt.close(fig)
    # 全年每日预测RMSE（图18与月度RMSE汇总共用）
    vd=np.array(valid_days)
    Lf, Pf = load_fc_store[vd,:144], pv_fc_store[vd,:144]
    dl=np.sqrt(np.mean((load[vd]-Lf)**2, axis=1))
    dp=np.sqrt(np.mean((pv[vd]-Pf)**2, axis=1))
    dn=np.sqrt(np.mean(((load[vd]-pv[vd])-(Lf-Pf))**2, axis=1))
    month_of=np.array([dates[i].month for i in valid_days])
    # 按实际存在月份汇总（正式运行即2--12月）
    months_sorted=sorted(set(month_of.tolist()))
    month_labels=[f"{m}月" for m in months_sorted]
    plan_cost=[]; emg_cost=[]; emg_energy=[]; sur_energy=[]
    rmse_load=[]; rmse_pv=[]; rmse_net=[]
    for m in months_sorted:
        ids=[i for i in valid_days if dates[i].month==m]
        plan_cost.append(sum(float(np.sum(plan_buy[i]*price)*cfg.dt) for i in ids))
        emg_cost.append(sum(float(np.sum(emg[i]*5*price)*cfg.dt) for i in ids))
        emg_energy.append(sum(float(np.sum(emg[i])*cfg.dt) for i in ids))
        sur_energy.append(sum(float(np.sum(sur[i])*cfg.dt) for i in ids))
        sel=month_of==m                   # 非有限值日剔除后再取均值（与原逻辑等价）
        rmse_load.append(float(np.mean(dl[sel & np.isfinite(dl)]))
                         if np.isfinite(dl[sel]).any() else np.nan)
        rmse_pv.append(float(np.mean(dp[sel & np.isfinite(dp)]))
                       if np.isfinite(dp[sel]).any() else np.nan)
        rmse_net.append(float(np.mean(dn[sel & np.isfinite(dn)]))
                        if np.isfinite(dn[sel]).any() else np.nan)
    # 2) 月度购电费用：堆叠展示费用组成
    fig,ax=plt.subplots(figsize=(10.8,5.4),dpi=300)
    xx=np.arange(len(months_sorted))
    ax.bar(xx,plan_cost,width=0.62,color=SOFT_GREEN,alpha=0.95,label="计划购电费",zorder=3,**BAR_EDGE)
    ax.bar(xx,emg_cost,width=0.62,bottom=plan_cost,color=SOFT_RED,alpha=0.95,label="紧急购电费",zorder=3,**BAR_EDGE)
    ax.set_xticks(xx); ax.set_xticklabels(month_labels,fontsize=TICK_FS)
    if not NO_TITLE:
        ax.set_title("2月起月度购电费用构成",fontsize=13,fontweight='bold',pad=12)
    ax.set_xlabel("月份",fontsize=AXIS_FS,color='black'); ax.set_ylabel("费用 / 元",fontsize=AXIS_FS,color='black')
    _style_axes(ax); _top_legend(ax, ncol=2)
    fig.tight_layout(); fig.savefig(fig_dir/"第二问月度购电费用构成.png",dpi=300,bbox_inches='tight'); plt.close(fig)
    # 3) 月度紧急购电量（柔粉红柱）
    _save_bar_fig(month_labels,emg_energy,"2月起月度紧急购电量","月份","紧急购电量 / kWh",
                  fig_dir/"第二问月度紧急购电量.png",color=SOFT_RED)
    from openpyxl import Workbook
    wb_rm = Workbook()
    ws_rm = wb_rm.active; ws_rm.title = '月度预测RMSE'
    ws_rm.append(['指标'] + month_labels)
    ws_rm.append(['负荷RMSE (kW)'] + [round(float(v), 2) for v in rmse_load])
    ws_rm.append(['光伏RMSE (kW)'] + [round(float(v), 2) for v in rmse_pv])
    ws_rm.append(['净负荷RMSE (kW)'] + [round(float(v), 2) for v in rmse_net])
    px_rm = fig_dir.parent / 'Q2_月度预测RMSE.xlsx'
    wb_rm.save(px_rm)
    #代表日优先选题目指定日期3月20日；若尚未计算则选最后一个有效日
    target=date(2025,3,20)
    sample_idx=next((i for i in valid_days if dates[i]==target),valid_days[-1])
    tt=np.arange(144)/6.0
    compare_days=[]
    for td in [date(2025,3,20),date(2025,6,21),
               date(2025,9,23),date(2025,12,21)]:
        idx=next((i for i in valid_days if dates[i]==td),None)
        if idx is not None:
            compare_days.append((idx,
                f"第二问{td.month}-{td.day}负荷与光伏预测对比.png",
                f"{td} 负荷与光伏预测对比"))
    for idx,fname,title in compare_days:
        _save_line_fig(tt,[load[idx],load_fc_store[idx,:144],pv[idx],pv_fc_store[idx,:144]],
                       ["真实负荷","预测负荷","真实光伏","预测光伏"],title,
                       "时刻 / h","功率 / kW",fig_dir/fname,
                       colors=[C_LOAD,C_BUY,C_PV,C_DIS],line_styles=['-','--','-','--'],
                       linewidths=[2.0,1.8,2.0,1.8],
                       xticks=XTICKS_DAY, xlim=(0,24))
        _save_line_fig(tt,[alpha_store[sample_idx]],
            ["动态 alpha"],
            f"{dates[sample_idx]} 动态非对称风险权重",
            "时刻 / h",
            "alpha",
            fig_dir / "第二问3-20动态非对称风险权重.png",
            colors=[C_ALPHA],
            xticks=XTICKS_DAY,
            xlim=(0, 24),
            ylim=(0, max(cfg.alpha_max + 0.5, 5)))
        _save_line_fig(tt,[reserve_store[sample_idx]],
            ["风险备用量"],
            f"{dates[sample_idx]} 风险备用量",
            "时刻 / h",
            "功率 / kW",
            fig_dir / "第二问3-20风险备用量.png",
            colors=[C_RES],
            xticks=XTICKS_DAY,
            xlim=(0, 24))

    soc_curve=np.empty(145); soc_curve[0]=soc_start[sample_idx]
    cur=soc_curve[0]
    for t in range(144):
        cur=cur+(cfg.eta_c*ch[sample_idx,t]-dis[sample_idx,t]/cfg.eta_d)*cfg.dt
        soc_curve[t+1]=cur
    #代表日SOC轨迹
    fig, ax = plt.subplots(figsize=(11, 5.5), dpi=300)
    t_soc = np.arange(145)/6.0
    ax.fill_between(t_soc, soc_curve, cfg.e_min, where=soc_curve >= cfg.e_min,
                    alpha=0.10, color=C_DIS, linewidth=0)
    ax.plot(t_soc, soc_curve, color=C_SOC, linewidth=2.8, marker='o', markersize=2.5,
            label='储电量', zorder=10)
    _limit_annot(ax, cfg.e_min, C_MIN, f'下限 {cfg.e_min:.0f}')
    _limit_annot(ax, cfg.e_max, C_MAX, f'上限 {cfg.e_max:.0f}')
    if not NO_TITLE:
        ax.set_title(f"{dates[sample_idx]} 储能设备电量变化", fontsize=13, fontweight='bold', pad=28)
    ax.set_xlabel('时刻 / h', fontsize=AXIS_FS, color='black')
    ax.set_ylabel('储电量 / kWh', fontsize=AXIS_FS, color='black')
    _style_axes(ax)
    ax.spines['left'].set_linewidth(1.5)
    h, l = ax.get_legend_handles_labels()
    ax.legend(h, l, loc='upper right', frameon=False, fontsize=15,
              bbox_to_anchor=(1.0, 0.90))
    ax.set_xticks(XTICKS_DAY); ax.set_xlim(0, 24)
    ax.set_ylim(max(0, cfg.e_min-500), cfg.e_max+500)
    fig.tight_layout()
    fig.savefig(fig_dir/"第二问3-20储能设备电量变化.png", dpi=300, bbox_inches='tight')
    plt.close(fig)
    a_ld,p_ld = np.mean(load[vd],axis=1), np.mean(load_fc_store[vd,:144],axis=1)
    a_pv,p_pv = np.mean(pv[vd],axis=1), np.mean(pv_fc_store[vd,:144],axis=1)
    for ys,lbls,ttl,fname,cols in (
            ([a_ld,p_ld],["实际日均负荷","预测日均负荷"],"2月起全年负荷预测结果对比",
             "第二问全年负荷预测结果对比.png",[C_LOAD,C_DIS]),
            ([a_pv,p_pv],["实际日均光伏","预测日均光伏"],"2月起全年光伏预测结果对比",
             "第二问全年光伏预测结果对比.png",[C_PV,C_DIS])):
        _save_line_fig(xs,ys,lbls,ttl,"日期","日均功率 / kW",fig_dir/fname,
                       date_axis=True,colors=cols,line_styles=['-','--'],
                       linewidths=[2.0,1.8])
    #风险强度系数
    rho_list   = np.array([0.70, 0.75, 0.80, 0.85, 0.90])
    total_wan  = np.array([1366.92, 1365.86, 1366.63, 1368.51, 1371.41])  # 万元
    best_i     = int(np.argmin(total_wan))
    fig, ax = plt.subplots(figsize=(10.8, 5.4), dpi=300)
    ax.plot(rho_list, total_wan * 1e4, color=C_BUY, linewidth=2.4,
            marker='o', markersize=9, markerfacecolor=C_BUY,
            markeredgecolor='white', markeredgewidth=1.5, label='全年总购电费', zorder=5)
    ax.plot(rho_list[best_i], total_wan[best_i] * 1e4, marker='*', markersize=22,
            color=C_CH, markeredgecolor=C_CH, linestyle='none',
            label=f'最优 $\\rho$={rho_list[best_i]:.2f}', zorder=6)
    ax.axvline(rho_list[best_i], linestyle='--', linewidth=1.2, color=C_CH, alpha=0.45)
    for x0, y0 in zip(rho_list, total_wan):
        ax.annotate(f'{y0:.2f}万', (x0, y0 * 1e4),
                    textcoords='offset points', xytext=(0, 14),
                    ha='center', fontsize=10.5,
                    color=C_CH if x0 == rho_list[best_i] else '#444444')
    if not NO_TITLE:
        ax.set_title(r'风险强度系数 $\rho$ 灵敏度分析（全年总购电费）',
                     fontsize=13, fontweight='bold', pad=28)
    ax.set_xlabel(r'风险强度系数 $\rho$（reserve_scale）', fontsize=AXIS_FS, color='black')
    ax.set_ylabel('全年总购电费用 / 元', fontsize=AXIS_FS, color='black')
    ax.set_xticks(rho_list)
    ax.tick_params(labelsize=TICK_FS)
    _style_axes(ax)
    hd, lb = ax.get_legend_handles_labels()
    ax.legend(hd, lb, loc='lower center', frameon=False, ncol=2,
              fontsize=LEGEND_FS, bbox_to_anchor=(0.5, 1.02), columnspacing=1.8)
    fig.tight_layout()
    fig.savefig(fig_dir/"第二问全年总购电费用随风险强度系数的变化.png",
                dpi=300, bbox_inches='tight')
    plt.close(fig)
    # 缓存绘图数据
    try:
        np.savez_compressed(
            Path(cfg.out_dir)/"q2_plot_data.npz",
            date_str=np.array([str(dates[i]) for i in range(run_D)]),
            price=price, load=load[:run_D], pv=pv[:run_D],
            load_fc=load_fc_store[:run_D], pv_fc=pv_fc_store[:run_D],
            plan_buy=plan_buy[:run_D], ch=ch[:run_D], dis=dis[:run_D],
            emg=emg[:run_D], sur=sur[:run_D],
            soc_start=soc_start[:run_D], soc_end=soc_end[:run_D],
            alpha=alpha_store[:run_D], reserve=reserve_store[:run_D],
            e_min=cfg.e_min, e_max=cfg.e_max, e0=cfg.e0,
            eta_c=cfg.eta_c, eta_d=cfg.eta_d, dt=cfg.dt,
        )
    except Exception as exc:
        logger.warning("绘图数据缓存失败（不影响图片）：%s", exc)
def write_result_summary(dates, price, plan_buy, emg, sur, ch, dis, soc_end,
                         alpha_store, run_D, cfg:Config, logger):
    #从2月1日开始计算
    submit_start=date(2025,2,1)
    valid=[i for i in range(run_D)
           if dates[i] >= submit_start and np.isfinite(plan_buy[i]).all()]
    if not valid:
        return
    plan_cost=sum(float(np.sum(plan_buy[i]*price)*cfg.dt) for i in valid)
    emg_cost=sum(float(np.sum(emg[i]*5*price)*cfg.dt) for i in valid)
    emg_energy=sum(float(np.sum(emg[i])*cfg.dt) for i in valid)
    sur_energy=sum(float(np.sum(sur[i])*cfg.dt) for i in valid)
    charge_energy=sum(float(np.sum(ch[i])*cfg.dt) for i in valid)
    discharge_energy=sum(float(np.sum(dis[i])*cfg.dt) for i in valid)
    alpha_mean=float(np.nanmean(alpha_store[valid]))
    zero_emg_days=sum(1 for i in valid if np.sum(emg[i])*cfg.dt <= 1e-8)
    lines=["问题2正式结果摘要（1月仅作为预热期，不纳入统计）",
        f"正式统计日期范围：{dates[valid[0]]} 至 {dates[valid[-1]]}",
        f"累计计划购电费：{plan_cost:.2f} 元",
        f"累计紧急购电费：{emg_cost:.2f} 元",
        f"累计总购电费：{plan_cost+emg_cost:.2f} 元",
        f"累计紧急购电量：{emg_energy:.2f} kWh",
        f"累计富余电量：{sur_energy:.2f} kWh",
        f"累计充电量：{charge_energy:.2f} kWh",
        f"累计放电量：{discharge_energy:.2f} kWh",
        f"正式期末储电量：{soc_end[valid[-1]]:.2f} kWh",
        f"动态alpha平均值：{alpha_mean:.4f}",
        f"无紧急购电天数：{zero_emg_days}/{len(valid)}",
        f"正式提交文件：{cfg.out_xlsx}",
        f"可视化目录：{cfg.fig_dir}",]
    path=Path(cfg.out_dir)/"Q2_结果摘要.txt"
    path.write_text("\n".join(lines),encoding="utf-8")
# 主程序
def main():
    cfg = CFG
    out_dir = Path(cfg.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    logger = setup_logger(out_dir)
    dates, price, load_prior, pv_prior, load, pv = load_inputs(cfg, logger)
    D = len(dates)
    nT = cfg.slots_per_day
    run_D = min(D, cfg.max_days) if cfg.max_days > 0 else D
    plan_buy, soc_hist = (np.full((D,nT), np.nan) for _ in range(2))
    plan_ch, plan_dis, ch, dis, emg, sur, reserve_store = (np.zeros((D,nT)) for _ in range(7))
    soc_start, soc_end = np.full(D, np.nan), np.full(D, np.nan)
    load_fc_store, pv_fc_store = (np.full((D,cfg.forecast_slots), np.nan) for _ in range(2))
    alpha_store, net_errors, cplus_cache, cminus_cache = (np.full((D,nT), np.nan) for _ in range(4))
    lp_templates = build_lp_templates(price, cfg)
    daily_rows = []
    soc = cfg.e0
    for d in range(run_D):
        soc_start[d] = soc
        # 0:00预测：只使用已经发生的历史数据
        lf = load_forecast(d, load, load_prior, cfg, logger)
        pf, pvinfo = pv_forecast(d, pv, pv_prior, cfg, logger)
        load_fc_store[d] = lf
        pv_fc_store[d] = pf
        #动态非对称风险备用量
        update_risk_cost_cache(d, cplus_cache, cminus_cache, net_errors, price,
                               soc_hist, plan_buy, load, pv, cfg)
        alpha, beta, reserve = dynamic_risk_weights(d, net_errors, cplus_cache, cminus_cache, cfg)
        alpha_store[d] = alpha
        reserve_store[d] = reserve
        #当天只求解一次日前LP
        da = solve_day_ahead(price, lf, pf, reserve, soc, cfg, lp_templates)
        plan_buy[d] = da["buy"]; plan_ch[d] = da["ch"]
        plan_dis[d] = da["dis"]
        #实际运行
        for t in range(nT):
            soc_hist[d,t] = soc
            sol = operate_one_slot(soc, plan_buy[d,t], load[d,t], pv[d,t],
                plan_ch[d,t], plan_dis[d,t], cfg)
            ch[d,t] = sol["ch"]; dis[d,t] = sol["dis"]; emg[d,t] = sol["emg"]
            sur[d,t] = sol["sur"]
            soc = sol["E"]
        soc_end[d] = soc
        #当日预测误差只在当天结束后进入历史，用于次日风险更新
        net_errors[d] = (load[d]-pv[d]) - (lf[:nT]-pf[:nT])
        plan_cost = float(np.sum(plan_buy[d]*price)*cfg.dt)
        emg_cost = float(np.sum(emg[d]*5*price)*cfg.dt)
        daily_rows.append([plan_cost, emg_cost, plan_cost+emg_cost,
                           float(np.sum(emg[d])*cfg.dt), float(np.sum(sur[d])*cfg.dt)])
    if run_D >= 32:
        fill_result2(cfg.template,cfg.out_xlsx,dates,price,plan_buy,ch,dis,soc_start,soc_end,emg,cfg,logger)
    else:
        logger.warning("本次尚未覆盖2月1日，不生成正式 result2。")
    # 题目指定日期若已计算，则单独导出
    try:
        write_specified_date_results(dates,price,plan_buy,ch,dis,emg,soc_start,soc_end,cfg,logger)
    except Exception as exc:
        logger.warning("指定日期结果单独导出失败：%s", exc)
    try:
        generate_visualizations(dates,price,load,pv,load_fc_store,pv_fc_store,
                                plan_buy,ch,dis,emg,sur,soc_start,soc_end,
                                alpha_store,reserve_store,run_D,cfg,logger)
        write_result_summary(dates,price,plan_buy,emg,sur,ch,dis,soc_end,
                             alpha_store,run_D,cfg,logger)
    except Exception as exc:
        logger.error("可视化或结果摘要生成失败，但不影响result2：%s", exc)
    if daily_rows:
        arr = np.asarray(daily_rows, dtype=float)
        print("累计计划购电费 %.2f 元；累计紧急购电费 %.2f 元；累计总费 %.2f 元；"
              "紧急购电量 %.2f kWh；富余电量 %.2f kWh"
              % (arr[:, 0].sum(), arr[:, 1].sum(), arr[:, 2].sum(),
                 arr[:, 3].sum(), arr[:, 4].sum()))
    print("最终 result2：%s" % (cfg.out_xlsx if run_D >= 32 else "本次未生成"))
if __name__ == "__main__":
    main()