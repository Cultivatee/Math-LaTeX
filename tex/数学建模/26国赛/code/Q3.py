from __future__ import annotations
import logging
from dataclasses import dataclass
from pathlib import Path
from datetime import datetime, timedelta, date
from typing import Tuple, Dict, Optional
import numpy as np
from scipy.optimize import linprog
from scipy import sparse
from statsmodels.tsa.ar_model import AutoReg
try:
    from statsmodels.tsa.seasonal import MSTL
    HAS_MSTL = True
except Exception:
    HAS_MSTL = False
try:
    from openpyxl import load_workbook
except ImportError as e:
    raise ImportError(
        "缺少 openpyxl。请在当前 Python 环境执行：pip install openpyxl"
    ) from e
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Arial Unicode MS", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["font.family"] = "sans-serif"
plt.rcParams["font.size"] = 12
plt.rcParams["axes.spines.top"] = False
plt.rcParams["axes.linewidth"] = 1.5
plt.rcParams["xtick.major.width"] = 1.5
plt.rcParams["ytick.major.width"] = 1.5
plt.rcParams["xtick.major.size"] = 6
plt.rcParams["ytick.major.size"] = 6
plt.rcParams["xtick.direction"] = "out"
plt.rcParams["ytick.direction"] = "out"
C_BUY, C_PV, C_DIS   = "#2E74B5", "#27AE60", "#E67E22"
C_LOAD, C_CH, C_W    = "#34495e", "#E74C3C", "#95A5A6"
C_SOC, C_PRICE, C_MIN = "#7F8C8D", "#E74C3C", "#2980B9"
#参数区
class Config:
    base_dir = str(Path(__file__).resolve().parent)
    att1 = str(Path(base_dir) / "附件1.xlsx")
    att2 = str(Path(base_dir) / "附件2.xlsx")
    att3 = str(Path(base_dir) / "附件3.xlsx")
    _tpl_candidates = [
        Path(base_dir) / "附件5" / "result3.xlsx",
        Path(base_dir) / "result3.xlsx",
    ]
    template = str(next(p for p in _tpl_candidates if p.exists()))
    out_dir = str(Path(base_dir) / "Q3结果")
    out_xlsx = str(Path(out_dir) / "result3_filled_forecast_driven.xlsx")
    fig_dir = str(Path(out_dir) / "Q3图片输出")
    slots_pday = 144; dt = 1.0/6.0; hori_slots = 144; fore_slots = 288
    eta_c = eta_d = 0.90; e_min, e_max = 1200.0, 10800.0; p_max = 5000.0; e0 = 6000.0
    l_days = 42; pv_days = 60; l_min_days = 21; m_days = 7
    ar_candi = (1, 2, 3, 6, 12, 18, 36); pv_candi = (1, 2, 3, 6, 12)
    k_candi = (1, 2, 3, 5, 7); clear_quant = 0.93
    pv_smooth = 2; rk_hisdays = 28; rk_decay = 0.90; rk_sam = 5
    alpha_cold = 4.0; alpha_min = 1.0; alpha_max = 8.0; beta_fixed = 1.0; c_eps = 1e-4
    tau_min = 0.55; tau_max = 0.88; reserve_scale = 0.75
    roll_slots = 144; use_bias = False
    stage_quant = 0.80; stage_scale = 0.75
    stage_mindays = 7; stage_rhisdays = 35; stage_decay = 0.94
    eps_cycle = 1e-5; eps_surplus = 1e-8; lp_method = "highs"
    update_hours = (6, 12, 18); run_comparison = True
    compare_schemes = ((), (6,), (12,), (18,), (6, 12, 18))
    compare_names = ("仅0:00", "0:00+6:00", "0:00+12:00", "0:00+18:00", "0:00+6:00+12:00+18:00")
    represent_dates = (date(2025,3,20), date(2025,6,21), date(2025,9,23), date(2025,12,21))
    max_days = 365
CFG = Config()
def setup_logger(out_dir: Path) -> logging.Logger:
    out_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("Q3")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    return logger
def excel_serial_to_date(x) -> datetime:
    # Excel 1900日期系统（含1900闰年历史兼容）
    return datetime(1899, 12, 30) + timedelta(days=float(x))
def cell_to_date(x):
    """兼容 openpyxl 返回的 date/datetime、Excel 序列号和常见日期字符串。"""
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    if isinstance(x, (int, float, np.integer, np.floating)):
        return excel_serial_to_date(x).date()
    if isinstance(x, str):
        txt=x.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(txt, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"无法识别日期单元格：{x!r}")
def safe_float_array(x) -> np.ndarray:
    return np.asarray([[float(v) if v is not None else np.nan for v in row] for row in x], dtype=float)
def weighted_quantile(values: np.ndarray, weights: np.ndarray, q: float) -> float:
    values = np.asarray(values, dtype=float)
    weights = np.asarray(weights, dtype=float)
    m = np.isfinite(values) & np.isfinite(weights) & (weights > 0)
    if m.sum() == 0:
        return 0.0
    v, w = values[m], weights[m]
    idx = np.argsort(v)
    v, w = v[idx], w[idx]
    cw = np.cumsum(w) / np.sum(w)
    return float(np.interp(q, cw, v))
def moving_average_circular(x: np.ndarray, half: int) -> np.ndarray:
    if half <= 0:
        return x.copy()
    n = len(x)
    out = np.zeros(n)
    for i in range(n):
        lo = max(0, i-half)
        hi = min(n, i+half+1)
        out[i] = np.mean(x[lo:hi])
    return out
def choose_ar_forecast(series: np.ndarray, horizon: int, candidates: Tuple[int, ...], fallback: float = 0.0) -> np.ndarray:
    y = np.asarray(series, dtype=float)
    y = y[np.isfinite(y)]
    if len(y) < 8 or np.std(y) < 1e-8:
        return np.full(horizon, float(np.mean(y) if len(y) else fallback))
    best = None
    for p in candidates:
        if len(y) <= p + 5:
            continue
        try:
            model = AutoReg(y, lags=p, trend="ct", old_names=False).fit()
            if np.isfinite(model.aic) and (best is None or model.aic < best[0]):
                best = (model.aic, model)
        except Exception:
            continue
    if best is None:
        return np.full(horizon, float(np.mean(y[-min(7, len(y)):])) )
    try:
        pred = np.asarray(best[1].predict(start=len(y), end=len(y)+horizon-1, dynamic=False), dtype=float)
        if np.any(~np.isfinite(pred)):
            raise ValueError("AR forecast nonfinite")
        return pred
    except Exception:
        return np.full(horizon, float(np.mean(y[-min(7, len(y)):])) )
def load_inputs(cfg: Config, logger):
    for p in (cfg.att1, cfg.att2, cfg.template):
        if not Path(p).exists():
            raise FileNotFoundError(f"未找到文件：{p}")
    # 附件1
    wb1 = load_workbook(cfg.att1, data_only=True, read_only=True)
    s1 = wb1["Sheet1"]
    rows1 = list(s1.iter_rows(min_row=2, max_row=145, min_col=2, max_col=4, values_only=True))
    price = np.array([float(r[0]) for r in rows1], dtype=float)
    load_prior = np.array([float(r[1]) for r in rows1], dtype=float)
    pv_prior = np.array([float(r[2]) for r in rows1], dtype=float)
    wb1.close()
    # 附件2
    wb2 = load_workbook(cfg.att2, data_only=True, read_only=True)
    sl = wb2["小区负载"]
    sp = wb2["光伏发电实际功率"]
    load_all = list(sl.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    pv_all = list(sp.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    wb2.close()
    dates = [cell_to_date(r[0]) for r in load_all]
    load_rows = [r[1:] for r in load_all]
    pv_rows = [r[1:] for r in pv_all]
    load = safe_float_array(load_rows)
    pv_actual = safe_float_array(pv_rows)
    assert load.shape == (365, 144), load.shape
    assert pv_actual.shape == (365, 144), pv_actual.shape
    assert len(price) == 144
    return dates, price, load_prior, pv_prior, load, pv_actual
_LOAD_MSTL_CACHE = {}
def load_forecast(day_idx: int, load: np.ndarray, prior: np.ndarray, cfg: Config, logger=None) -> np.ndarray:
    H = cfg.fore_slots
    if day_idx == 0:
        return np.tile(prior, 2)
    hist_days = min(day_idx, cfg.l_days)
    hist_mat = load[day_idx-hist_days:day_idx]
    # 冷启动：同一时刻近期日加权 + 附件1先验
    if hist_days < cfg.l_min_days or not HAS_MSTL:
        k = hist_mat.shape[0]
        weights = np.exp(-0.10*np.arange(k-1, -1, -1))
        weights /= weights.sum()
        shape = np.sum(hist_mat * weights[:, None], axis=0)
        if hist_days >= 7:
            shape = 0.55*shape + 0.45*load[day_idx-7]
        w_prior = max(0.0, 1.0 - hist_days / cfg.l_min_days)
        shape = (1-w_prior)*shape + w_prior*prior
        return np.tile(shape, 2)
    # 仅每隔若干天完整重估一次 MSTL；数据已确认无明显异常，因此关闭 robust 迭代
    need_refit = (
        not _LOAD_MSTL_CACHE
        or day_idx - _LOAD_MSTL_CACHE.get("refit_day", -10**9) >= cfg.m_days
    )
    if need_refit:
        y = hist_mat.reshape(-1)
        try:
            mstl = MSTL(y, periods=(144, 1008), stl_kwargs={"robust": False}).fit()
            trend = np.asarray(mstl.trend)
            seasonal = np.asarray(mstl.seasonal)
            if seasonal.ndim == 1:
                s_daily = seasonal
                s_weekly = np.zeros_like(seasonal)
            else:
                s_daily = seasonal[:, 0]
                s_weekly = seasonal[:, 1]
            resid = np.asarray(mstl.resid)
            # 保存稳定的日周期与周周期结构，供未来几天复用
            daily_tail = s_daily[-min(len(s_daily), 28*144):]
            weekly_tail = s_weekly[-min(len(s_weekly), 6*1008):]
            daily_pattern = np.array([
                np.nanmean(daily_tail[i::144]) for i in range(144)
            ])
            weekly_pattern = np.array([
                np.nanmean(weekly_tail[i::1008]) for i in range(1008)
            ])
            _LOAD_MSTL_CACHE.clear()
            _LOAD_MSTL_CACHE.update({
                "refit_day": day_idx,
                "origin_day": day_idx - hist_days,
                "daily_pattern": daily_pattern,
                "weekly_pattern": weekly_pattern,
                "trend_last": float(trend[-1]),
                "trend_slope": float(np.clip(
                    np.polyfit(
                        np.arange(min(432, len(trend)), dtype=float),
                        trend[-min(432, len(trend)):],
                        1
                    )[0], -5.0, 5.0
                )),
                "resid_tail": resid[-min(len(resid), 14*144):].copy(),
            })
        except Exception as e:
            if logger:
                logger.warning("MSTL失败(day=%d)，回退季节加权预测：%s", day_idx, e)
            k = hist_mat.shape[0]
            weights = np.exp(-0.10*np.arange(k-1, -1, -1))
            weights /= weights.sum()
            shape = np.sum(hist_mat * weights[:, None], axis=0)
            if hist_days >= 7:
                shape = 0.5*shape + 0.5*load[day_idx-7]
            return np.tile(shape, 2)
    # 使用最近一次 MSTL 的季节结构，并用最新 3 天数据快速校正趋势和残差
    daily_pattern = _LOAD_MSTL_CACHE["daily_pattern"]
    weekly_pattern = _LOAD_MSTL_CACHE["weekly_pattern"]
    origin_day = _LOAD_MSTL_CACHE["origin_day"]
    recent_days = min(hist_days, 3)
    recent = load[day_idx-recent_days:day_idx].reshape(-1)
    global_start = (day_idx - recent_days) * 144
    gidx = global_start + np.arange(len(recent))
    sd_hist = daily_pattern[gidx % 144]
    sw_hist = weekly_pattern[(gidx - origin_day*144) % 1008]
    deseason = recent - sd_hist - sw_hist
    m = min(432, len(deseason))
    x = np.arange(m, dtype=float)
    coef = np.polyfit(x, deseason[-m:], 1)
    slope = float(np.clip(coef[0], -5.0, 5.0))
    level = float(deseason[-1])
    trend_future = level + slope*np.arange(1, H+1)
    future_gidx = day_idx*144 + np.arange(H)
    sd_future = daily_pattern[future_gidx % 144]
    sw_future = weekly_pattern[(future_gidx - origin_day*144) % 1008]
    # 残差 AR 仍每天轻量更新，不需要重新做 MSTL
    hist_axis = np.arange(len(deseason), dtype=float)
    fitted_trend_hist = level + slope*(hist_axis - (len(deseason)-1))
    resid_recent = deseason - fitted_trend_hist
    ar_future = choose_ar_forecast(
        resid_recent[-min(len(resid_recent), 3*144):],
        H,
        cfg.ar_candi,
        0.0,
    )
    pred = trend_future + sd_future + sw_future + ar_future
    return np.maximum(pred, 0.0)
def build_clear_envelope(day_idx: int, pv: np.ndarray, prior: np.ndarray, cfg: Config) -> np.ndarray:
    if day_idx == 0:
        base = prior.copy()
        return np.maximum(base, 0.0)
    w = min(day_idx, cfg.pv_days)
    hist = pv[day_idx-w:day_idx]
    q = np.quantile(hist, cfg.clear_quant, axis=0)
    q = moving_average_circular(q, cfg.pv_smooth)
    # 样本短时与附件1先验混合
    prior_w = max(0.0, 1.0 - w/21.0)
    base = (1-prior_w)*q + prior_w*prior
    # 夜间稳定置0：历史90%以上为0即置0
    zero_ratio = np.mean(hist <= 1e-8, axis=0)
    base[zero_ratio > 0.90] = 0.0
    return np.maximum(base, 0.0)
def pv_forecast(day_idx: int, pv: np.ndarray, prior: np.ndarray, cfg: Config, logger=None) -> Tuple[np.ndarray, Dict]:
    H = cfg.fore_slots
    if day_idx == 0:
        return np.tile(prior, 2), {"kappa_d": 1.0, "kappa_next": 1.0, "envelope": prior.copy()}
    B_d = build_clear_envelope(day_idx, pv, prior, cfg)
    # 下一日包络：年周期变化慢，先用当前包络；用近期日能量变化率做轻微尺度外推
    recent_energy = np.sum(pv[max(0,day_idx-14):day_idx], axis=1)*cfg.dt
    scale_next = 1.0
    if len(recent_energy) >= 7:
        x = np.arange(len(recent_energy))
        slope = np.polyfit(x, recent_energy, 1)[0]
        denom = max(np.mean(recent_energy), 1.0)
        scale_next = float(np.clip(1.0 + slope/denom, 0.95, 1.05))
    B_next = B_d * scale_next
    # 用当前滚动包络近似反演近期日透射率；只使用过去日期
    w = min(day_idx, cfg.pv_days)
    hist = pv[day_idx-w:day_idx]
    denom = max(np.sum(B_d)*cfg.dt, 1e-6)
    kappas = np.clip(np.sum(hist, axis=1)*cfg.dt/denom, 0.0, 1.0)
    kpred = choose_ar_forecast(kappas, 2, cfg.k_candi, fallback=float(np.mean(kappas)))
    kpred = np.clip(kpred, 0.0, 1.0)
    base_d = B_d*kpred[0]
    base_next = B_next*kpred[1]
    # 残差构造：每个历史日用其日能量对应kappa相对于当前包络重构
    residual_days = hist - kappas[:, None]*B_d[None, :]
    resid_series = residual_days.reshape(-1)
    rpred = choose_ar_forecast(resid_series[-min(len(resid_series), 21*144):], H, cfg.pv_candi, 0.0)
    pred = np.concatenate([base_d, base_next]) + rpred
    pred = np.maximum(pred, 0.0)
    # 夜间稳定无发电时段置零
    zero_mask = B_d <= 1e-8
    for h in range(H):
        if zero_mask[h % 144]:
            pred[h] = 0.0
    return pred, {"kappa_d": float(kpred[0]), "kappa_next": float(kpred[1]), "envelope": B_d}
def compute_historical_marginal_cost_day(j: int,
                                         available_day_exclusive: int,
                                         net_errors: np.ndarray,
                                         price: np.ndarray,
                                         soc_hist: np.ndarray,
                                         planned_hist: np.ndarray,
                                         load: np.ndarray,
                                         pv: np.ndarray,
                                         cfg: Config) -> Tuple[np.ndarray, np.ndarray]:
    nT = cfg.slots_pday
    cplus = np.full(nT, np.nan, dtype=float)
    cminus = np.full(nT, np.nan, dtype=float)
    if j < 0 or j >= available_day_exclusive:
        return cplus, cminus
    # 只把当前允许访问的历史展平；后续窗口均是连续切片
    n_avail = available_day_exclusive * nT
    plan_f = planned_hist[:available_day_exclusive].reshape(-1)
    load_f = load[:available_day_exclusive].reshape(-1)
    pv_f = pv[:available_day_exclusive].reshape(-1)
    soc_f = soc_hist[:available_day_exclusive].reshape(-1)
    err_f = net_errors[:available_day_exclusive].reshape(-1)
    slot_price = np.tile(price, available_day_exclusive)
    base = j * nT
    eff_rt = cfg.eta_c * cfg.eta_d
    for t in range(nT):
        e = net_errors[j, t]
        if not np.isfinite(e):
            continue
        lam_t = price[t]
        s0 = soc_hist[j, t] if np.isfinite(soc_hist[j, t]) else cfg.e0
        g0 = base + t
        g1 = g0 + 1
        g2 = min(g0 + cfg.hori_slots, n_avail - 1)
        if e > 0:
            c_em = 5.0 * lam_t
            c_bat = np.inf
            if s0 > cfg.e_min + 1e-6 and g1 <= g2:
                sl = slice(g1, g2 + 1)
                margin = plan_f[sl] + pv_f[sl] - load_f[sl]
                sf = soc_f[sl]
                pf = slot_price[sl]
                soc_ok = (~np.isfinite(sf)) | (sf < cfg.e_max - 1e-6)
                feasible = soc_ok & ((margin > 0.0) | (pf < lam_t))
                if np.any(feasible):
                    c_bat = float(np.min(pf[feasible]) / eff_rt)
            cplus[t] = min(c_em, c_bat) if np.isfinite(c_bat) else c_em
        elif e < 0:
            recovery = 0.0
            if s0 < cfg.e_max - 1e-6 and g1 <= g2:
                sl = slice(g1, g2 + 1)
                ef = err_f[sl]
                pf = slot_price[sl]
                m = np.isfinite(ef) & (ef > 0.0)
                if np.any(m):
                    recovery = float(np.max(5.0 * pf[m] * eff_rt))
            cminus[t] = max(0.0, lam_t - recovery)
    return cplus, cminus
def update_risk_cost_cache(day_idx: int,
                           cplus_cache: np.ndarray,
                           cminus_cache: np.ndarray,
                           net_errors: np.ndarray,
                           price: np.ndarray,
                           soc_hist: np.ndarray,
                           planned_hist: np.ndarray,
                           load: np.ndarray,
                           pv: np.ndarray,
                           cfg: Config) -> None:
    if day_idx <= 0:
        return
    # 昨天：只能使用到昨天24:00，保持严格无未来信息泄露
    j = day_idx - 1
    cp, cm = compute_historical_marginal_cost_day(
        j, day_idx, net_errors, price, soc_hist, planned_hist, load, pv, cfg)
    cplus_cache[j] = cp
    cminus_cache[j] = cm
    # 前天：此时“前天时刻 + 未来24h”已全部落在已发生历史内，覆盖为最终缓存
    if day_idx >= 2:
        j = day_idx - 2
        cp, cm = compute_historical_marginal_cost_day(
            j, day_idx, net_errors, price, soc_hist, planned_hist, load, pv, cfg)
        cplus_cache[j] = cp
        cminus_cache[j] = cm
def dynamic_risk_weights(day_idx: int,
                         net_errors: np.ndarray,
                         cplus_cache: np.ndarray,
                         cminus_cache: np.ndarray,
                         cfg: Config) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    nT = cfg.slots_pday
    alpha = np.full(nT, cfg.alpha_cold, dtype=float)
    beta = np.full(nT, cfg.beta_fixed, dtype=float)
    reserve = np.zeros(nT, dtype=float)
    if day_idx <= 0:
        return alpha, beta, reserve
    j0 = max(0, day_idx - cfg.rk_hisdays)
    hist_days = np.arange(j0, day_idx)
    if hist_days.size == 0:
        return alpha, beta, reserve
    ages = day_idx - hist_days
    base_w = cfg.rk_decay ** (ages - 1)
    errs = net_errors[hist_days]          # (HIST, 144)
    cp_all = cplus_cache[hist_days]
    cm_all = cminus_cache[hist_days]
    for t in range(nT):
        et = errs[:, t]
        # 正向误差历史
        mp = (et > 0) & np.isfinite(et) & np.isfinite(cp_all[:, t])
        # 负向误差历史
        mn = (et < 0) & np.isfinite(et) & np.isfinite(cm_all[:, t])
        if np.count_nonzero(mp) >= cfg.rk_sam and np.count_nonzero(mn) >= cfg.rk_sam:
            cp = float(np.average(cp_all[mp, t], weights=base_w[mp]))
            cm = float(np.average(cm_all[mn, t], weights=base_w[mn]))
            raw = cp / max(cm, cfg.c_eps)
            alpha[t] = min(cfg.alpha_max, max(cfg.alpha_min, raw))
        tau = alpha[t] / (alpha[t] + beta[t])
        tau = min(cfg.tau_max, max(cfg.tau_min, tau))
        if np.count_nonzero(mp) >= cfg.rk_sam:
            reserve[t] = cfg.reserve_scale * weighted_quantile(et[mp], base_w[mp], tau)
    return alpha, beta, reserve
@dataclass
class LPTemplates:
    da_Aeq: sparse.csr_matrix
    da_c: np.ndarray
    da_bounds: list
    mu_terminal: float
def build_lp_templates(price: np.ndarray, cfg: Config) -> LPTemplates:
    n = cfg.slots_pday
    N = 5 * n
    ib, ic, idis, isu, ie = 0, n, 2*n, 3*n, 4*n
    rows, cols, data = [], [], []
    r = 0
    # 功率平衡：buy + dis - ch - sur = 风险校准后的预测净负荷
    for t in range(n):
        for col, val in ((ib+t,1.0),(idis+t,1.0),(ic+t,-1.0),(isu+t,-1.0)):
            rows.append(r); cols.append(col); data.append(val)
        r += 1
    # SOC递推
    for t in range(n):
        rows.append(r); cols.append(ie+t); data.append(1.0)
        if t > 0:
            rows.append(r); cols.append(ie+t-1); data.append(-1.0)
        rows.append(r); cols.append(ic+t); data.append(-cfg.eta_c*cfg.dt)
        rows.append(r); cols.append(idis+t); data.append(cfg.dt/cfg.eta_d)
        r += 1
    Aeq = sparse.csr_matrix((data,(rows,cols)), shape=(r,N))
    mu = cfg.eta_d * float(np.min(price))
    c = np.zeros(N)
    c[ib:ib+n] = price * cfg.dt
    c[ic:ic+n] = cfg.eps_cycle * cfg.dt
    c[idis:idis+n] = cfg.eps_cycle * cfg.dt
    c[isu:isu+n] = cfg.eps_surplus * cfg.dt
    c[ie+n-1] = -mu
    bounds = ([(0,None)]*n + [(0,cfg.p_max)]*n + [(0,cfg.p_max)]*n +
              [(0,None)]*n + [(cfg.e_min,cfg.e_max)]*n)
    return LPTemplates(Aeq, c, bounds, mu)
def solve_day_ahead(price: np.ndarray,
                    load_fc: np.ndarray,
                    pv_fc: np.ndarray,
                    reserve: np.ndarray,
                    soc0: float,
                    cfg: Config,
                    lp: LPTemplates) -> Dict[str,np.ndarray]:
    """每天0:00求解一次全天计划购电及计划储能LP。"""
    n = cfg.slots_pday
    net_req = load_fc[:n] - pv_fc[:n] + reserve
    b = np.empty(2*n, dtype=float)
    b[:n] = net_req
    b[n:] = 0.0
    b[n] = soc0
    res = linprog(lp.da_c, A_eq=lp.da_Aeq, b_eq=b, bounds=lp.da_bounds, method=cfg.lp_method)
    if not res.success:
        raise RuntimeError(f"日前LP失败: {res.message}")
    x = res.x
    ib, ic, idis, isu, ie = 0,n,2*n,3*n,4*n
    # 加回常数 mu*soc0，得到论文中完整目标函数值
    paper_obj = float(res.fun + lp.mu_terminal * soc0)
    return {"buy":x[ib:ib+n], "ch":x[ic:ic+n], "dis":x[idis:idis+n],
            "sur":x[isu:isu+n], "E":x[ie:ie+n], "obj":paper_obj}
def operate_one_slot(soc0: float, buy: float, load_actual: float, pv_actual: float,
                     plan_ch: float, plan_dis: float, cfg: Config) -> Dict[str,float]:
    dt = cfg.dt
    # 先把计划动作投影到当前SOC可行范围
    ch = max(0.0, min(float(plan_ch), cfg.p_max))
    dis = max(0.0, min(float(plan_dis), cfg.p_max))
    max_ch_soc = max(0.0, (cfg.e_max - soc0) / (cfg.eta_c * dt))
    max_dis_soc = max(0.0, (soc0 - cfg.e_min) * cfg.eta_d / dt)
    ch = min(ch, max_ch_soc)
    dis = min(dis, max_dis_soc)
    # 基于真实负荷/光伏检查功率平衡
    gap = float(load_actual - (buy + pv_actual + dis - ch))
    emg = 0.0
    sur = 0.0
    if gap > 1e-10:
        # 1. 先取消充电，相当于立即释放可用供电
        cut_ch = min(ch, gap)
        ch -= cut_ch
        gap -= cut_ch
        # 2. 再增加放电；紧急购电价格为5倍，故在可行范围内优先用储能
        max_dis_total = min(cfg.p_max, max(0.0, (soc0 - cfg.e_min) * cfg.eta_d / dt))
        extra_dis = min(gap, max(0.0, max_dis_total - dis))
        dis += extra_dis
        gap -= extra_dis
        # 3. 储能仍不足时才紧急购电
        emg = max(0.0, gap)
    elif gap < -1e-10:
        excess = -gap
        # 1. 先取消不必要的计划放电
        cut_dis = min(dis, excess)
        dis -= cut_dis
        excess -= cut_dis
        # 2. 再增加充电
        max_ch_total = min(cfg.p_max, max(0.0, (cfg.e_max - soc0) / (cfg.eta_c * dt)))
        extra_ch = min(excess, max(0.0, max_ch_total - ch))
        ch += extra_ch
        excess -= extra_ch
        # 3. 电池无法继续吸收的部分作为富余功率
        sur = max(0.0, excess)
    soc1 = soc0 + (cfg.eta_c*ch - dis/cfg.eta_d)*dt
    soc1 = min(cfg.e_max, max(cfg.e_min, soc1))
    return {"ch":ch, "dis":dis, "emg":emg, "sur":sur, "E":soc1}
def load_attachment3(cfg: Config, dates, logger):
    if not Path(cfg.att3).exists():
        # 兼容用户下载后的文件名
        alt = Path(cfg.base_dir) / "附件3(1).xlsx"
        if alt.exists():
            cfg.att3 = str(alt)
        else:
            raise FileNotFoundError(f"未找到附件3：{cfg.att3}")
    wb = load_workbook(cfg.att3, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, max_col=26, values_only=True))
    wb.close()
    date_to_idx = {d:i for i,d in enumerate(dates)}
    out = np.full((len(dates), 4, 24), np.nan, dtype=float)
    stage_map = {"0:00":0, "6:00":1, "12:00":2, "18:00":3}
    current_date = None
    for r in rows:
        raw_date, raw_stage = r[0], r[1]
        if raw_date not in (None, ""):
            current_date = cell_to_date(raw_date)
        if current_date is None:
            continue
        stage_txt = str(raw_stage).strip() if raw_stage is not None else ""
        if stage_txt not in stage_map or current_date not in date_to_idx:
            continue
        vals = np.array([0.0 if v is None else float(v) for v in r[2:26]], dtype=float)
        out[date_to_idx[current_date], stage_map[stage_txt], :] = np.maximum(vals, 0.0)
    if np.isnan(out).any():
        bad = int(np.isnan(out).sum())
        raise ValueError(f"附件3读取不完整，仍有 {bad} 个缺失预报值。")
    return out
def pv_stage_forecast_horizon_10min(day_idx: int, stage_hour: int, pv_hourly: np.ndarray,
                                    pv_actual: np.ndarray, cfg: Config) -> np.ndarray:
    stage_idx = {0:0, 6:1, 12:2, 18:3}[stage_hour]
    hourly = pv_hourly[day_idx, stage_idx]
    start_slot = stage_hour * 6
    if stage_hour == 0:
        # 0点附近光伏通常为0；用0作为当前时刻锚点，不泄漏未来信息。
        anchor = 0.0
    else:
        # 当前时刻可用的信息仅到发布时刻之前一个10分钟区间。
        anchor = float(pv_actual[day_idx, start_slot - 1])
    xp = np.arange(0, 25, dtype=float)
    fp = np.concatenate([[anchor], hourly])
    target_rel = (np.arange(cfg.roll_slots) + 1) / 6.0
    pred = np.interp(target_rel, xp, fp)
    return np.maximum(pred, 0.0)
def load_stage_forecast_horizon(base_fc_full: np.ndarray, actual_today: np.ndarray,
                                stage_hour: int, cfg: Config) -> np.ndarray:
    start = stage_hour * 6
    H = cfg.roll_slots
    seg = np.asarray(base_fc_full[start:start+H], dtype=float).copy()
    if len(seg) != H:
        raise ValueError("负荷48小时预测长度不足，无法构造24小时MPC窗口。")
    if cfg.use_bias and stage_hour > 0:
        m = min(start, 18)
        # 注意：偏差估计只使用发布时刻以前已经实现的负荷。
        base_day = np.asarray(base_fc_full[:cfg.slots_pday], dtype=float)
        resid = actual_today[start-m:start] - base_day[start-m:start]
        w = np.exp(np.linspace(-1.5, 0.0, m))
        w /= w.sum()
        bias = float(np.sum(w * resid))
        decay = np.exp(-np.arange(H) / 36.0)
        seg += bias * decay
    return np.maximum(seg, 0.0)
def actual_horizon_from_arrays(day_idx: int, stage_hour: int, arr: np.ndarray,
                               cfg: Config) -> Optional[np.ndarray]:
    start = stage_hour * 6
    H = cfg.roll_slots
    flat = arr.reshape(-1)
    g0 = day_idx * cfg.slots_pday + start
    if g0 + H > len(flat):
        return None
    return np.asarray(flat[g0:g0+H], dtype=float)
def stage_reserve_from_history(day_idx: int, stage_idx: int,
                               stage_net_error_hist: np.ndarray,
                               fallback_reserve_day: np.ndarray,
                               stage_hour: int, cfg: Config) -> np.ndarray:
    H = cfg.roll_slots
    # 为确保任意18:00发布的24h预报已经完全实现，统一只使用 d-2 及更早记录。
    end = max(0, day_idx - 1)
    start = max(0, end - cfg.stage_rhisdays)
    hist = stage_net_error_hist[start:end, stage_idx, :]
    # 问题二备用仅有当天144点。对跨日窗口按时钟位置循环展开作为冷启动基线。
    fb = np.asarray(fallback_reserve_day, dtype=float)
    clock = (stage_hour * 6 + np.arange(H)) % cfg.slots_pday
    fallback = fb[clock]
    if hist.shape[0] < cfg.stage_mindays:
        return fallback.copy()
    out = np.zeros(H, dtype=float)
    age = np.arange(hist.shape[0]-1, -1, -1)
    day_w = cfg.stage_decay ** age
    for k in range(H):
        vals = hist[:, k]
        mask = np.isfinite(vals)
        if int(np.sum(mask)) < cfg.stage_mindays:
            out[k] = fallback[k]
            continue
        q = weighted_quantile(vals[mask], day_w[mask], cfg.stage_quant)
        out[k] = cfg.stage_scale * max(0.0, float(q))
    # 防止极端小样本产生不合理备用；上限取储能/购电系统量级内的保守值。
    return np.clip(out, 0.0, 5000.0)
def archive_stage_forecast_errors(forecast_day: int,
                                  load_fc_full_store: np.ndarray,
                                  pv_hourly: np.ndarray,
                                  load: np.ndarray, pv: np.ndarray,
                                  stage_net_error_hist: np.ndarray,
                                  cfg: Config) -> None:
    """在某日24小时预报窗口完全实现后，归档0/6/12/18四个时点的净负荷误差。"""
    if forecast_day < 0 or forecast_day >= len(load_fc_full_store):
        return
    base_full = load_fc_full_store[forecast_day]
    if not np.all(np.isfinite(base_full)):
        return
    for s, h in enumerate((0, 6, 12, 18)):
        lf = load_stage_forecast_horizon(base_full, load[forecast_day], h, cfg)
        pf = pv_stage_forecast_horizon_10min(forecast_day, h, pv_hourly, pv, cfg)
        la = actual_horizon_from_arrays(forecast_day, h, load, cfg)
        pa = actual_horizon_from_arrays(forecast_day, h, pv, cfg)
        if la is None or pa is None:
            continue
        stage_net_error_hist[forecast_day, s, :] = (la - pa) - (lf - pf)
def solve_rolling_mpc_24h(price: np.ndarray,
                          load_fc_horizon: np.ndarray,
                          pv_fc_horizon: np.ndarray,
                          reserve_horizon: np.ndarray,
                          soc0: float,
                          prev_buy_day: np.ndarray,
                          start_slot: int,
                          cfg: Config) -> Dict[str, np.ndarray]:
    H = cfg.roll_slots
    if not (len(load_fc_horizon) == len(pv_fc_horizon) == len(reserve_horizon) == H):
        raise ValueError("24小时MPC输入必须全部为144个10分钟时段。")
    current_remaining = cfg.slots_pday - start_slot
    # 变量：[buy, ch, dis, sur, E, up, down]，每组H个变量。
    N = 7 * H
    ib, ic, idis, isu, ie, iup, idn = 0, H, 2*H, 3*H, 4*H, 5*H, 6*H
    rows, cols, data, b = [], [], [], []
    r = 0
    net = np.asarray(load_fc_horizon) - np.asarray(pv_fc_horizon) + np.asarray(reserve_horizon)
    # 功率平衡
    for j in range(H):
        for col, val in ((ib+j,1.0),(idis+j,1.0),(ic+j,-1.0),(isu+j,-1.0)):
            rows.append(r); cols.append(col); data.append(val)
        b.append(float(net[j])); r += 1
    # SOC递推
    for j in range(H):
        rows.append(r); cols.append(ie+j); data.append(1.0)
        if j > 0:
            rows.append(r); cols.append(ie+j-1); data.append(-1.0)
            rhs = 0.0
        else:
            rhs = float(soc0)
        rows.append(r); cols.append(ic+j); data.append(-cfg.eta_c*cfg.dt)
        rows.append(r); cols.append(idis+j); data.append(cfg.dt/cfg.eta_d)
        b.append(rhs); r += 1
    # 仅当天剩余区间存在“与上一版已确认计划的调整关系”
    for j in range(current_remaining):
        rows.append(r); cols.append(ib+j); data.append(1.0)
        rows.append(r); cols.append(iup+j); data.append(-1.0)
        rows.append(r); cols.append(idn+j); data.append(1.0)
        b.append(float(prev_buy_day[start_slot+j])); r += 1
    Aeq = sparse.csr_matrix((data,(rows,cols)), shape=(r,N))
    clock = (start_slot + np.arange(H)) % cfg.slots_pday
    ph = np.asarray(price)[clock]
    c = np.zeros(N)
    # 当前日剩余区间只计相对既有合同的调整净费用。
    c[iup:iup+current_remaining] = 1.5 * ph[:current_remaining] * cfg.dt
    c[idn:idn+current_remaining] = -0.5 * ph[:current_remaining] * cfg.dt
    # 下一日区间尚未签订，以正常电价形成“影子购电成本”，用于跨日SOC前瞻。
    if current_remaining < H:
        c[ib+current_remaining:ib+H] = ph[current_remaining:] * cfg.dt
    c[ic:ic+H] = cfg.eps_cycle * cfg.dt
    c[idis:idis+H] = cfg.eps_cycle * cfg.dt
    c[isu:isu+H] = cfg.eps_surplus * cfg.dt
    # 24h窗口末端仍赋予剩余电量机会价值，降低滚动窗口末端效应。
    mu = cfg.eta_d * float(np.min(price))
    c[ie+H-1] = -mu
    bounds = []
    bounds += [(0,None)] * H
    bounds += [(0,cfg.p_max)] * H
    bounds += [(0,cfg.p_max)] * H
    bounds += [(0,None)] * H
    bounds += [(cfg.e_min,cfg.e_max)] * H
    # up/down只允许当前日剩余时段为正，跨日影子区间固定为0。
    bounds += [(0,None) if j < current_remaining else (0,0) for j in range(H)]
    bounds += [(0,None) if j < current_remaining else (0,0) for j in range(H)]
    res = linprog(c, A_eq=Aeq, b_eq=np.asarray(b), bounds=bounds, method=cfg.lp_method)
    if not res.success:
        raise RuntimeError(f"{start_slot//6}:00 24小时MPC失败: {res.message}")
    x = res.x
    up = x[iup:iup+H]
    dn = x[idn:idn+H]
    adj_net = float(np.sum(
        (1.5*ph[:current_remaining]*up[:current_remaining]
         - 0.5*ph[:current_remaining]*dn[:current_remaining]) * cfg.dt
    ))
    return {
        "buy": x[ib:ib+H],
        "ch": x[ic:ic+H],
        "dis": x[idis:idis+H],
        "sur": x[isu:isu+H],
        "E": x[ie:ie+H],
        "up": up,
        "down": dn,
        "adj_net_cost": adj_net,
        "up_energy": float(np.sum(up[:current_remaining]) * cfg.dt),
        "down_energy": float(np.sum(dn[:current_remaining]) * cfg.dt),
        "current_remaining": current_remaining,
    }
def fill_result3(template_path, out_path, dates, price, initial_buy, effective_buy,
                 buy_cost_after_adjust, ch, dis, soc_start, soc_end, emg, cfg, logger):
    from copy import copy
    wb=load_workbook(template_path)
    sp=wb["计划购电量"]; sa=wb["调整购电量"]; sc=wb["充放电量"]; se=wb["紧急购电量"]
    date_to_idx={d:i for i,d in enumerate(dates)}
    computed=[d for d in dates if d>=date(2025,2,1) and np.isfinite(initial_buy[date_to_idx[d]]).all()]
    def fill_buy_sheet(ws, arr, cost_arr):
        for r in range(2, ws.max_row+1):
            raw=ws.cell(r,1).value
            if raw is None: continue
            try: d=cell_to_date(raw)
            except Exception: continue
            if d not in date_to_idx: continue
            i=date_to_idx[d]
            if not np.isfinite(arr[i]).all(): continue
            vals=arr[i]*cfg.dt
            for c,v in enumerate(vals,start=2): ws.cell(r,c).value=float(v)
            ws.cell(r,146).value=float(np.sum(vals))
            ws.cell(r,147).value=float(cost_arr[i])
    initial_cost=np.nansum(initial_buy*price[None,:],axis=1)*cfg.dt
    fill_buy_sheet(sp,initial_buy,initial_cost)
    fill_buy_sheet(sa,effective_buy,buy_cost_after_adjust)
    charge_proto=[{"height":sc.row_dimensions[rr].height,
                   "cells":[sc.cell(rr,c) for c in range(1,7)]} for rr in range(2,8)]
    emg_proto=[{"height":se.row_dimensions[rr].height,
                "cells":[se.cell(rr,c) for c in range(1,4)]} for rr in range(2,min(5,se.max_row+1))]
    def cpstyle(src,dst):
        if src.has_style: dst._style=copy(src._style)
        dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border)
        dst.alignment=copy(src.alignment); dst.protection=copy(src.protection)
        dst.number_format=src.number_format
    if sc.max_row>=2: sc.delete_rows(2,sc.max_row-1)
    intervals=[(0,4,"0:00-4:00"),(4,8,"4:00-8:00"),(8,12,"8:00-12:00"),
               (12,16,"12:00-16:00"),(16,20,"16:00-20:00"),(20,24,"20:00-24:00")]
    rr=2
    for d in computed:
        i=date_to_idx[d]
        for k,(h0,h1,label) in enumerate(intervals):
            proto=charge_proto[k]
            sc.row_dimensions[rr].height=proto["height"]
            for c in range(1,7): cpstyle(proto["cells"][c-1],sc.cell(rr,c))
            sc.cell(rr,1).value=d if k==0 else None; sc.cell(rr,2).value=label
            a,b=h0*6,h1*6
            sc.cell(rr,3).value=float(np.sum(ch[i,a:b])*cfg.dt)
            sc.cell(rr,4).value=float(np.sum(dis[i,a:b])*cfg.dt)
            if k==0: sc.cell(rr,5).value="0:00"; sc.cell(rr,6).value=float(soc_start[i])
            elif k==1: sc.cell(rr,5).value="24:00"; sc.cell(rr,6).value=float(soc_end[i])
            rr+=1
    if se.max_row>=2: se.delete_rows(2,se.max_row-1)
    rr=2
    for d in computed:
        i=date_to_idx[d]
        events=_emergency_events(emg[i],cfg.dt) or [(None,None)]
        for k,(period,energy) in enumerate(events):
            proto=emg_proto[min(k,len(emg_proto)-1)]
            se.row_dimensions[rr].height=proto["height"]
            for c in range(1,4):cpstyle(proto["cells"][c-1],se.cell(rr,c))
            se.cell(rr,1).value=d if k==0 else None; se.cell(rr,2).value=period; se.cell(rr,3).value=energy
            rr+=1
    for ws in (sc,se):
        for r in range(2,ws.max_row+1):
            if ws.cell(r,1).value is not None: ws.cell(r,1).number_format='yyyy/m/d'
    Path(out_path).parent.mkdir(parents=True,exist_ok=True)
    wb.save(out_path); wb.close()
PAPER_TABLE1_INTERVALS = ("10:00-10:10", "12:00-12:10", "14:00-14:10",
                          "16:00-16:10", "18:00-18:10", "20:00-20:10")
PAPER_TABLE2_INTERVALS = ("0:00-4:00", "4:00-8:00", "8:00-12:00",
                          "12:00-16:00", "16:00-20:00", "20:00-24:00")
def _interval_start_slot(label: str) -> int:
    hh, mm = label.split("-")[0].split(":")
    return (int(hh) * 60 + int(mm)) // 10 - 1
def _slot_label(k: int) -> str:
    mins = k * 10
    if mins == 1440:
        return "24:00"
    return f"{(mins // 60) % 24}:{mins % 60:02d}"
def _emergency_events(emg_row: np.ndarray, dt: float) -> list:
    """把逐时段紧急购电功率聚合成[(时间段, 购电量kWh), ...]。"""
    events = []
    t = 0
    n = len(emg_row)
    while t < n:
        if emg_row[t] > 1e-6:
            st = t
            ee = 0.0
            while t < n and emg_row[t] > 1e-6:
                ee += float(emg_row[t]) * dt
                t += 1
            events.append((f"{_slot_label(st)}-{_slot_label(t)}", ee))
        else:
            t += 1
    return events
def export_paper_tables(dates, price, initial_buy, effective_buy, ch, dis, emg,
                        soc_start, soc_end, plan_cost, adj_cost, emg_cost,
                        run_D, cfg, logger):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    out_path = Path(cfg.out_dir) / "Q3指定日期结果.xlsx"
    d2i = {d: i for i, d in enumerate(dates)}
    targets = [d for d in cfg.represent_dates if d in d2i and d2i[d] < run_D]
    if not targets:
        logger.warning("【论文表格】表3指定日期不在本次运行范围内，跳过导出。")
        return None
    dt = cfg.dt
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    title_font = Font(bold=True, size=12)
    center = Alignment(horizontal="center", vertical="center")
    num = "#,##0.00"
    def put(ws, r, c, v, font=None, align=None, bd=True, fmt=None):
        cell = ws.cell(row=r, column=c, value=v)
        if font is not None:
            cell.font = font
        if align is not None:
            cell.alignment = align
        if bd:
            cell.border = border
        if fmt is not None:
            cell.number_format = fmt
        return cell
    wb = Workbook()
    wb.remove(wb.active)
    # 列宽
    colw = {"A": 15, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14}
    # ---------------- 表1 ----------------
    ws = wb.create_sheet("表1_购电量")
    for col, w in colw.items():
        ws.column_dimensions[col].width = w
    r = 1
    for d in targets:
        i = d2i[d]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        put(ws, r, 1, f"表1 微网在指定时间段的购电量及全天的购电量和购电费（{d.strftime('%Y-%m-%d')}）",
            title_font, center, bd=False)
        r += 1
        for c, v in enumerate(["时间段", "购电量(kWh)", "时间段", "购电量(kWh)", "时间段", "购电量(kWh)"], start=1):
            put(ws, r, c, v, bold, center)
        r += 1
        vals = [float(effective_buy[i, _interval_start_slot(lb)]) * dt for lb in PAPER_TABLE1_INTERVALS]
        grid = [(0, 1, 2), (3, 4, 5)]
        for tri in grid:
            for col_off, idx in enumerate(tri):
                put(ws, r, 1 + 2 * col_off, PAPER_TABLE1_INTERVALS[idx], None, center)
                put(ws, r, 2 + 2 * col_off, vals[idx], None, center, fmt=num)
            r += 1
        day_buy = float(np.nansum(effective_buy[i]) * dt)
        day_cost = float(plan_cost[i] + adj_cost[i] + emg_cost[i])
        put(ws, r, 1, "全天购电量", bold, center)
        put(ws, r, 2, day_buy, None, center, fmt=num)
        put(ws, r, 3, "全天购电费(元)", bold, center)
        put(ws, r, 4, day_cost, None, center, fmt=num)
        put(ws, r, 5, None, None, center)
        put(ws, r, 6, None, None, center)
        r += 2
    # ---------------- 表2 ----------------
    ws = wb.create_sheet("表2_充放电量")
    for col, w in colw.items():
        ws.column_dimensions[col].width = w
    r = 1
    ivs = [(0, 4), (4, 8), (8, 12), (12, 16), (16, 20), (20, 24)]
    for d in targets:
        i = d2i[d]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        put(ws, r, 1, f"表2 储能设备在指定时间段的充放电量及0:00和24:00的储电量（{d.strftime('%Y-%m-%d')}）",
            title_font, center, bd=False)
        r += 1
        for c, v in enumerate(["时间段", "充电量(kWh)", "放电量(kWh)", "时间段", "充电量(kWh)", "放电量(kWh)"], start=1):
            put(ws, r, c, v, bold, center)
        r += 1
        ch_e = [float(np.sum(ch[i, h0 * 6:h1 * 6]) * dt) for h0, h1 in ivs]
        dis_e = [float(np.sum(dis[i, h0 * 6:h1 * 6]) * dt) for h0, h1 in ivs]
        for a, b in ((0, 1), (2, 3), (4, 5)):
            put(ws, r, 1, PAPER_TABLE2_INTERVALS[a], None, center)
            put(ws, r, 2, ch_e[a], None, center, fmt=num)
            put(ws, r, 3, dis_e[a], None, center, fmt=num)
            put(ws, r, 4, PAPER_TABLE2_INTERVALS[b], None, center)
            put(ws, r, 5, ch_e[b], None, center, fmt=num)
            put(ws, r, 6, dis_e[b], None, center, fmt=num)
            r += 1
        put(ws, r, 1, "0:00 储电量", bold, center)
        put(ws, r, 2, float(soc_start[i]), None, center, fmt=num)
        put(ws, r, 3, "24:00 储电量", bold, center)
        put(ws, r, 4, float(soc_end[i]), None, center, fmt=num)
        put(ws, r, 5, None, None, center)
        put(ws, r, 6, None, None, center)
        r += 2
    # ---------------- 表3 ----------------
    ws = wb.create_sheet("表3_紧急购电量")
    for col, w in colw.items():
        ws.column_dimensions[col].width = w
    events_per_date = [_emergency_events(emg[d2i[d]], dt) for d in targets]
    ncol = 2 * len(targets)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    put(ws, 1, 1, "表3 微网在指定日期的紧急购电量", title_font, center, bd=False)
    for j, d in enumerate(targets):
        c0 = 1 + 2 * j
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + 1)
        put(ws, 2, c0, d.strftime("%Y.%m.%d"), bold, center)
        put(ws, 2, c0 + 1, None, bold, center)
    for j in range(len(targets)):
        put(ws, 3, 1 + 2 * j, "时间段", bold, center)
        put(ws, 3, 2 + 2 * j, "购电量(kWh)", bold, center)
    nrow = max(3, max((len(e) for e in events_per_date), default=0))
    r = 4
    for k in range(nrow):
        for j, evs in enumerate(events_per_date):
            c0 = 1 + 2 * j
            if k < len(evs):
                put(ws, r, c0, evs[k][0], None, center)
                put(ws, r, c0 + 1, evs[k][1], None, center, fmt=num)
            else:
                put(ws, r, c0, None, None, center)
                put(ws, r, c0 + 1, None, None, center)
        r += 1
    # ---------------- 附表：费用与“0点计划/最终”对照 ----------------
    ws = wb.create_sheet("附_费用与计划对照")
    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 18
    heads = ["日期", "计划全天购电量(kWh)", "最终全天购电量(kWh)", "计划购电费(元)",
             "调整净费用(元)", "紧急购电费(元)", "全天总费用(元)", "紧急购电量(kWh)"]
    for c, v in enumerate(heads, start=1):
        put(ws, 1, c, v, bold, center)
    r = 2
    for d in targets:
        i = d2i[d]
        row = [d.strftime("%Y-%m-%d"),
               float(np.nansum(initial_buy[i]) * dt),
               float(np.nansum(effective_buy[i]) * dt),
               float(plan_cost[i]), float(adj_cost[i]), float(emg_cost[i]),
               float(plan_cost[i] + adj_cost[i] + emg_cost[i]),
               float(np.sum(emg[i]) * dt)]
        for c, v in enumerate(row, start=1):
            put(ws, r, c, v, None, center if c > 1 else None, fmt=(None if c == 1 else num))
        r += 1
    r += 2
    put(ws, r, 1, "指定时段：0:00计划购电量与最终购电量对照", bold)
    r += 1
    for c, v in enumerate(["日期", "时间段", "计划购电量(kWh)", "最终购电量(kWh)"], start=1):
        put(ws, r, c, v, bold, center)
    r += 1
    for d in targets:
        i = d2i[d]
        for lb in PAPER_TABLE1_INTERVALS:
            s = _interval_start_slot(lb)
            put(ws, r, 1, d.strftime("%Y-%m-%d"), None, center)
            put(ws, r, 2, lb, None, center)
            put(ws, r, 3, float(initial_buy[i, s]) * dt, None, center, fmt=num)
            put(ws, r, 4, float(effective_buy[i, s]) * dt, None, center, fmt=num)
            r += 1
    # ---------------- 说明 ----------------
    ws = wb.create_sheet("说明")
    ws.column_dimensions["A"].width = 110
    notes = [
        "问题3 论文表格（表1/表2/表3）—— 指定日期：2025.3.20、2025.6.21、2025.9.23、2025.12.21",
        "数据来源：Q3结果/result3_filled_forecast_driven.xlsx 同源的计算结果（调整后的最终策略）。",
        "表1 购电量：取“调整后的最终购电量”（0:00计划经6:00/12:00/18:00预报调整后的结果）。",
        "表1 全天购电费：= 计划购电费 + 调整净费用 + 紧急购电费（问题3总购电费用口径）。",
        "表2 充放电量：为指定4小时时间段的最终充电量与放电量；0:00/24:00 储电量为当日首末储电量。",
        "表3 紧急购电量：为最终紧急购电的时间段与电量；留空表示该日期没有发生紧急购电。",
        "附表：给出每日0:00计划与最终购电量的对照及费用拆分，便于论文中解释“调整”的作用。",
    ]
    for i, t in enumerate(notes, start=1):
        put(ws, i, 1, t, bold if i == 1 else None, bd=False)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    wb.close()
    return str(out_path)
NO_TITLE  = True   # 图内不带标题
AXIS_FS   = 15     # 坐标轴标签字号（黑色）
TICK_FS   = 12     # 刻度字号
LEGEND_FS = 17     # 图例字号
SOFT_RED, SOFT_GREEN, SOFT_BLUE = '#E9A6A1', '#8BCF8B', '#7097CA'  # figures4papers 柔和色
BAR_EDGE  = dict(edgecolor='black', linewidth=1.2)
def _style_y(ax):
    ax.grid(True,axis='y',linestyle='--',linewidth=0.7,alpha=0.3)
    ax.spines['top'].set_visible(False); ax.spines['right'].set_visible(False)
    ax.tick_params(labelsize=TICK_FS)
def _set_labels(ax,x,y):
    ax.set_xlabel(x,fontsize=AXIS_FS,color='black')
    ax.set_ylabel(y,fontsize=AXIS_FS,color='black')
def _legend_top(ax,ncol=None):
    h,l=ax.get_legend_handles_labels()
    if h:
        ax.legend(h,l,loc='lower center',frameon=False,ncol=ncol or len(h),
                  fontsize=LEGEND_FS,bbox_to_anchor=(0.5,1.02),columnspacing=1.8)
def save_q3_outputs(dates, price, initial_buy, effective_buy, ch, dis, emg, sur,
                    soc_start, soc_end, plan_cost, adj_cost, emg_cost,
                    up_energy, down_energy,
                    run_D, cfg, logger):
    out=Path(cfg.out_dir); fig=Path(cfg.fig_dir); out.mkdir(parents=True,exist_ok=True); fig.mkdir(parents=True,exist_ok=True)
    if run_D<32:return
    idx=[i for i in range(31,run_D)]
    x=[dates[i] for i in idx]
    # 01 SOC日末轨迹图已按需删除（月度聚合仍需 idx/x）
    months=range(2,13); pc=[]; ac=[]; ec=[]; tc=[]
    for m in months:
        ids=[i for i in idx if dates[i].month==m]
        p0=sum(plan_cost[i] for i in ids); aa=sum(adj_cost[i] for i in ids); ee=sum(emg_cost[i] for i in ids)
        pc.append(p0); ac.append(aa); ec.append(ee); tc.append(p0+aa+ee)
    xx=np.arange(11)
    # 02 月度总费用：不再把数量级差异很大的三个分量硬堆叠在一起。
    f,a=plt.subplots(figsize=(12,5),dpi=300)
    a.bar(xx,np.array(tc)/1e4,width=0.55,color=SOFT_BLUE,alpha=0.95,zorder=3,**BAR_EDGE)
    a.set_xticks(xx,[f'{m}月' for m in months])
    if not NO_TITLE:
        a.set_title('完整滚动方案的月度总购电费用', fontsize=13, fontweight='bold')
    _set_labels(a,'月份','总费用 / 万元')
    _style_y(a)
    f.tight_layout(); f.savefig(fig/'第三问月度总购电费用.png',dpi=300,bbox_inches='tight'); plt.close(f)
    # 03 调整净费用与紧急购电费用单独展示，避免被计划购电费淹没。
    f,a=plt.subplots(figsize=(12,5),dpi=300)
    a.bar(xx-0.19,np.array(ac)/1e4,width=.38,color=SOFT_GREEN,alpha=0.95,label='调整净费用',zorder=3,**BAR_EDGE)
    a.bar(xx+0.19,np.array(ec)/1e4,width=.38,color=SOFT_RED,alpha=0.95,label='紧急购电费',zorder=3,**BAR_EDGE)
    a.axhline(0,lw=.9,color='black')
    a.set_xticks(xx,[f'{m}月' for m in months])
    if not NO_TITLE:
        a.set_title('月度调整净费用与紧急购电费用', fontsize=13, fontweight='bold')
    _set_labels(a,'月份','费用 / 万元')
    _style_y(a); _legend_top(a,ncol=2)
    f.tight_layout(); f.savefig(fig/'第三问月度调整净费用与紧急购电费用.png',dpi=300,bbox_inches='tight'); plt.close(f)
    formal=idx
    total_plan=sum(plan_cost[i] for i in formal); total_adj=sum(adj_cost[i] for i in formal); total_emg=sum(emg_cost[i] for i in formal)
    with open(out/'Q3_结果摘要.txt','w',encoding='utf-8') as f:
        f.write('问题3正式结果摘要（1月仅作为预热期）\n'+'='*55+'\n')
        f.write(f'正式统计日期范围：{dates[31]} 至 {dates[run_D-1]}\n')
        f.write(f'累计初始计划购电费：{total_plan:.2f} 元\n')
        f.write(f'累计调整净费用：{total_adj:.2f} 元\n')
        f.write(f'累计紧急购电费：{total_emg:.2f} 元\n')
        f.write(f'累计总费用：{total_plan+total_adj+total_emg:.2f} 元\n')
        f.write(f'累计紧急购电量：{sum(np.sum(emg[i])*cfg.dt for i in formal):.2f} kWh\n')
        f.write(f'累计富余电量：{sum(np.sum(sur[i])*cfg.dt for i in formal):.2f} kWh\n')
        f.write(f'累计向上调整电量：{sum(up_energy[i] for i in formal):.2f} kWh\n')
        f.write(f'累计向下调整电量：{sum(down_energy[i] for i in formal):.2f} kWh\n')
        f.write(f'期末储电量：{soc_end[run_D-1]:.2f} kWh\n')
def simulate_update_scheme(update_hours, dates, price, load, pv,
                           load_fc_full_store, pv_hourly,
                           fallback_reserve_store, stage_net_error_hist,
                           run_D, cfg, lp):
    n = cfg.slots_pday
    update_hours = tuple(sorted(update_hours))
    plan_cost=np.zeros(run_D); adj_cost=np.zeros(run_D); emg_cost=np.zeros(run_D)
    emg_energy=np.zeros(run_D); surplus_energy=np.zeros(run_D)
    up_energy=np.zeros(run_D); down_energy=np.zeros(run_D)
    soc_start=np.zeros(run_D); soc_end=np.zeros(run_D)
    emg_days=np.zeros(run_D,dtype=int)
    daily_emg=np.zeros(run_D)
    soc=cfg.e0
    stage_idx_map={0:0,6:1,12:2,18:3}
    for d in range(run_D):
        soc_start[d]=soc
        lf_full=load_fc_full_store[d]
        lf0=lf_full[:n]
        pf0=pv_stage_forecast_horizon_10min(d,0,pv_hourly,pv,cfg)[:n]
        fallback_reserve=fallback_reserve_store[d]
        # 0:00计划仍按当天合同生成；0点风险备用沿用问题二动态模块，
        # 保证与第二问的基准逻辑连续。
        da=solve_day_ahead(price,lf0,pf0,fallback_reserve,soc,cfg,lp)
        initial_buy=da['buy'].copy()
        current_buy=da['buy'].copy()
        current_ch=da['ch'].copy()
        current_dis=da['dis'].copy()
        plan_cost[d]=float(np.sum(initial_buy*price)*cfg.dt)
        stages=(0,)+update_hours
        for si,h in enumerate(stages):
            start=h*6
            if h>0:
                lf_h=load_stage_forecast_horizon(lf_full,load[d],h,cfg)
                pf_h=pv_stage_forecast_horizon_10min(d,h,pv_hourly,pv,cfg)
                reserve_h=stage_reserve_from_history(
                    d,stage_idx_map[h],stage_net_error_hist,
                    fallback_reserve,h,cfg
                )
                m=n-start
                adj=solve_rolling_mpc_24h(
                    price,lf_h,pf_h,reserve_h,soc,current_buy,start,cfg
                )
                current_buy[start:]=adj['buy'][:m]
                current_ch[start:]=adj['ch'][:m]
                current_dis[start:]=adj['dis'][:m]
                adj_cost[d]+=adj['adj_net_cost']
                up_energy[d]+=adj['up_energy']
                down_energy[d]+=adj['down_energy']
            end=(stages[si+1]*6) if si+1<len(stages) else n
            for t in range(start,end):
                sol=operate_one_slot(
                    soc,current_buy[t],load[d,t],pv[d,t],
                    current_ch[t],current_dis[t],cfg
                )
                soc=sol['E']
                emg_energy[d]+=sol['emg']*cfg.dt
                surplus_energy[d]+=sol['sur']*cfg.dt
                emg_cost[d]+=sol['emg']*5*price[t]*cfg.dt
        soc_end[d]=soc
        daily_emg[d]=emg_energy[d]
        emg_days[d]=1 if emg_energy[d]>1e-6 else 0
    return {
        'plan_cost':plan_cost,'adj_cost':adj_cost,'emg_cost':emg_cost,
        'emg_energy':emg_energy,'surplus_energy':surplus_energy,
        'up_energy':up_energy,'down_energy':down_energy,
        'soc_start':soc_start,'soc_end':soc_end,'emg_days':emg_days,
        'daily_emg':daily_emg,
    }
def run_update_time_comparison(dates, price, load, pv,
                               load_fc_full_store, pv_hourly,
                               fallback_reserve_store, stage_net_error_hist,
                               run_D, cfg, lp, logger):
    """比较仅0点、单独增加6/12/18点以及完整滚动方案。重点突出0点与12点。"""
    if run_D<32:
        return None
    out=Path(cfg.out_dir); fig=Path(cfg.fig_dir)
    out.mkdir(parents=True,exist_ok=True); fig.mkdir(parents=True,exist_ok=True)
    formal=slice(31,run_D)
    results=[]; raw={}
    for name,hours in zip(cfg.compare_names,cfg.compare_schemes):
        r=simulate_update_scheme(
            hours,dates,price,load,pv,load_fc_full_store,pv_hourly,
            fallback_reserve_store,stage_net_error_hist,run_D,cfg,lp
        )
        raw[name]=r
        plan=float(np.sum(r['plan_cost'][formal]))
        adj=float(np.sum(r['adj_cost'][formal]))
        emgc=float(np.sum(r['emg_cost'][formal]))
        total=plan+adj+emgc
        daily=np.asarray(r['emg_energy'][31:run_D],dtype=float)
        results.append({
            '方案':name,
            '更新时点':','.join(map(str,hours)) if hours else '无日内更新',
            '初始计划购电费':plan, '调整净费用':adj,
            '紧急购电费':emgc, '总费用':total,
            '紧急购电量kWh':float(np.sum(daily)),
            '富余电量kWh':float(np.sum(r['surplus_energy'][formal])),
            '向上调整电量kWh':float(np.sum(r['up_energy'][formal])),
            '向下调整电量kWh':float(np.sum(r['down_energy'][formal])),
            '紧急购电天数':int(np.sum(daily>1e-6)),
            '单日最大紧急购电量kWh':float(np.max(daily)) if len(daily) else 0.0,
            '日紧急购电量P95kWh':float(np.quantile(daily,0.95)) if len(daily) else 0.0,
            '期末SOC':float(r['soc_end'][run_D-1]),
        })
    names=[x['方案'] for x in results]
    x=np.arange(len(names))
    # 04：相对0点方案总费用变化，比从0开始画总费用更有辨识度。
    base_total=results[0]['总费用']
    vals=np.array([r['总费用']-base_total for r in results])/1e4
    f,a=plt.subplots(figsize=(10.8,5.4),dpi=300)
    a.bar(x,vals,width=0.40,color=SOFT_GREEN,alpha=0.95,zorder=3,**BAR_EDGE); a.axhline(0,lw=.9,color='black')
    a.set_xticks(x,names); a.tick_params(labelsize=TICK_FS)
    if not NO_TITLE:
        a.set_title('不同预报更新时点的经济性增益', fontsize=13, fontweight='bold')
    a.set_xlabel('预报更新时点方案', fontsize=AXIS_FS, color='black')
    a.set_ylabel('相对仅0:00方案的总费用变化 / 万元', fontsize=AXIS_FS, color='black')
    for i,v in enumerate(vals):
        a.text(i,v+(0.15 if v>=0 else -0.15),f'{v:+.2f}',ha='center',
               va='bottom' if v>=0 else 'top',fontsize=11)
    _style_y(a)
    f.tight_layout(); f.savefig(fig/'第三问不同预报更新时点的经济性增益.png',dpi=300,bbox_inches='tight'); plt.close(f)
    # 05-08 预报时点方案对比图已按需删除（数据仍随 CSV 输出）
    # 09 月度调整净费用图已按需删除；月度聚合保留供图10使用
    full=raw[cfg.compare_names[-1]]
    months=range(2,13); xx=np.arange(11)
    adjm=[]; upm=[]; dnm=[]
    for m in months:
        ids=[i for i in range(31,run_D) if dates[i].month==m]
        adjm.append(float(np.sum(full['adj_cost'][ids])))
        upm.append(float(np.sum(full['up_energy'][ids])))
        dnm.append(float(np.sum(full['down_energy'][ids])))
    # 10：完整方案月度向上/向下调整量
    f,a=plt.subplots(figsize=(12,5),dpi=300)
    a.bar(xx-.19,np.array(upm)/1000,width=.38,color=SOFT_BLUE,alpha=0.95,label='向上调整',zorder=3,**BAR_EDGE)
    a.bar(xx+.19,np.array(dnm)/1000,width=.38,color=SOFT_RED,alpha=0.95,label='向下调整',zorder=3,**BAR_EDGE)
    a.set_xticks(xx,[f'{m}月' for m in months])
    if not NO_TITLE:
        a.set_title('完整滚动方案的月度向上与向下调整电量', fontsize=13, fontweight='bold')
    _set_labels(a,'月份','调整电量 / MWh')
    _style_y(a); _legend_top(a,ncol=2)
    f.tight_layout(); f.savefig(fig/'第三问向上与向下调整电量.png',dpi=300,bbox_inches='tight'); plt.close(f)
    return results
def save_representative_rolling_figures(dates, price, load, pv,
                                        load_fc_full_store, pv_hourly,
                                        fallback_reserve_store, stage_net_error_hist,
                                        soc_start, run_D, cfg, lp, logger):
    """指定日期可视化：只展示有解释力的第三问新增信息。"""
    fig=Path(cfg.fig_dir); fig.mkdir(parents=True,exist_ok=True)
    date_to_idx={d:i for i,d in enumerate(dates)}
    n=cfg.slots_pday
    hours_x=(np.arange(n)+0.5)/6.0
    stage_idx_map={0:0,6:1,12:2,18:3}
    for target in cfg.represent_dates:
        if target not in date_to_idx:
            continue
        d=date_to_idx[target]
        if d>=run_D or not np.isfinite(soc_start[d]):
            continue
        lf_full=load_fc_full_store[d]
        lf0=lf_full[:n]
        pf0=pv_stage_forecast_horizon_10min(d,0,pv_hourly,pv,cfg)[:n]
        fallback=fallback_reserve_store[d]
        soc=float(soc_start[d])
        da=solve_day_ahead(price,lf0,pf0,fallback,soc,cfg,lp)
        initial_buy=da['buy'].copy(); current_buy=da['buy'].copy()
        current_ch=da['ch'].copy(); current_dis=da['dis'].copy()
        executed=np.zeros(n)
        adjustments={6:np.zeros(n),12:np.zeros(n),18:np.zeros(n)}
        stages=(0,6,12,18)
        for si,h in enumerate(stages):
            start=h*6
            if h>0:
                before=current_buy.copy()
                lf_h=load_stage_forecast_horizon(lf_full,load[d],h,cfg)
                pf_h=pv_stage_forecast_horizon_10min(d,h,pv_hourly,pv,cfg)
                reserve_h=stage_reserve_from_history(
                    d,stage_idx_map[h],stage_net_error_hist,
                    fallback,h,cfg
                )
                adj=solve_rolling_mpc_24h(price,lf_h,pf_h,reserve_h,soc,current_buy,start,cfg)
                m=n-start
                current_buy[start:]=adj['buy'][:m]; current_ch[start:]=adj['ch'][:m]
                current_dis[start:]=adj['dis'][:m]
                adjustments[h][start:]=current_buy[start:]-before[start:]
            end=(stages[si+1]*6) if si+1<len(stages) else n
            executed[start:end]=current_buy[start:end]
            for t in range(start,end):
                sol=operate_one_slot(
                    soc,current_buy[t],load[d,t],pv[d,t],
                    current_ch[t],current_dis[t],cfg
                )
                soc=sol['E']
        tag=f'{target.month}-{target.day}'
        # 11A：0点计划与最终执行，避免5条高度重合曲线。
        f,a=plt.subplots(figsize=(12,5.4),dpi=300)
        a.plot(hours_x,initial_buy,lw=2.0,color=C_BUY,label='0:00初始计划')
        a.plot(hours_x,executed,lw=1.9,color=C_DIS,label='最终执行购电策略')
        a.fill_between(hours_x,initial_buy,executed,alpha=.14,color=C_BUY,label='滚动调整区间')
        for h in (6,12,18): a.axvline(h,ls=':',lw=.9,color='#888888',zorder=1)
        a.set_xlim(0,24); a.set_xticks(np.arange(0,25,2))
        if not NO_TITLE:
            a.set_title(f'{target} 初始计划与最终执行购电策略', fontsize=13, fontweight='bold')
        _set_labels(a,'时刻 / h','购电功率 / kW')
        _style_y(a)
        hd,lb=a.get_legend_handles_labels()
        a.legend([hd[1],hd[2],hd[0]],[lb[1],lb[2],lb[0]],loc='lower center',
                 frameon=False,ncol=3,fontsize=LEGEND_FS,bbox_to_anchor=(0.5,1.02),
                 columnspacing=1.8)
        f.tight_layout(); f.savefig(fig/f'第三问{tag}初始计划与最终执行购电策略.png',dpi=300,bbox_inches='tight'); plt.close(f)
        # 11B/12 图已按需删除
def main():
    cfg=CFG
    out=Path(cfg.out_dir); out.mkdir(parents=True,exist_ok=True)
    logger=setup_logger(out)
    # 兼容本次上传文件名：默认文件缺失时回退到带(3)后缀的副本
    for attr,alt in (('att1','附件1(3).xlsx'),('att2','附件2(3).xlsx'),('template','result3.xlsx')):
        cur=getattr(cfg,attr)
        if (not cur or not Path(cur).exists()) and (Path(cfg.base_dir)/alt).exists():
            setattr(cfg,attr,str(Path(cfg.base_dir)/alt))
    dates,price,load_prior,pv_prior,load,pv=load_inputs(cfg,logger)
    pv_hourly=load_attachment3(cfg,dates,logger)
    D=len(dates); n=cfg.slots_pday
    run_D=min(D,cfg.max_days) if cfg.max_days>0 else D
    initial_buy, effective_buy, soc_hist = (np.full((D,n),np.nan) for _ in range(3))
    ch, dis, emg, sur = (np.zeros((D,n)) for _ in range(4))
    soc_start, soc_end = np.full(D,np.nan), np.full(D,np.nan)
    net_errors, cplus_cache, cminus_cache = (np.full((D,n),np.nan) for _ in range(3))
    fallback_reserve_store = np.zeros((D,n))
    # 0点负荷预测保留未来48小时，用于6/12/18点跨日24小时MPC。
    load_fc_full_store = np.full((D,cfg.fore_slots),np.nan)
    pv_fc0_store = np.full((D,n),np.nan)
    # [日期, 发布时点0/6/12/18, 未来144个10分钟时段]。
    stage_net_error_hist = np.full((D,4,cfg.roll_slots),np.nan)
    plan_cost=np.zeros(D); adj_cost=np.zeros(D); emg_cost=np.zeros(D)
    buy_cost_after_adjust=np.zeros(D)
    up_energy=np.zeros(D); down_energy=np.zeros(D)
    lp=build_lp_templates(price,cfg)
    soc=cfg.e0
    stage_idx_map={0:0,6:1,12:2,18:3}
    for d in range(run_D):
        soc_start[d]=soc
        # 到第d天0:00时，d-2日各发布时点的24小时预报已经完全实现，可安全归档。
        if d>=2:
            archive_stage_forecast_errors(
                d-2,load_fc_full_store,pv_hourly,load,pv,
                stage_net_error_hist,cfg
            )
        # 0:00负荷预测：保留完整48小时。
        lf_full=load_forecast(d,load,load_prior,cfg,logger)
        if len(lf_full)<cfg.fore_slots:
            raise ValueError('负荷预测未返回完整48小时序列。')
        lf_full=np.asarray(lf_full[:cfg.fore_slots],dtype=float)
        load_fc_full_store[d]=lf_full
        lf0=lf_full[:n].copy()
        # 附件3 0:00未来24小时光伏预报。
        pf0_full=pv_stage_forecast_horizon_10min(d,0,pv_hourly,pv,cfg)
        pf0=pf0_full[:n].copy()
        pv_fc0_store[d]=pf0
        # 0:00计划沿用问题二动态风险模块，保证问题二→问题三模型连续。
        update_risk_cost_cache(
            d,cplus_cache,cminus_cache,net_errors,
            price,soc_hist,initial_buy,load,pv,cfg
        )
        fallback_reserve=dynamic_risk_weights(
            d,net_errors,cplus_cache,cminus_cache,cfg
        )[2]
        fallback_reserve_store[d]=fallback_reserve
        da=solve_day_ahead(price,lf0,pf0,fallback_reserve,soc,cfg,lp)
        initial_buy[d]=da['buy']; current_buy=da['buy'].copy()
        current_ch=da['ch'].copy(); current_dis=da['dis'].copy()
        effective_buy[d]=current_buy
        plan_cost[d]=float(np.sum(initial_buy[d]*price)*cfg.dt)
        # 0-6、6-12、12-18、18-24四段执行。
        stage_hours=(0,6,12,18)
        for si,h in enumerate(stage_hours):
            start=h*6
            if h>0:
                # 未来24小时预测窗口，允许跨越午夜。
                lf_h=load_stage_forecast_horizon(lf_full,load[d],h,cfg)
                pf_h=pv_stage_forecast_horizon_10min(d,h,pv_hourly,pv,cfg)
                reserve_h=stage_reserve_from_history(
                    d,stage_idx_map[h],stage_net_error_hist,
                    fallback_reserve,h,cfg
                )
                adj=solve_rolling_mpc_24h(
                    price,lf_h,pf_h,reserve_h,
                    soc,current_buy,start,cfg
                )
                # result3只允许修改当天剩余计划；跨日部分仅用于MPC前瞻，不落入当天合同。
                m=n-start
                current_buy[start:]=adj['buy'][:m]; current_ch[start:]=adj['ch'][:m]
                current_dis[start:]=adj['dis'][:m]
                adj_cost[d]+=adj['adj_net_cost']
                up_energy[d]+=adj['up_energy']
                down_energy[d]+=adj['down_energy']
            end=(stage_hours[si+1]*6) if si+1<len(stage_hours) else n
            effective_buy[d,start:end]=current_buy[start:end]
            for t in range(start,end):
                soc_hist[d,t]=soc
                sol=operate_one_slot(
                    soc,current_buy[t],load[d,t],pv[d,t],
                    current_ch[t],current_dis[t],cfg
                )
                ch[d,t]=sol['ch']; dis[d,t]=sol['dis']; emg[d,t]=sol['emg']
                sur[d,t]=sol['sur']; soc=sol['E']
        soc_end[d]=soc
        emg_cost[d]=float(np.sum(emg[d]*5*price)*cfg.dt)
        buy_cost_after_adjust[d]=plan_cost[d]+adj_cost[d]
        # 0点预测误差仅用于问题二继承的风险模块与预测评价。
        net_errors[d]=(load[d]-pv[d])-(lf0-pf0)
    # 保存checkpoint，供Q3_2.py等下游脚本复用基线轨迹与阶段风险备用。
    np.savez(
        Path(cfg.out_dir)/'q3_checkpoint_forecast_driven.npz',
        soc_start=soc_start[:run_D], soc_end=soc_end[:run_D], last_day=run_D-1,
        ch=ch[:run_D], dis=dis[:run_D], emg=emg[:run_D], sur=sur[:run_D],
        initial_buy=initial_buy[:run_D], effective_buy=effective_buy[:run_D],
        adj_cost=adj_cost[:run_D], up_energy=up_energy[:run_D],
        down_energy=down_energy[:run_D],
        fallback_reserve=fallback_reserve_store[:run_D],
        stage_net_error_hist=stage_net_error_hist[:run_D]
    )
    if run_D>=32:
        fill_result3(
            cfg.template,cfg.out_xlsx,dates,price,
            initial_buy,effective_buy,buy_cost_after_adjust,
            ch,dis,soc_start,soc_end,emg,cfg,logger
        )
    save_q3_outputs(
        dates,price,initial_buy,effective_buy,ch,dis,emg,sur,
        soc_start,soc_end,plan_cost,adj_cost,emg_cost,
        up_energy,down_energy,
        run_D,cfg,logger
    )
    # 论文表格：按表1/表2/表3格式给出表3指定日期的结果，导出到Excel。
    if run_D>=32:
        export_paper_tables(
            dates,price,initial_buy,effective_buy,ch,dis,emg,
            soc_start,soc_end,plan_cost,adj_cost,emg_cost,
            run_D,cfg,logger
        )
    if run_D>=32:
        save_representative_rolling_figures(
            dates,price,load,pv,
            load_fc_full_store,pv_hourly,
            fallback_reserve_store,stage_net_error_hist,
            soc_start,run_D,cfg,lp,logger
        )
    if cfg.run_comparison and run_D>=32:
        run_update_time_comparison(
            dates,price,load,pv,
            load_fc_full_store,pv_hourly,
            fallback_reserve_store,stage_net_error_hist,
            run_D,cfg,lp,logger
        )
    print('问题三线性规划版运行结束。')
    if run_D>=32:
        formal=slice(31,run_D)
        tp=float(np.sum(plan_cost[formal]))
        ta=float(np.sum(adj_cost[formal]))
        te=float(np.sum(emg_cost[formal]))
        print(f'正式期（{dates[31]} 至 {dates[run_D-1]}）累计：初始计划购电费 {tp:.2f} 元，'
              f'调整净费用 {ta:.2f} 元，紧急购电费 {te:.2f} 元，总费用 {tp+ta+te:.2f} 元')
        print(f'result3：{cfg.out_xlsx}')
if __name__ == "__main__":
    main()