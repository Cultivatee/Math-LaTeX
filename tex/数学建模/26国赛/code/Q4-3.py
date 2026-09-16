from __future__ import annotations; import os, sys, csv, math, time, logging
from dataclasses import dataclass; from pathlib import Path
from datetime import datetime, timedelta, date
from typing import List, Tuple, Dict, Optional; import numpy as np
from scipy.optimize import linprog; from scipy import sparse
try:
    from tqdm import tqdm
except Exception:
    def tqdm(it, **kwargs): return it
from statsmodels.tsa.ar_model import AutoReg
try:
    from statsmodels.tsa.seasonal import MSTL; HAS_MSTL=True
except Exception:
    HAS_MSTL=False
try:
    from openpyxl import load_workbook
except ImportError as e:
    raise ImportError("缺少 openpyxl。请执行：pip install openpyxl") from e
class Config:
    base_dir=str(Path(__file__).resolve().parent)
    att1=str(Path(base_dir)/"附件1.xlsx"); att2=str(Path(base_dir)/"附件2.xlsx")
    att3=str(Path(base_dir)/"附件3.xlsx"); att4=str(Path(base_dir)/"附件4.xlsx")
    template=str(Path(base_dir)/"附件5"/"result4-3.xlsx")
    if not Path(template).exists():
        for _name in ("result4-3.xlsx","result4-3(1).xlsx"):
            _cand=Path(base_dir)/_name
            if _cand.exists(): template=str(_cand); break
    out_xlsx=str(Path(base_dir)/"Q4-3结果"/"result4-3_filled.xlsx")
    out_dir=str(Path(base_dir)/"Q4-3结果")
    slots_per_day=144; dt=1.0/6.0; horizon_slots=144; forecast_slots=288
    rolling_horizon_slots=144
    eta_c=0.90; eta_d=0.90; e_min=1200.0; e_max=10800.0; p_max=5000.0; e0=6000.0
    load_window_days=42; price_window_days=42
    load_min_mstl_days=21; price_min_mstl_days=21; mstl_refit_days=7
    ar_lags_candidates=(1,2,3,6,12,18,36)
    price_ar_lags_candidates=(1,2,3,6,12,18,36)
    # Q3正式模型风险参数保持不变
    risk_history_days=28; risk_decay=0.90; risk_min_samples=5
    alpha_cold=4.0; alpha_min=1.0; alpha_max=8.0; beta_fixed=1.0
    c_eps=1e-4; tau_min=0.55; tau_max=0.88; reserve_scale=0.75
    use_intraday_load_bias=False
    stage_reserve_quantile=0.80; stage_reserve_scale=0.75
    stage_reserve_min_days=7; stage_reserve_history_days=35
    stage_reserve_decay=0.94
    # 日内电价滚动修正：仅用发布时刻之前最近3h真实电价误差
    price_intraday_recent_slots=18; price_intraday_decay_hours=4.0
    price_intraday_bias_clip=0.35
    eps_cycle=1e-5; eps_surplus=1e-8; lp_method="highs"
    run_comparison=True; comparison_schemes=((),(6,),(12,),(18,),(6,12,18))
    comparison_names=( "仅0:00", "0:00+6:00", "0:00+12:00", "0:00+18:00",
        "0:00+6:00+12:00+18:00")
    focus_scheme_hours=(12,); focus_scheme_name="0:00+12:00"
    representative_dates=( date(2025, 3, 20), date(2025, 6, 21), date(2025, 9,
        23), date(2025, 12, 21))
    max_days=365; checkpoint_every_days=30; verbose_each_day=False
    reuse_q42_forecast_cache=True
    q42_cache=str(Path(base_dir)/"Q4-2结果"/"q4_2_checkpoint.npz")
CFG=Config()
def setup_logger(out_dir:Path)->logging.Logger:
    out_dir.mkdir(parents=True,exist_ok=True); logger=logging.getLogger("Q4-3")
    logger.setLevel(logging.INFO); logger.handlers.clear()
    fmt=logging.Formatter("%(message)s")
    sh=logging.StreamHandler(sys.stdout); sh.setFormatter(fmt)
    fh=logging.FileHandler(out_dir/"q4_3_run.log",mode="w",encoding="utf-8")
    fh.setFormatter(fmt)
    logger.addHandler(sh); logger.addHandler(fh); return logger
def excel_serial_to_date(x) -> datetime:
    return datetime(1899, 12, 30) + timedelta(days=float(x))
def cell_to_date(x):
    """兼容 openpyxl 返回的 date/datetime、Excel 序列号和常见日期字符串。"""
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    try:
        from datetime import date as _date
        if isinstance(x, _date):
            return x
    except Exception:
        pass
    if isinstance(x, (int, float, np.integer, np.floating)):
        return excel_serial_to_date(x).date()
    if isinstance(x, str):
        txt=x.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(txt, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"无法识别日期单元格：{x!r}")
def safe_float_array(x) -> np.ndarray:
    return np.asarray( [[float(v) if v is not None else np.nan for v in row] for
        row in x], dtype=float)
def weighted_quantile( values: np.ndarray, weights: np.ndarray,
    q: float) -> float:
    values = np.asarray(values, dtype=float)
    weights = np.asarray(weights, dtype=float)
    m = np.isfinite(values) & np.isfinite(weights) & (weights > 0)
    if m.sum() == 0:
        return 0.0
    v, w = values[m], weights[m]; idx = np.argsort(v); v, w = v[idx], w[idx]
    cw = np.cumsum(w) / np.sum(w); return float(np.interp(q, cw, v))
def moving_average_circular(x: np.ndarray, half: int) -> np.ndarray:
    if half <= 0:
        return x.copy()
    n = len(x); out = np.zeros(n)
    for i in range(n):
        lo = max(0, i-half); hi = min(n, i+half+1); out[i] = np.mean(x[lo:hi])
    return out
def choose_ar_forecast( series: np.ndarray, horizon: int, candidates: Tuple[int,
    ...], fallback: float = 0.0) -> np.ndarray:
    """用AIC在少量AutoReg候选中选阶；失败时返回fallback/近期均值。"""
    y = np.asarray(series, dtype=float); y = y[np.isfinite(y)]
    if len(y) < 8 or np.std(y) < 1e-8:
        return np.full(horizon, float(np.mean(y) if len(y) else fallback))
    best = None
    for p in candidates:
        if len(y) <= p + 5:
            continue
        try:
            model = AutoReg(y, lags=p, trend="ct", old_names=False).fit()
            if np.isfinite(model.aic) and (best is None or model.aic < best[0]):
                best = (model.aic, model)
        except Exception:
            continue
    if best is None:
        return np.full(horizon, float(np.mean(y[-min(7, len(y)):])) )
    try:
        pred = np.asarray( best[1].predict(start=len(y), end=len(y)+ horizon- 1,
            dynamic=False), dtype=float)
        if np.any(~np.isfinite(pred)):
            raise ValueError("AR forecast nonfinite")
        return pred
    except Exception:
        return np.full(horizon, float(np.mean(y[-min(7, len(y)):])) )
def load_inputs(cfg:Config,logger):
    logger.info("【数据读取】正在读取附件1、附件2、附件3、附件4……")
    for attr,names in [ ('att1', ('附件1.xlsx', '附件1(3).xlsx', '附件1(2).xlsx')),
        ('att2', ('附件2.xlsx', '附件2(3).xlsx', '附件2(1).xlsx')), ('att3',
        ('附件3.xlsx', '附件3(1).xlsx')), ('att4', ('附件4.xlsx', '附件4(2).xlsx'))]:
        if not Path(getattr(cfg,attr)).exists():
            for name in names:
                cand=Path(cfg.base_dir)/name
                if cand.exists(): setattr(cfg,attr,str(cand)); break
    for p in (cfg.att1,cfg.att2,cfg.att3,cfg.att4):
        if not Path(p).exists(): raise FileNotFoundError(f"未找到文件：{p}")
    wb1=load_workbook(cfg.att1,data_only=True,read_only=True)
    ws1=wb1["Sheet1"] if "Sheet1" in wb1.sheetnames else wb1[wb1.sheetnames[0]]
    rows1=list( ws1.iter_rows(min_row=2, max_row=145, min_col=2, max_col=4,
        values_only=True))
    wb1.close()
    price_prior=np.asarray([float(r[0]) for r in rows1],dtype=float)
    load_prior=np.asarray([float(r[1]) for r in rows1],dtype=float)
    wb2=load_workbook(cfg.att2,data_only=True,read_only=True)
    sl=wb2["小区负载"]; sp=wb2["光伏发电实际功率"]
    load_all=list( sl.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145,
        values_only=True))
    pv_all=list( sp.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145,
        values_only=True))
    wb2.close(); dates=[cell_to_date(r[0]) for r in load_all]
    load=safe_float_array([r[1:] for r in load_all])
    pv=safe_float_array([r[1:] for r in pv_all])
    wb4=load_workbook(cfg.att4,data_only=True,read_only=True)
    ws4=wb4["Sheet1"] if "Sheet1" in wb4.sheetnames else wb4[wb4.sheetnames[0]]
    pr_rows=list( ws4.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145,
        values_only=True))
    wb4.close(); price_dates=[cell_to_date(r[0]) for r in pr_rows]
    price_actual=safe_float_array([r[1:] for r in pr_rows])
    if load.shape!=( 365, 144) or pv.shape!=( 365, 144) or price_actual.shape!=(
        365, 144):
        raise ValueError(
            f"数据维度异常：load={load.shape}, pv={pv.shape}, price={price_actual.shape}")
    if dates!=price_dates: raise ValueError("附件2与附件4日期未完全对齐。")
    if np.any( ~np.isfinite(price_actual)) or np.any( price_actual< 0):
        raise ValueError( "附件4存在缺失或负电价。")
    logger.info( "【数据读取】完成：365天×144个10分钟时段；电价范围 %.4f~%.4f 元/kWh。",
        float(np.min(price_actual)), float(np.max(price_actual)))
    return dates,price_prior,load_prior,load,pv,price_actual
_LOAD_MSTL_CACHE = {}
def load_forecast( day_idx: int, load: np.ndarray, prior: np.ndarray,
    cfg: Config, logger=None) -> np.ndarray:
    H = cfg.forecast_slots
    if day_idx == 0:
        return np.tile(prior, 2)
    hist_days = min(day_idx, cfg.load_window_days)
    hist_mat = load[day_idx-hist_days:day_idx]
    # 冷启动：同一时刻近期日加权 + 附件1先验
    if hist_days < cfg.load_min_mstl_days or not HAS_MSTL:
        k = hist_mat.shape[0]; weights = np.exp(-0.10*np.arange(k-1, -1, -1))
        weights /= weights.sum()
        shape = np.sum(hist_mat * weights[:, None], axis=0)
        if hist_days >= 7:
            shape = 0.55*shape + 0.45*load[day_idx-7]
        w_prior = max(0.0, 1.0 - hist_days / cfg.load_min_mstl_days)
        shape = (1-w_prior)*shape + w_prior*prior; return np.tile(shape, 2)
    # 仅每隔若干天完整重估一次 MSTL；数据已确认无明显异常，因此关闭 robust 迭代
    need_refit = ( not _LOAD_MSTL_CACHE or day_idx -
        _LOAD_MSTL_CACHE.get("refit_day", - 10**9) >= cfg.mstl_refit_days )
    if need_refit:
        y = hist_mat.reshape(-1)
        try:
            mstl = MSTL( y, periods=(144, 1008),
                stl_kwargs={"robust": False}).fit( )
            trend = np.asarray(mstl.trend); seasonal = np.asarray(mstl.seasonal)
            if seasonal.ndim == 1:
                s_daily = seasonal; s_weekly = np.zeros_like(seasonal)
            else:
                s_daily = seasonal[:, 0]; s_weekly = seasonal[:, 1]
            resid = np.asarray(mstl.resid)
            # 保存稳定的日周期与周周期结构，供未来几天复用
            daily_tail = s_daily[-min(len(s_daily), 28*144):]
            weekly_tail = s_weekly[-min(len(s_weekly), 6*1008):]
            daily_pattern = np.array( [ np.nanmean(daily_tail[i::144]) for
                i in range(144) ])
            weekly_pattern = np.array( [ np.nanmean(weekly_tail[i::1008]) for
                i in range(1008) ])
            _LOAD_MSTL_CACHE.clear()
            _LOAD_MSTL_CACHE.update( { "refit_day": day_idx,
                "origin_day": day_idx - hist_days,
                "daily_pattern": daily_pattern,"weekly_pattern": weekly_pattern,
                "trend_last": float(trend[- 1]),
                "trend_slope": float(np.clip( np.polyfit( np.arange(min(432,
                len(trend)), dtype=float), trend[- min(432, len(trend)):],
                1 )[0], - 5.0, 5.0 )), "resid_tail": resid[- min(len(resid), 14*
                144):].copy(), })
            if logger:
                logger.info("负荷预测：第 %d 天重新拟合 MSTL（日/周周期）。", day_idx + 1)
        except Exception as e:
            if logger:
                logger.warning("MSTL失败(day=%d)，回退季节加权预测：%s", day_idx, e)
            k = hist_mat.shape[0]
            weights = np.exp(-0.10*np.arange(k-1, -1, -1))
            weights /= weights.sum()
            shape = np.sum(hist_mat * weights[:, None], axis=0)
            if hist_days >= 7:
                shape = 0.5*shape + 0.5*load[day_idx-7]
            return np.tile(shape, 2)
    # 使用最近一次 MSTL 的季节结构，并用最新 3 天数据快速校正趋势和残差
    daily_pattern = _LOAD_MSTL_CACHE["daily_pattern"]
    weekly_pattern = _LOAD_MSTL_CACHE["weekly_pattern"]
    origin_day = _LOAD_MSTL_CACHE["origin_day"]; recent_days = min(hist_days, 3)
    recent = load[day_idx-recent_days:day_idx].reshape(-1)
    global_start = (day_idx - recent_days) * 144
    gidx = global_start + np.arange(len(recent))
    sd_hist = daily_pattern[gidx % 144]
    sw_hist = weekly_pattern[(gidx - origin_day*144) % 1008]
    deseason = recent - sd_hist - sw_hist; m = min(432, len(deseason))
    x = np.arange(m, dtype=float); coef = np.polyfit(x, deseason[-m:], 1)
    slope = float(np.clip(coef[0], -5.0, 5.0)); level = float(deseason[-1])
    trend_future = level + slope*np.arange(1, H+1)
    future_gidx = day_idx*144 + np.arange(H)
    sd_future = daily_pattern[future_gidx % 144]
    sw_future = weekly_pattern[(future_gidx - origin_day*144) % 1008]
    # 残差 AR 仍每天轻量更新，不需要重新做 MSTL
    hist_axis = np.arange(len(deseason), dtype=float)
    fitted_trend_hist = level + slope*(hist_axis - (len(deseason)-1))
    resid_recent = deseason - fitted_trend_hist
    ar_future = choose_ar_forecast( resid_recent[- min(len(resid_recent), 3*
        144):], H, cfg.ar_lags_candidates, 0.0, )
    pred = trend_future + sd_future + sw_future + ar_future
    return np.maximum(pred, 0.0)
_PRICE_MSTL_CACHE = {}
def price_forecast( day_idx: int, price_actual: np.ndarray, prior: np.ndarray,
    cfg: Config, logger=None) -> np.ndarray:
    H = cfg.forecast_slots
    if day_idx == 0:
        return np.tile(prior, 2)
    hist_days = min(day_idx, cfg.price_window_days)
    hist_mat = price_actual[day_idx-hist_days:day_idx]
    # 冷启动：同一时刻近期日加权 + 附件1先验
    if hist_days < cfg.price_min_mstl_days or not HAS_MSTL:
        k = hist_mat.shape[0]; weights = np.exp(-0.10*np.arange(k-1, -1, -1))
        weights /= weights.sum()
        shape = np.sum(hist_mat * weights[:, None], axis=0)
        if hist_days >= 7:
            shape = 0.55*shape + 0.45*price_actual[day_idx-7]
        w_prior = max(0.0, 1.0 - hist_days / cfg.price_min_mstl_days)
        shape = (1-w_prior)*shape + w_prior*prior; return np.tile(shape, 2)
    # 仅每隔若干天完整重估一次 MSTL；数据已确认无明显异常，因此关闭 robust 迭代
    need_refit = ( not _PRICE_MSTL_CACHE or day_idx -
        _PRICE_MSTL_CACHE.get("refit_day", - 10**9) >= cfg.mstl_refit_days )
    if need_refit:
        y = hist_mat.reshape(-1)
        try:
            mstl = MSTL( y, periods=(144, 1008),
                stl_kwargs={"robust": False}).fit( )
            trend = np.asarray(mstl.trend); seasonal = np.asarray(mstl.seasonal)
            if seasonal.ndim == 1:
                s_daily = seasonal; s_weekly = np.zeros_like(seasonal)
            else:
                s_daily = seasonal[:, 0]; s_weekly = seasonal[:, 1]
            resid = np.asarray(mstl.resid)
            # 保存稳定的日周期与周周期结构，供未来几天复用
            daily_tail = s_daily[-min(len(s_daily), 28*144):]
            weekly_tail = s_weekly[-min(len(s_weekly), 6*1008):]
            daily_pattern = np.array( [ np.nanmean(daily_tail[i::144]) for
                i in range(144) ])
            weekly_pattern = np.array( [ np.nanmean(weekly_tail[i::1008]) for
                i in range(1008) ])
            _PRICE_MSTL_CACHE.clear()
            _PRICE_MSTL_CACHE.update( { "refit_day": day_idx,
                "origin_day": day_idx - hist_days,
                "daily_pattern": daily_pattern,"weekly_pattern": weekly_pattern,
                "trend_last": float(trend[- 1]),
                "trend_slope": float(np.clip( np.polyfit( np.arange(min(432,
                len(trend)), dtype=float), trend[- min(432, len(trend)):],
                1 )[0], - 5.0, 5.0 )), "resid_tail": resid[- min(len(resid), 14*
                144):].copy(), })
            if logger:
                logger.info("电价预测：第 %d 天重新拟合 MSTL（日/周周期）。", day_idx + 1)
        except Exception as e:
            if logger:
                logger.warning("MSTL失败(day=%d)，回退季节加权预测：%s", day_idx, e)
            k = hist_mat.shape[0]
            weights = np.exp(-0.10*np.arange(k-1, -1, -1))
            weights /= weights.sum()
            shape = np.sum(hist_mat * weights[:, None], axis=0)
            if hist_days >= 7:
                shape = 0.5*shape + 0.5*price_actual[day_idx-7]
            return np.tile(shape, 2)
    # 使用最近一次 MSTL 的季节结构，并用最新 3 天数据快速校正趋势和残差
    daily_pattern = _PRICE_MSTL_CACHE["daily_pattern"]
    weekly_pattern = _PRICE_MSTL_CACHE["weekly_pattern"]
    origin_day = _PRICE_MSTL_CACHE["origin_day"]
    recent_days = min(hist_days, 3)
    recent = price_actual[day_idx-recent_days:day_idx].reshape(-1)
    global_start = (day_idx - recent_days) * 144
    gidx = global_start + np.arange(len(recent))
    sd_hist = daily_pattern[gidx % 144]
    sw_hist = weekly_pattern[(gidx - origin_day*144) % 1008]
    deseason = recent - sd_hist - sw_hist; m = min(432, len(deseason))
    x = np.arange(m, dtype=float); coef = np.polyfit(x, deseason[-m:], 1)
    slope = float(np.clip(coef[0], -5.0, 5.0)); level = float(deseason[-1])
    trend_future = level + slope*np.arange(1, H+1)
    future_gidx = day_idx*144 + np.arange(H)
    sd_future = daily_pattern[future_gidx % 144]
    sw_future = weekly_pattern[(future_gidx - origin_day*144) % 1008]
    # 残差 AR 仍每天轻量更新，不需要重新做 MSTL
    hist_axis = np.arange(len(deseason), dtype=float)
    fitted_trend_hist = level + slope*(hist_axis - (len(deseason)-1))
    resid_recent = deseason - fitted_trend_hist
    ar_future = choose_ar_forecast( resid_recent[- min(len(resid_recent), 3*
        144):], H, cfg.price_ar_lags_candidates, 0.0, )
    pred = trend_future + sd_future + sw_future + ar_future
    return np.maximum(pred, 0.0)
def compute_historical_marginal_cost_day( j: int, available_day_exclusive: int,
    net_errors: np.ndarray, price_actual: np.ndarray, soc_hist: np.ndarray,
    planned_hist: np.ndarray, load: np.ndarray, pv: np.ndarray,
    cfg: Config) -> Tuple[ np.ndarray, np.ndarray]:
    nT = cfg.slots_per_day; cplus = np.full(nT, np.nan, dtype=float)
    cminus = np.full(nT, np.nan, dtype=float)
    if j < 0 or j >= available_day_exclusive:
        return cplus, cminus
    # 只把当前允许访问的历史展平；后续窗口均是连续切片
    n_avail = available_day_exclusive * nT
    plan_f = planned_hist[:available_day_exclusive].reshape(-1)
    load_f = load[:available_day_exclusive].reshape(-1)
    pv_f = pv[:available_day_exclusive].reshape(-1)
    soc_f = soc_hist[:available_day_exclusive].reshape(-1)
    err_f = net_errors[:available_day_exclusive].reshape(-1)
    slot_price = price_actual[:available_day_exclusive].reshape(-1)
    base = j * nT; eff_rt = cfg.eta_c * cfg.eta_d
    for t in range(nT):
        e = net_errors[j, t]
        if not np.isfinite(e):
            continue
        lam_t = price_actual[j, t]
        s0 = soc_hist[j, t] if np.isfinite(soc_hist[j, t]) else cfg.e0
        g0 = base + t; g1 = g0 + 1
        g2 = min(g0 + cfg.horizon_slots, n_avail - 1)
        if e > 0:
            c_em = 5.0 * lam_t; c_bat = np.inf
            if s0 > cfg.e_min + 1e-6 and g1 <= g2:
                sl = slice(g1, g2 + 1)
                margin = plan_f[sl] + pv_f[sl] - load_f[sl]; sf = soc_f[sl]
                pf = slot_price[sl]
                soc_ok = (~np.isfinite(sf)) | (sf < cfg.e_max - 1e-6)
                feasible = soc_ok & ((margin > 0.0) | (pf < lam_t))
                if np.any(feasible):
                    c_bat = float(np.min(pf[feasible]) / eff_rt)
            cplus[t] = min(c_em, c_bat) if np.isfinite(c_bat) else c_em
        elif e < 0:
            recovery = 0.0
            if s0 < cfg.e_max - 1e-6 and g1 <= g2:
                sl = slice(g1, g2 + 1); ef = err_f[sl]; pf = slot_price[sl]
                m = np.isfinite(ef) & (ef > 0.0)
                if np.any(m):
                    recovery = float(np.max(5.0 * pf[m] * eff_rt))
            cminus[t] = max(0.0, lam_t - recovery)
    return cplus, cminus
def update_risk_cost_cache( day_idx: int, cplus_cache: np.ndarray,
    cminus_cache: np.ndarray, net_errors: np.ndarray, price_actual: np.ndarray,
    soc_hist: np.ndarray, planned_hist: np.ndarray, load: np.ndarray,
    pv: np.ndarray, cfg: Config) -> None:
    if day_idx <= 0:
        return
    # 昨天：只能使用到昨天24:00，保持严格无未来信息泄露
    j = day_idx - 1
    cp, cm = compute_historical_marginal_cost_day( j, day_idx, net_errors,
        price_actual, soc_hist, planned_hist, load, pv, cfg)
    cplus_cache[j] = cp; cminus_cache[j] = cm
    # 前天：此时“前天时刻 + 未来24h”已全部落在已发生历史内，覆盖为最终缓存
    if day_idx >= 2:
        j = day_idx - 2
        cp, cm = compute_historical_marginal_cost_day( j, day_idx, net_errors,
            price_actual, soc_hist, planned_hist, load, pv, cfg)
        cplus_cache[j] = cp; cminus_cache[j] = cm
def dynamic_risk_weights( day_idx: int, net_errors: np.ndarray,
    cplus_cache: np.ndarray, cminus_cache: np.ndarray, cfg: Config) -> Tuple[
    np.ndarray, np.ndarray, np.ndarray]:
    nT = cfg.slots_per_day; alpha = np.full(nT, cfg.alpha_cold, dtype=float)
    beta = np.full(nT, cfg.beta_fixed, dtype=float)
    reserve = np.zeros(nT, dtype=float)
    if day_idx <= 0:
        return alpha, beta, reserve
    j0 = max(0, day_idx - cfg.risk_history_days)
    hist_days = np.arange(j0, day_idx)
    if hist_days.size == 0:
        return alpha, beta, reserve
    ages = day_idx - hist_days; base_w = cfg.risk_decay ** (ages - 1)
    errs = net_errors[hist_days]  # (HIST, 144)
    cp_all = cplus_cache[hist_days]; cm_all = cminus_cache[hist_days]
    for t in range(nT):
        et = errs[:, t]
        # 正向误差历史
        mp = (et > 0) & np.isfinite(et) & np.isfinite(cp_all[:, t])
        # 负向误差历史
        mn = (et < 0) & np.isfinite(et) & np.isfinite(cm_all[:, t])
        if np.count_nonzero( mp) >= cfg.risk_min_samples and np.count_nonzero(
            mn) >= cfg.risk_min_samples:
            cp = float(np.average(cp_all[mp, t], weights=base_w[mp]))
            cm = float(np.average(cm_all[mn, t], weights=base_w[mn]))
            raw = cp / max(cm, cfg.c_eps)
            alpha[t] = min(cfg.alpha_max, max(cfg.alpha_min, raw))
        tau = alpha[t] / (alpha[t] + beta[t])
        tau = min(cfg.tau_max, max(cfg.tau_min, tau))
        if np.count_nonzero(mp) >= cfg.risk_min_samples:
            reserve[ t] = cfg.reserve_scale * weighted_quantile( et[mp],
                base_w[mp], tau)
    return alpha, beta, reserve
def operate_one_slot( soc0: float, buy: float, load_actual: float,
    pv_actual: float, plan_ch: float, plan_dis: float, cfg: Config) -> Dict[str,
    float]:
    dt = cfg.dt
    # 先把计划动作投影到当前SOC可行范围
    ch = max(0.0, min(float(plan_ch), cfg.p_max))
    dis = max(0.0, min(float(plan_dis), cfg.p_max))
    max_ch_soc = max(0.0, (cfg.e_max - soc0) / (cfg.eta_c * dt))
    max_dis_soc = max(0.0, (soc0 - cfg.e_min) * cfg.eta_d / dt)
    ch = min(ch, max_ch_soc); dis = min(dis, max_dis_soc)
    # 基于真实负荷/光伏检查功率平衡
    gap = float(load_actual - (buy + pv_actual + dis - ch)); emg = 0.0
    sur = 0.0
    if gap > 1e-10:
        cut_ch = min(ch, gap); ch -= cut_ch; gap -= cut_ch
        max_dis_total = min( cfg.p_max, max(0.0, (soc0 - cfg.e_min) *cfg.eta_d /
            dt))
        extra_dis = min(gap, max(0.0, max_dis_total - dis)); dis += extra_dis
        gap -= extra_dis

        emg = max(0.0, gap)
    elif gap < -1e-10:
        excess = -gap
        cut_dis = min(dis, excess); dis -= cut_dis; excess -= cut_dis
        max_ch_total = min( cfg.p_max, max(0.0, (cfg.e_max - soc0) /(cfg.eta_c *
            dt)))
        extra_ch = min(excess, max(0.0, max_ch_total - ch)); ch += extra_ch
        excess -= extra_ch
        sur = max(0.0, excess)
    soc1 = soc0 + (cfg.eta_c*ch - dis/cfg.eta_d)*dt
    soc1 = min(cfg.e_max, max(cfg.e_min, soc1))
    return {"ch":ch, "dis":dis, "emg":emg, "sur":sur, "E":soc1}
def load_attachment3(cfg: Config, dates, logger):
    if not Path(cfg.att3).exists():
        alt = Path(cfg.base_dir) / "附件3(1).xlsx"
        if alt.exists():
            cfg.att3 = str(alt)
        else:
            raise FileNotFoundError(f"未找到附件3：{cfg.att3}")
    wb = load_workbook(cfg.att3, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, max_col=26, values_only=True))
    wb.close(); date_to_idx = {d:i for i,d in enumerate(dates)}
    out = np.full((len(dates), 4, 24), np.nan, dtype=float)
    stage_map = {"0:00":0, "6:00":1, "12:00":2, "18:00":3}; current_date = None
    for r in rows:
        raw_date, raw_stage = r[0], r[1]
        if raw_date not in (None, ""):
            current_date = cell_to_date(raw_date)
        if current_date is None:
            continue
        stage_txt = str(raw_stage).strip() if raw_stage is not None else ""
        if stage_txt not in stage_map or current_date not in date_to_idx:
            continue
        vals = np.array( [0.0 if v is None else float(v) for v in r[2:26]],
            dtype=float)
        out[ date_to_idx[current_date], stage_map[stage_txt], :] = np.maximum(
            vals, 0.0)
    if np.isnan(out).any():
        bad = int(np.isnan(out).sum())
        raise ValueError(f"附件3读取不完整，仍有 {bad} 个缺失预报值。")
    logger.info("【附件3】读取完成：%d天 × 4个预报时刻 × 24小时。", len(dates)); return out
def pv_stage_forecast_horizon_10min( day_idx: int, stage_hour: int,
    pv_hourly: np.ndarray, pv_actual: np.ndarray, cfg: Config) -> np.ndarray:
    stage_idx = {0:0, 6:1, 12:2, 18:3}[stage_hour]
    hourly = pv_hourly[day_idx, stage_idx]; start_slot = stage_hour * 6
    if stage_hour == 0:
        anchor = 0.0
    else:
        anchor = float(pv_actual[day_idx, start_slot - 1])
    xp = np.arange(0, 25, dtype=float); fp = np.concatenate([[anchor], hourly])
    target_rel = (np.arange(cfg.rolling_horizon_slots) + 1) / 6.0
    pred = np.interp(target_rel, xp, fp); return np.maximum(pred, 0.0)
def load_stage_forecast_horizon( base_fc_full: np.ndarray,
    actual_today: np.ndarray, stage_hour: int, cfg: Config) -> np.ndarray:
    start = stage_hour * 6; H = cfg.rolling_horizon_slots
    seg = np.asarray(base_fc_full[start:start+H], dtype=float).copy()
    if len(seg) != H:
        raise ValueError("负荷48小时预测长度不足，无法构造24小时MPC窗口。")
    if cfg.use_intraday_load_bias and stage_hour > 0:
        m = min(start, 18)
        # 注意：偏差估计只使用发布时刻以前已经实现的负荷。
        base_day = np.asarray(base_fc_full[:cfg.slots_per_day], dtype=float)
        resid = actual_today[start-m:start] - base_day[start-m:start]
        w = np.exp(np.linspace(-1.5, 0.0, m)); w /= w.sum()
        bias = float(np.sum(w * resid)); decay = np.exp(-np.arange(H) / 36.0)
        seg += bias * decay
    return np.maximum(seg, 0.0)
def actual_horizon_from_arrays( day_idx: int, stage_hour: int, arr: np.ndarray,
    cfg: Config) -> Optional[ np.ndarray]:
    start = stage_hour * 6; H = cfg.rolling_horizon_slots
    flat = arr.reshape(-1); g0 = day_idx * cfg.slots_per_day + start
    if g0 + H > len(flat):
        return None
    return np.asarray(flat[g0:g0+H], dtype=float)
def stage_reserve_from_history( day_idx: int, stage_idx: int,
    stage_net_error_hist: np.ndarray, fallback_reserve_day: np.ndarray,
    stage_hour: int, cfg: Config) -> np.ndarray:
    H = cfg.rolling_horizon_slots
    # 为确保任意18:00发布的24h预报已经完全实现，统一只使用 d-2 及更早记录。
    end = max(0, day_idx - 1)
    start = max(0, end - cfg.stage_reserve_history_days)
    hist = stage_net_error_hist[start:end, stage_idx, :]
    # 问题二备用仅有当天144点。对跨日窗口按时钟位置循环展开作为冷启动基线。
    fb = np.asarray(fallback_reserve_day, dtype=float)
    clock = (stage_hour * 6 + np.arange(H)) % cfg.slots_per_day
    fallback = fb[clock]
    if hist.shape[0] < cfg.stage_reserve_min_days:
        return fallback.copy()
    out = np.zeros(H, dtype=float); age = np.arange(hist.shape[0]-1, -1, -1)
    day_w = cfg.stage_reserve_decay ** age
    for k in range(H):
        vals = hist[:, k]; mask = np.isfinite(vals)
        if int(np.sum(mask)) < cfg.stage_reserve_min_days:
            out[k] = fallback[k]; continue
        q = weighted_quantile( vals[mask], day_w[mask],
            cfg.stage_reserve_quantile)
        out[k] = cfg.stage_reserve_scale * max(0.0, float(q))
    # 防止极端小样本产生不合理备用；上限取储能/购电系统量级内的保守值。
    return np.clip(out, 0.0, 5000.0)
def archive_stage_forecast_errors( forecast_day: int,
    load_fc_full_store: np.ndarray, pv_hourly: np.ndarray, load: np.ndarray,
    pv: np.ndarray, stage_net_error_hist: np.ndarray, cfg: Config) -> None:
    """在某日24小时预报窗口完全实现后，归档0/6/12/18四个时点的净负荷误差。"""
    if forecast_day < 0 or forecast_day >= len(load_fc_full_store):
        return
    base_full = load_fc_full_store[forecast_day]
    if not np.all(np.isfinite(base_full)):
        return
    for s, h in enumerate((0, 6, 12, 18)):
        lf = load_stage_forecast_horizon(base_full, load[forecast_day], h, cfg)
        pf = pv_stage_forecast_horizon_10min( forecast_day, h, pv_hourly, pv,
            cfg)
        la = actual_horizon_from_arrays(forecast_day, h, load, cfg)
        pa = actual_horizon_from_arrays(forecast_day, h, pv, cfg)
        if la is None or pa is None:
            continue
        stage_net_error_hist[forecast_day, s, :] = (la - pa) - (lf - pf)
def price_stage_forecast_horizon( price_fc_full:np.ndarray,
    price_actual_today:np.ndarray, stage_hour:int, cfg:Config)->np.ndarray:
    start=stage_hour*6; H=cfg.rolling_horizon_slots
    seg=np.asarray(price_fc_full[start:start+H],dtype=float).copy()
    if len(seg)!=H: raise ValueError("48h电价预测长度不足，无法构造24h滚动窗口。")
    if stage_hour<=0: return np.maximum(seg,0.0)
    m=min(start,cfg.price_intraday_recent_slots)
    base_today=np.asarray(price_fc_full[:cfg.slots_per_day],dtype=float)
    resid=np.asarray( price_actual_today[start- m:start],
        dtype=float)-base_today[ start- m:start]
    w=np.exp(np.linspace(-1.5,0.0,m)); w/=w.sum(); bias=float(np.sum(w*resid))
    bias=float( np.clip(bias, - cfg.price_intraday_bias_clip,
        cfg.price_intraday_bias_clip))
    tau=max(cfg.price_intraday_decay_hours*6.0,1.0)
    decay=np.exp(-np.arange(H,dtype=float)/tau)
    return np.maximum(seg+bias*decay,0.0)
def try_load_q42_forecast_cache(cfg:Config,D:int,logger):
    if not cfg.reuse_q42_forecast_cache: return None,None
    p=Path(cfg.q42_cache)
    if not p.exists():
        logger.info("未找到Q4-2缓存，将按同一方法重建负荷与电价预测。"); return None,None
    try:
        z=np.load(p,allow_pickle=False); lf=np.asarray(z['load_fc'],dtype=float)
        pf=np.asarray(z['price_fc'],dtype=float)
        if lf.shape[ 0]<D or pf.shape[ 0]<D or lf.shape[
            1]<cfg.forecast_slots or pf.shape[ 1]<cfg.forecast_slots:
            raise ValueError(f"缓存维度不足 load={lf.shape}, price={pf.shape}")
        logger.info("【预测缓存】已复用Q4-2的0:00负荷/48h电价预测，跳过全年MSTL重建。")
        return lf[ :D, :cfg.forecast_slots].copy( ),pf[ :D,
            :cfg.forecast_slots].copy( )
    except Exception as e:
        logger.warning("【预测缓存】读取失败：%s；将重新预测。",e); return None,None
@dataclass
class DayAheadTemplates:
    Aeq:sparse.csr_matrix; bounds:list
def build_day_ahead_templates(cfg:Config)->DayAheadTemplates:
    n=cfg.slots_per_day; N=5*n; ib,ic,idis,isu,ie=0,n,2*n,3*n,4*n
    rows=[]; cols=[]; data=[]; r=0
    for t in range(n):
        for col,val in ((ib+t,1.0),(idis+t,1.0),(ic+t,-1.0),(isu+t,-1.0)):
            rows.append(r); cols.append(col); data.append(val)
        r+=1
    for t in range(n):
        rows.append(r); cols.append(ie+t); data.append(1.0)
        if t>0: rows.append(r); cols.append(ie+t-1); data.append(-1.0)
        rows.append(r); cols.append(ic+t); data.append(-cfg.eta_c*cfg.dt)
        rows.append(r); cols.append(idis+t); data.append(cfg.dt/cfg.eta_d); r+=1
    Aeq=sparse.csr_matrix((data,(rows,cols)),shape=(r,N))
    bounds=( [(0, None)]* n+ [(0, cfg.p_max)]* n+ [(0, cfg.p_max)]* n+ [(0,
        None)]* n+ [(cfg.e_min, cfg.e_max)]* n)
    return DayAheadTemplates(Aeq,bounds)
def solve_day_ahead_price( price_fc_full:np.ndarray, load_fc:np.ndarray,
    pv_fc:np.ndarray, reserve:np.ndarray, soc0:float, cfg:Config,
    lp:DayAheadTemplates):
    n=cfg.slots_per_day; pr=np.asarray(price_fc_full,dtype=float)
    if len(pr)<2*n: raise ValueError("日前电价预测必须包含未来48h。")
    price_today=pr[:n]; price_next=pr[n:2*n]; c=np.zeros(5*n)
    c[:n]=price_today*cfg.dt; c[n:3*n]=cfg.eps_cycle*cfg.dt
    c[3*n:4*n]=cfg.eps_surplus*cfg.dt
    mu=cfg.eta_d*float(np.min(price_next)); c[5*n-1]=-mu
    b=np.empty(2*n,dtype=float)
    b[:n]=np.asarray(load_fc[:n])-np.asarray(pv_fc[:n])+np.asarray(reserve)
    b[n:]=0.0; b[n]=soc0
    res=linprog(c,A_eq=lp.Aeq,b_eq=b,bounds=lp.bounds,method=cfg.lp_method)
    if not res.success: raise RuntimeError(f"0:00日前LP失败：{res.message}")
    x=res.x
    return { 'buy':x[:n], 'ch':x[n:2* n], 'dis':x[2* n:3* n], 'sur':x[3* n:4*n],
        'E':x[4* n:5* n], 'mu':mu}
def solve_rolling_mpc_24h_price( price_h:np.ndarray, load_fc_horizon:np.ndarray,
    pv_fc_horizon:np.ndarray, reserve_horizon:np.ndarray, soc0:float,
    prev_buy_day:np.ndarray, start_slot:int, cfg:Config):
    H=cfg.rolling_horizon_slots; price_h=np.asarray(price_h,dtype=float)
    if not( len(price_h)== len(load_fc_horizon)== len(pv_fc_horizon)==
        len(reserve_horizon)== H):
        raise ValueError("24h滚动MPC全部输入必须为144点。")
    current_remaining=cfg.slots_per_day-start_slot
    N=7*H; ib,ic,idis,isu,ie,iup,idn=0,H,2*H,3*H,4*H,5*H,6*H
    rows=[]; cols=[]; data=[]; b=[]; r=0
    net=np.asarray( load_fc_horizon)-np.asarray( pv_fc_horizon)+np.asarray(
        reserve_horizon)
    for j in range(H):
        for col,val in ((ib+j,1.0),(idis+j,1.0),(ic+j,-1.0),(isu+j,-1.0)):
            rows.append(r); cols.append(col); data.append(val)
        b.append(float(net[j])); r+=1
    for j in range(H):
        rows.append(r); cols.append(ie+j); data.append(1.0)
        if j>0: rows.append(r); cols.append(ie+j-1); data.append(-1.0); rhs=0.0
        else: rhs=float(soc0)
        rows.append(r); cols.append(ic+j); data.append(-cfg.eta_c*cfg.dt)
        rows.append(r); cols.append(idis+j); data.append(cfg.dt/cfg.eta_d)
        b.append(rhs); r+=1
    for j in range(current_remaining):
        rows.append(r); cols.append(ib+j); data.append(1.0)
        rows.append(r); cols.append(iup+j); data.append(-1.0)
        rows.append(r); cols.append(idn+j); data.append(1.0)
        b.append(float(prev_buy_day[start_slot+j])); r+=1
    Aeq=sparse.csr_matrix((data,(rows,cols)),shape=(r,N)); c=np.zeros(N)
    c[iup:iup+current_remaining]=1.5*price_h[:current_remaining]*cfg.dt
    c[idn:idn+current_remaining]=-0.5*price_h[:current_remaining]*cfg.dt
    if current_remaining<H: c[ ib+ current_remaining:ib+ H]=price_h[
        current_remaining:]*cfg.dt
    c[ic:ic+H]=cfg.eps_cycle*cfg.dt; c[idis:idis+H]=cfg.eps_cycle*cfg.dt
    c[isu:isu+H]=cfg.eps_surplus*cfg.dt
    # 窗口末端机会价值：未来24h预测价格中最低可替代购电价
    mu=cfg.eta_d*float(np.min(price_h)); c[ie+H-1]=-mu; bounds=[]
    bounds += [(0,None)]*H; bounds += [(0,cfg.p_max)]*H
    bounds += [(0,cfg.p_max)]*H
    bounds += [(0,None)]*H; bounds += [(cfg.e_min,cfg.e_max)]*H
    bounds += [(0,None) if j<current_remaining else (0,0) for j in range(H)]
    bounds += [(0,None) if j<current_remaining else (0,0) for j in range(H)]
    res=linprog( c, A_eq=Aeq, b_eq=np.asarray(b), bounds=bounds,
        method=cfg.lp_method)
    if not res.success: raise RuntimeError(
        f"{start_slot//6}:00滚动LP失败：{res.message}")
    x=res.x; up=x[iup:iup+H]; dn=x[idn:idn+H]
    return { 'buy':x[ib:ib+ H], 'ch':x[ic:ic+ H], 'dis':x[idis:idis+ H],
        'sur':x[isu:isu+ H], 'E':x[ie:ie+ H], 'up':up, 'down':dn,
        'up_energy':float(np.sum(up[:current_remaining])* cfg.dt),
        'down_energy':float(np.sum(dn[:current_remaining])* cfg.dt), 'mu':mu,
        'current_remaining':current_remaining}
def _copy_style(src,dst):
    from copy import copy
    if src.has_style: dst._style=copy(src._style)
    dst.font=copy(src.font); dst.fill=copy(src.fill)
    dst.border=copy(src.border); dst.alignment=copy(src.alignment)
    dst.protection=copy(src.protection); dst.number_format=src.number_format
def fill_result4_3( template_path, out_path, dates, price_actual, initial_buy,
    effective_buy, buy_cost_after_adjust, ch, dis, soc_start, soc_end, emg, cfg,
    logger):
    if not Path(template_path).exists():
        logger.warning("【模板缺失】未找到 %s；CSV/TXT/图仍正常输出。",template_path); return
    wb=load_workbook(template_path); sp=wb['计划购电量']; sa=wb['调整购电量']
    sc=wb['充放电量']; se=wb['紧急购电量']; date_to_idx={d:i for i,d in enumerate(dates)}
    computed=[ d for d in dates if d>= date(2025, 2, 1) and
        np.isfinite(initial_buy[date_to_idx[d]]).all()]
    init_cost=np.nansum(initial_buy*price_actual,axis=1)*cfg.dt
    def fill_buy(ws,arr,costs):
        for r in range(2,ws.max_row+1):
            raw=ws.cell(r,1).value
            if raw is None: continue
            try:d=cell_to_date(raw)
            except Exception:continue
            if d not in date_to_idx: continue
            i=date_to_idx[d]
            if not np.isfinite(arr[i]).all(): continue
            vals=arr[i]*cfg.dt
            for c,v in enumerate(vals,start=2): ws.cell(r,c).value=float(v)
            ws.cell(r,146).value=float(np.sum(vals))
            ws.cell(r,147).value=float(costs[i])
    fill_buy(sp,initial_buy,init_cost)
    fill_buy(sa,effective_buy,buy_cost_after_adjust)
    charge_proto=[ {'height':sc.row_dimensions[r].height, 'cells':[sc.cell(r,
        c) for c in range(1, 7)]} for r in range(2, 8)]
    emg_proto=[ {'height':se.row_dimensions[r].height, 'cells':[se.cell(r,c) for
        c in range(1, 4)]} for r in range(2, min(5, se.max_row+ 1))]
    if sc.max_row>=2: sc.delete_rows(2,sc.max_row-1)
    intervals=[ (0, 4, '0:00-4:00'), (4, 8, '4:00-8:00'), (8, 12, '8:00-12:00'),
        (12, 16, '12:00-16:00'), (16, 20, '16:00-20:00'), (20, 24,
        '20:00-24:00')]
    rr=2
    for d in computed:
        i=date_to_idx[d]
        for k,(h0,h1,label) in enumerate(intervals):
            proto=charge_proto[k]; sc.row_dimensions[rr].height=proto['height']
            for c in range(1,7): _copy_style(proto['cells'][c-1],sc.cell(rr,c))
            sc.cell(rr,1).value=d if k==0 else None; sc.cell(rr,2).value=label
            a,b=h0*6,h1*6; sc.cell(rr,3).value=float(np.sum(ch[i,a:b])*cfg.dt)
            sc.cell(rr,4).value=float(np.sum(dis[i,a:b])*cfg.dt)
            if k==0: sc.cell( rr, 5).value='0:00'; sc.cell( rr, 6).value=float(
                soc_start[i])
            elif k==1: sc.cell( rr, 5).value='24:00'; sc.cell( rr,
                6).value=float( soc_end[i])
            rr+=1
    def fmtslot(k):
        mins=k*10
        if mins==1440:return '24:00'
        return f'{(mins//60)%24}:{mins%60:02d}'
    if se.max_row>=2:se.delete_rows(2,se.max_row-1)
    rr=2
    for d in computed:
        i=date_to_idx[d]; events=[]; t=0
        while t<144:
            if emg[i,t]>1e-6:
                st=t; ee=0.0
                while t<144 and emg[i,t]>1e-6: ee+=emg[i,t]*cfg.dt; t+=1
                events.append((f'{fmtslot(st)}-{fmtslot(t)}',ee))
            else:t+=1
        if not events:events=[(None,None)]
        for k,(period,energy) in enumerate(events):
            proto=emg_proto[min(k,len(emg_proto)-1)]
            se.row_dimensions[rr].height=proto['height']
            for c in range(1,4):_copy_style(proto['cells'][c-1],se.cell(rr,c))
            se.cell(rr,1).value=d if k==0 else None; se.cell(rr,2).value=period
            se.cell(rr,3).value=energy; rr+=1
    for ws in (sc,se):
        for r in range(2,ws.max_row+1):
            if ws.cell( r, 1).value is not None: ws.cell( r,
                1).number_format='yyyy/m/d'
    Path(out_path).parent.mkdir(parents=True,exist_ok=True); wb.save(out_path)
    wb.close(); logger.info("【结果输出】result4-3已保存：%s",out_path)
# 题目表1要求给出的指定时间段（10分钟）
PAPER_TABLE1_INTERVALS = ( "10:00-10:10", "12:00-12:10", "14:00-14:10",
    "16:00-16:10", "18:00-18:10", "20:00-20:10")
# 题目表2要求给出的指定时间段（4小时）
PAPER_TABLE2_INTERVALS = ( "0:00-4:00", "4:00-8:00", "8:00-12:00","12:00-16:00",
    "16:00-20:00", "20:00-24:00")
def _interval_start_slot(label: str) -> int:
    """把'10:00-10:10'这类标签映射到行内列序号(1-based列-2)。标签起点=(slot+1)*10分钟。"""
    hh, mm = label.split("-")[0].split(":")
    return (int(hh) * 60 + int(mm)) // 10 - 1
def _slot_label(k: int) -> str:
    mins = k * 10
    if mins == 1440:
        return "24:00"
    return f"{(mins // 60) % 24}:{mins % 60:02d}"
def _emergency_events(emg_row: np.ndarray, dt: float) -> list:
    """把逐时段紧急购电功率聚合成[(时间段, 购电量kWh), ...]。"""; events = []; t = 0
    n = len(emg_row)
    while t < n:
        if emg_row[t] > 1e-6:
            st = t; ee = 0.0
            while t < n and emg_row[t] > 1e-6:
                ee += float(emg_row[t]) * dt; t += 1
            events.append((f"{_slot_label(st)}-{_slot_label(t)}", ee))
        else:
            t += 1
    return events
def export_paper_tables_q43( dates, price_actual, initial_buy, effective_buy,ch,
    dis, emg, soc_start, soc_end, plan_cost, adj_cost, emg_cost, run_D, cfg,
    logger):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    out_path = Path(cfg.out_dir) / "Q4-3_论文表格_指定日期.xlsx"
    d2i = {d: i for i, d in enumerate(dates)}
    targets = [ d for d in cfg.representative_dates if d in d2i and d2i[d] <
        run_D]
    if not targets:
        logger.warning("【论文表格】表3指定日期不在本次运行范围内，跳过导出。"); return None
    dt = cfg.dt; thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True); title_font = Font(bold=True, size=12)
    center = Alignment(horizontal="center", vertical="center"); num = "#,##0.00"
    def put(ws, r, c, v, font=None, align=None, bd=True, fmt=None):
        cell = ws.cell(row=r, column=c, value=v)
        if font is not None:
            cell.font = font
        if align is not None:
            cell.alignment = align
        if bd:
            cell.border = border
        if fmt is not None:
            cell.number_format = fmt
        return cell
    wb = Workbook(); wb.remove(wb.active)
    colw = { "A": 15, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14,
        "H": 14}
    # 表1 
    ws = wb.create_sheet("表1_购电量")
    for col, w in colw.items():
        ws.column_dimensions[col].width = w
    r = 1
    for d in targets:
        i = d2i[d]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        put( ws, r, 1, f"表1 微网在指定时间段的购电量及全天的购电量和购电费（{d.strftime('%Y-%m-%d')}）",
            title_font, center, bd=False)
        r += 1
        for c, v in enumerate( ["时间段", "购电量(kWh)", "时间段", "购电量(kWh)", "时间段",
            "购电量(kWh)"], start=1):
            put(ws, r, c, v, bold, center)
        r += 1
        vals = [ float(effective_buy[i, _interval_start_slot(lb)]) * dt for
            lb in PAPER_TABLE1_INTERVALS]
        for tri in ((0, 1, 2), (3, 4, 5)):
            for col_off, idx in enumerate(tri):
                put( ws, r, 1 + 2 * col_off, PAPER_TABLE1_INTERVALS[idx], None,
                    center)
                put(ws, r, 2 + 2 * col_off, vals[idx], None, center, fmt=num)
            r += 1
        day_buy = float(np.nansum(effective_buy[i]) * dt)
        day_cost = float(plan_cost[i] + adj_cost[i] + emg_cost[i])
        put(ws, r, 1, "全天购电量", bold, center)
        put(ws, r, 2, day_buy, None, center, fmt=num)
        put(ws, r, 3, "全天购电费(元)", bold, center)
        put(ws, r, 4, day_cost, None, center, fmt=num)
        put(ws, r, 5, None, None, center); put(ws, r, 6, None, None, center)
        r += 2
    # 表2 
    ws = wb.create_sheet("表2_充放电量")
    for col, w in colw.items():
        ws.column_dimensions[col].width = w
    r = 1; ivs = [(0, 4), (4, 8), (8, 12), (12, 16), (16, 20), (20, 24)]
    for d in targets:
        i = d2i[d]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        put( ws, r, 1,
            f"表2 储能设备在指定时间段的充放电量及0:00和24:00的储电量（{d.strftime('%Y-%m-%d')}）",
            title_font, center, bd=False)
        r += 1
        for c, v in enumerate( ["时间段", "充电量(kWh)", "放电量(kWh)", "时间段","充电量(kWh)",
            "放电量(kWh)"], start=1):
            put(ws, r, c, v, bold, center)
        r += 1
        ch_e = [float(np.sum(ch[i, h0 * 6:h1 * 6]) * dt) for h0, h1 in ivs]
        dis_e = [float(np.sum(dis[i, h0 * 6:h1 * 6]) * dt) for h0, h1 in ivs]
        for a, b in ((0, 1), (2, 3), (4, 5)):
            put(ws, r, 1, PAPER_TABLE2_INTERVALS[a], None, center)
            put(ws, r, 2, ch_e[a], None, center, fmt=num)
            put(ws, r, 3, dis_e[a], None, center, fmt=num)
            put(ws, r, 4, PAPER_TABLE2_INTERVALS[b], None, center)
            put(ws, r, 5, ch_e[b], None, center, fmt=num)
            put(ws, r, 6, dis_e[b], None, center, fmt=num); r += 1
        put(ws, r, 1, "0:00 储电量", bold, center)
        put(ws, r, 2, float(soc_start[i]), None, center, fmt=num)
        put(ws, r, 3, "24:00 储电量", bold, center)
        put(ws, r, 4, float(soc_end[i]), None, center, fmt=num)
        put(ws, r, 5, None, None, center); put(ws, r, 6, None, None, center)
        r += 2
    # 表3 
    ws = wb.create_sheet("表3_紧急购电量")
    for col, w in colw.items():
        ws.column_dimensions[col].width = w
    events_per_date = [_emergency_events(emg[d2i[d]], dt) for d in targets]
    ncol = 2 * len(targets)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    put(ws, 1, 1, "表3 微网在指定日期的紧急购电量", title_font, center, bd=False)
    for j, d in enumerate(targets):
        c0 = 1 + 2 * j
        ws.merge_cells( start_row=2, start_column=c0, end_row=2, end_column=c0 +
            1)
        put(ws, 2, c0, d.strftime("%Y.%m.%d"), bold, center)
        put(ws, 2, c0 + 1, None, bold, center)
    for j in range(len(targets)):
        put(ws, 3, 1 + 2 * j, "时间段", bold, center)
        put(ws, 3, 2 + 2 * j, "购电量(kWh)", bold, center)
    nrow = max(3, max((len(e) for e in events_per_date), default=0)); r = 4
    for k in range(nrow):
        for j, evs in enumerate(events_per_date):
            c0 = 1 + 2 * j
            if k < len(evs):
                put(ws, r, c0, evs[k][0], None, center)
                put(ws, r, c0 + 1, evs[k][1], None, center, fmt=num)
            else:
                put(ws, r, c0, None, None, center)
                put(ws, r, c0 + 1, None, None, center)
        r += 1
    # 附表：费用与“0点计划/最终”对照 
    ws = wb.create_sheet("附_费用与计划对照")
    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 18
    heads = [ "日期", "计划全天购电量(kWh)", "最终全天购电量(kWh)", "计划购电费(元)", "调整净费用(元)",
        "紧急购电费(元)", "全天总费用(元)", "紧急购电量(kWh)"]
    for c, v in enumerate(heads, start=1):
        put(ws, 1, c, v, bold, center)
    r = 2
    for d in targets:
        i = d2i[d]
        row = [ d.strftime("%Y-%m-%d"), float(np.nansum(initial_buy[i]) * dt),
            float(np.nansum(effective_buy[i]) * dt), float(plan_cost[i]),
            float(adj_cost[i]), float(emg_cost[i]), float(plan_cost[i] +
            adj_cost[i] + emg_cost[i]), float(np.sum(emg[i]) * dt)]
        for c, v in enumerate(row, start=1):
            put( ws, r, c, v, None, center if c > 1 else None, fmt=(None if c ==
                1 else num))
        r += 1
    r += 2; put(ws, r, 1, "指定时段：0:00计划购电量与最终购电量对照", bold); r += 1
    for c, v in enumerate(["日期", "时间段", "计划购电量(kWh)", "最终购电量(kWh)"], start=1):
        put(ws, r, c, v, bold, center)
    r += 1
    for d in targets:
        i = d2i[d]
        for lb in PAPER_TABLE1_INTERVALS:
            s = _interval_start_slot(lb)
            put(ws, r, 1, d.strftime("%Y-%m-%d"), None, center)
            put(ws, r, 2, lb, None, center)
            put(ws, r, 3, float(initial_buy[i, s]) * dt, None, center, fmt=num)
            put( ws, r, 4, float(effective_buy[i, s]) * dt, None, center,
                fmt=num)
            r += 1
    # 说明 
    ws = wb.create_sheet("说明"); ws.column_dimensions["A"].width = 110
    notes = ["问题4-3 论文表格（表1/表2/表3）—— 指定日期：2025.3.20、2025.6.21、2025.9.23、2025.12.21",
        "波动电价场景：计划购电按0:00预测电价制定，结算与紧急购电（5倍）均按真实电价计算。",
        "表1 购电量：取“调整后的最终购电量”（0:00计划经6:00/12:00/18:00预报调整后的结果）。",
        "表1 全天购电费：= 计划购电费 + 调整净费用 + 紧急购电费（问题4-3总购电费用口径）。",
        "表2 充放电量：为指定4小时时间段的最终充电量与放电量；0:00/24:00 储电量为当日首末储电量。",
        "表3 紧急购电量：为最终紧急购电的时间段与电量；留空表示该日期没有发生紧急购电。",
        "附表：给出每日0:00计划与最终购电量的对照及费用拆分，便于论文中解释“调整”的作用。", ]
    for i, t in enumerate(notes, start=1):
        put(ws, i, 1, t, bold if i == 1 else None, bd=False)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True); wb.save(out_path)
    wb.close()
    logger.info( "【论文表格】表1/表2/表3已导出：%s（日期：%s）", out_path,
        "、".join(d.strftime("%Y-%m-%d") for d in targets))
    return str(out_path)
def save_outputs( dates, price_actual, price_fc0, price_stage_fc, initial_buy,
    effective_buy, ch, dis, emg, sur, soc_start, soc_end, plan_cost, adj_cost,
    emg_cost, up_energy, down_energy, load_rmse, pv_rmse0, run_D, cfg, logger):
    out=Path(cfg.out_dir)
    out.mkdir(parents=True,exist_ok=True)
    rows=[]
    for d in range(run_D):
        rows.append( [str(dates[d]), soc_start[d], soc_end[d], plan_cost[d],
            adj_cost[d], emg_cost[d], plan_cost[d]+ adj_cost[d]+ emg_cost[d],
            float(np.sum(emg[d])* cfg.dt), float(np.sum(sur[d])* cfg.dt),
            up_energy[d], down_energy[d], load_rmse[d], pv_rmse0[d]])
    with open( out/ 'q4_3_daily_summary.csv', 'w', newline='',
        encoding='utf-8-sig') as f:
        w=csv.writer(f)
        w.writerow( ['日期', '日初SOC', '日末SOC', '初始计划购电费', '调整净费用', '紧急购电费', '总费用',
            '紧急购电量kWh', '富余电量kWh', '向上调整电量kWh', '向下调整电量kWh', '负荷RMSE',
            '0点光伏RMSE'])
        w.writerows(rows)
    if run_D<32:return
    valid=[i for i in range(31,run_D)]; formal=np.array(valid,dtype=int)
    total_plan=float(np.sum(plan_cost[formal]))
    total_adj=float(np.sum(adj_cost[formal]))
    total_emg=float(np.sum(emg_cost[formal]))
    emg_energy=float(np.sum(emg[formal])*cfg.dt)
    sur_energy=float(np.sum(sur[formal])*cfg.dt)
    zero=int(np.sum(np.sum(emg[formal],axis=1)*cfg.dt<1e-6))
    daily_emg=np.sum(emg[formal],axis=1)*cfg.dt
    p95=float(np.quantile(daily_emg,.95)); mx=float(np.max(daily_emg))
    a=price_actual[formal].reshape(-1); p=price_fc0[formal,:144].reshape(-1)
    mask=np.isfinite(a)&np.isfinite(p)
    mae=float(np.mean(np.abs(a[mask]-p[mask])))
    rmse=float(np.sqrt(np.mean((a[mask]-p[mask])**2)))
    mape=float( np.mean(np.abs((a[mask]- p[mask])/ np.maximum(np.abs(a[mask]),
        1e-3)))* 100)
    lines=[ '问题4-3正式结果摘要（电价预测驱动，1月仅预热）', '='* 68,
        f'正式统计日期范围：{dates[31]} 至 {dates[run_D- 1]}',
        f'累计初始计划购电费（真实电价结算）：{total_plan:.2f} 元',
        f'累计调整净费用（真实电价结算）：{total_adj:.2f} 元',
        f'累计紧急购电费（5倍真实电价）：{total_emg:.2f} 元', f'累计总费用：{total_plan+ total_adj+total_emg:.2f} 元', f'累计紧急购电量：{emg_energy:.2f} kWh',
        f'累计富余电量：{sur_energy:.2f} kWh',
        f'累计向上调整电量：{float(np.sum(up_energy[formal])):.2f} kWh',
        f'累计向下调整电量：{float(np.sum(down_energy[formal])):.2f} kWh',
        f'日紧急购电P95：{p95:.2f} kWh', f'最大单日紧急购电：{mx:.2f} kWh',
        f'无紧急购电天数：{zero}/{len(formal)}', f'期末储电量：{soc_end[run_D- 1]:.2f} kWh',
        f'0:00电价预测MAE：{mae:.6f} 元/kWh', f'0:00电价预测RMSE：{rmse:.6f} 元/kWh',
        f'0:00电价预测MAPE：{mape:.3f}%', '', f'正式结果文件：{cfg.out_xlsx}']
    (out/'Q4-3_结果摘要.txt').write_text('\n'.join(lines),encoding='utf-8')
    # 各发布时点未来24h电价预测误差
    metrics=[]
    for si,h in enumerate((0,6,12,18)):
        aa=[]; pp=[]
        for d in valid:
            real=actual_horizon_from_arrays(d,h,price_actual,cfg)
            pred=price_stage_fc[d,si]
            if real is None or not np.all(np.isfinite(pred)):continue
            aa.append(real); pp.append(pred)
        if aa:
            aa=np.concatenate(aa); pp=np.concatenate(pp); er=pp-aa
            metrics.append( [h, float(np.mean(np.abs(er))),
                float(np.sqrt(np.mean(er* er))), float(np.mean(np.abs(er)/
                np.maximum(np.abs(aa), 1e-3))* 100)])
    with open( out/ 'Q4-3_电价分阶段预测指标.csv', 'w', newline='',
        encoding='utf-8-sig') as f:
        w=csv.writer(f); w.writerow(['发布时点/h','MAE元每kWh','RMSE元每kWh','MAPE%'])
        w.writerows(metrics)
    # 指定日期表1/2/3简表
    target_rows=[]; date_to_idx={d:i for i,d in enumerate(dates)}
    for dd in cfg.representative_dates:
        i=date_to_idx.get(dd)
        if i is None:continue
        target_rows.append( [str(dd), float(np.sum(initial_buy[i])* cfg.dt),
            float(np.sum(effective_buy[i])* cfg.dt), plan_cost[i], adj_cost[i],
            emg_cost[i], float(np.sum(emg[i])* cfg.dt), soc_start[i],
            soc_end[i]])
    with open(out/'Q4-3_指定日期结果汇总.csv','w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f)
        w.writerow( ['日期', '初始计划购电量kWh', '最终调整购电量kWh', '初始计划购电费元', '调整净费用元',
            '紧急购电费元', '紧急购电量kWh', '0:00SOC', '24:00SOC'])
        w.writerows(target_rows)
    logger.info("Q4-3摘要及CSV结果已保存：%s", out)
def simulate_update_scheme_q43( update_hours, dates, price_actual, load, pv,
    load_fc_full_store, price_fc_full_store, pv_hourly, fallback_reserve_store,
    stage_net_error_hist, run_D, cfg, lp, collect_response=False):
    n=cfg.slots_per_day; update_hours=tuple(sorted(update_hours))
    plan_cost=np.zeros(run_D); adj_cost=np.zeros(run_D)
    emg_cost=np.zeros(run_D)
    emg_energy=np.zeros(run_D); surplus_energy=np.zeros(run_D)
    up_energy=np.zeros(run_D); down_energy=np.zeros(run_D)
    soc_start=np.zeros(run_D); soc_end=np.zeros(run_D)
    response_records=[]; soc=cfg.e0; stage_idx_map={0:0,6:1,12:2,18:3}
    for d in range(run_D):
        soc_start[d]=soc; lf_full=np.asarray(load_fc_full_store[d],dtype=float)
        pr_full=np.asarray(price_fc_full_store[d],dtype=float); lf0=lf_full[:n]
        pf0=pv_stage_forecast_horizon_10min(d,0,pv_hourly,pv,cfg)[:n]
        fallback=np.asarray(fallback_reserve_store[d],dtype=float)
        da=solve_day_ahead_price(pr_full,lf0,pf0,fallback,soc,cfg,lp)
        initial_buy=da['buy'].copy(); current_buy=da['buy'].copy()
        current_ch=da['ch'].copy(); current_dis=da['dis'].copy()
        plan_cost[d]=float(np.sum(initial_buy*price_actual[d])*cfg.dt)
        # 上一阶段“已知信息”基准：用于12:00重点响应分析。
        prev_pv_day=pf0.copy(); prev_net_day=(lf0-pf0+fallback).copy()
        prev_price_day=pr_full[:n].copy(); stages=(0,)+update_hours
        for si,h in enumerate(stages):
            start=h*6
            if h>0:
                lf_h=load_stage_forecast_horizon(lf_full,load[d],h,cfg)
                pf_h=pv_stage_forecast_horizon_10min(d,h,pv_hourly,pv,cfg)
                reserve_h=stage_reserve_from_history( d, stage_idx_map[h],
                    stage_net_error_hist, fallback, h, cfg)
                pr_h=price_stage_forecast_horizon(pr_full,price_actual[d],h,cfg)
                m=n-start; buy_before=current_buy[start:].copy()
                ch_before=current_ch[start:].copy()
                dis_before=current_dis[start:].copy()
                pv_before=prev_pv_day[start:].copy()
                net_before=prev_net_day[start:].copy()
                price_before=prev_price_day[start:].copy()
                net_after=(lf_h[:m]-pf_h[:m]+reserve_h[:m]).copy()
                price_after=pr_h[:m].copy()
                adj=solve_rolling_mpc_24h_price( pr_h, lf_h, pf_h, reserve_h,
                    soc, current_buy, start, cfg)
                buy_after=adj['buy'][:m].copy(); ch_after=adj['ch'][:m].copy()
                dis_after=adj['dis'][:m].copy(); current_buy[start:]=buy_after
                current_ch[start:]=ch_after; current_dis[start:]=dis_after
                # 与正式Q4-3一致：真实价格只用于事后结算。
                upv=adj['up'][:m]; dnv=adj['down'][:m]
                real_today=price_actual[d,start:]
                adj_cost[ d]+=float( np.sum((1.5* real_today* upv- 0.5*
                    real_today* dnv)* cfg.dt))
                up_energy[d]+=adj['up_energy']
                down_energy[d]+=adj['down_energy']
                if collect_response:
                    for j in range(m):
                        delta_buy=float(buy_after[j]-buy_before[j])
                        response_records.append( { '日期':str(dates[d]), '更新时点':h,
                            '目标时刻':f"{(start+ j)//6:02d}:{((start+ j)%6)*10:02d}", '旧光伏预报kW':float(pv_before[j]),
                            '新光伏预报kW':float(pf_h[j]), '光伏预报修正量kW':float(pf_h[j]-
                            pv_before[j]), '旧预测电价元每kWh':float(price_before[j]),
                            '新预测电价元每kWh':float(price_after[j]),
                            '电价预测修正量元每kWh':float(price_after[j]-
                            price_before[j]), '旧有效净负荷kW':float(net_before[j]),
                            '新有效净负荷kW':float(net_after[j]),
                            '有效净负荷修正量kW':float(net_after[j]- net_before[j]),
                            '旧购电计划kW':float(buy_before[j]),
                            '新购电计划kW':float(buy_after[j]), '购电调整量kW':delta_buy,
                            '向上调整kW':max(delta_buy, 0.0), '向下调整kW':max(-
                            delta_buy, 0.0), '原充电计划kW':float(ch_before[j]),
                            '新充电计划kW':float(ch_after[j]),
                            '原放电计划kW':float(dis_before[j]),
                            '新放电计划kW':float(dis_after[j]), })
                prev_pv_day[start:]=pf_h[:m]; prev_net_day[start:]=net_after
                prev_price_day[start:]=price_after
            end=(stages[si+1]*6) if si+1<len(stages) else n
            for t in range(start,end):
                sol=operate_one_slot( soc, current_buy[t], load[d, t], pv[d, t],
                    current_ch[t], current_dis[t], cfg)
                soc=sol['E']; emg_energy[d]+=sol['emg']*cfg.dt
                surplus_energy[d]+=sol['sur']*cfg.dt
                emg_cost[d]+=sol['emg']*5.0*price_actual[d,t]*cfg.dt
        soc_end[d]=soc
    return { 'plan_cost':plan_cost, 'adj_cost':adj_cost, 'emg_cost':emg_cost,
        'emg_energy':emg_energy, 'surplus_energy':surplus_energy,
        'up_energy':up_energy, 'down_energy':down_energy, 'soc_start':soc_start,
        'soc_end':soc_end, 'response_records':response_records, }
def run_update_time_comparison_q43( dates, price_actual, load, pv,
    load_fc_full_store, price_fc_full_store, pv_hourly, fallback_reserve_store,
    stage_net_error_hist, run_D, cfg, lp, logger):
    """复现原Q3五方案对照，并扩展到“光伏+电价”同步滚动更新。"""
    if run_D<32:
        logger.info('测试天数不足32天，跳过正式期方案比较。'); return None
    out=Path(cfg.out_dir)
    out.mkdir(parents=True,exist_ok=True)
    formal=slice(31,run_D); results=[]; raw={}
    logger.info('开始五方案比较；负荷预测、价格基础预测、风险参数和储能参数保持一致。')
    for name,hours in zip(cfg.comparison_names,cfg.comparison_schemes):
        logger.info('  正在计算方案：%s',name)
        r=simulate_update_scheme_q43( hours, dates, price_actual, load, pv,
            load_fc_full_store, price_fc_full_store, pv_hourly,
            fallback_reserve_store, stage_net_error_hist, run_D, cfg, lp,
            collect_response=(tuple(hours)== tuple(cfg.focus_scheme_hours)))
        raw[name]=r; plan=float(np.sum(r['plan_cost'][formal]))
        adj=float(np.sum(r['adj_cost'][formal]))
        emgc=float(np.sum(r['emg_cost'][formal]))
        daily=np.asarray(r['emg_energy'][31:run_D],dtype=float)
        results.append( { '方案':name, '更新时点':','.join(map(str, hours)) if
            hours else '无日内更新', '初始计划购电费':plan, '调整净费用':adj, '紧急购电费':emgc,
            '总费用':plan+ adj+ emgc, '紧急购电量kWh':float(np.sum(daily)),
            '富余电量kWh':float(np.sum(r['surplus_energy'][formal])),
            '向上调整电量kWh':float(np.sum(r['up_energy'][formal])),
            '向下调整电量kWh':float(np.sum(r['down_energy'][formal])),
            '紧急购电天数':int(np.sum(daily> 1e-6)),
            '单日最大紧急购电量kWh':float(np.max(daily)) if len(daily) else 0.0,
            '日紧急购电量P95kWh':float(np.quantile(daily, 0.95)) if len(daily) else
            0.0, '期末SOC':float(r['soc_end'][run_D- 1]), })
    csv_path=out/'Q4-3_既有预报时点方案对照.csv'
    with open(csv_path,'w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(results[0].keys())); w.writeheader()
        w.writerows(results)
    base=results[0]; txt=out/'Q4-3_既有预报时点方案对照结论.txt'
    with open(txt,'w',encoding='utf-8') as f:
        f.write('Q4-3既有预报时点价值对照——波动电价下24小时滚动线性规划\n'+'='*68+'\n')
        f.write('说明：固定负荷预测、0:00基础电价预测、风险参数和储能参数；仅改变6/12/18哪些时点允许同步更新光伏预报与电价预报。\n')
        for r in results:
            f.write(
                f"{r['方案']}：总费用 {r['总费用']:.2f} 元，紧急购电 {r['紧急购电量kWh']:.2f} kWh，P95 {r['日紧急购电量P95kWh']:.2f} kWh，"
                f"最大单日 {r['单日最大紧急购电量kWh']:.2f} kWh，富余 {r['富余电量kWh']:.2f} kWh。\n")
        f.write('\n相对仅0:00方案：\n')
        for r in results[1:]:
            f.write( f"{r['方案']}：总费用变化 {r['总费用']-base['总费用']:+.2f} 元，紧急购电量减少 {base['紧急购电量kWh']-r['紧急购电量kWh']:.2f} kWh，富余电量减少 {base['富余电量kWh']- r['富余电量kWh']:.2f} kWh。\n")
    focus=raw.get(cfg.focus_scheme_name)
    if focus is not None and focus.get('response_records'):
        recs=[ r for r in focus['response_records'] if r['更新时点']== 12 and
            str(r['日期'])>= str(dates[31])]
        if recs:
            detail_path=out/'Q4-3_0点与12点预报决策响应明细.csv'
            with open(detail_path,'w',newline='',encoding='utf-8-sig') as f:
                w=csv.DictWriter(f,fieldnames=list(recs[0].keys()))
                w.writeheader(); w.writerows(recs)
            dpv=np.array([r['光伏预报修正量kW'] for r in recs],dtype=float)
            dprice=np.array([r['电价预测修正量元每kWh'] for r in recs],dtype=float)
            dnet=np.array([r['有效净负荷修正量kW'] for r in recs],dtype=float)
            dbuy=np.array([r['购电调整量kW'] for r in recs],dtype=float)
            def corr( a, b): return float( np.corrcoef(a, b)[0, 1]) if np.std(
                a)>1e-9 and np.std( b)>1e-9 else np.nan
            cpv=corr(dpv,dbuy); cprice=corr(dprice,dbuy); cnet=corr(dnet,dbuy)
            up_kwh=float(np.sum(np.maximum(dbuy,0))*cfg.dt)
            down_kwh=float(np.sum(np.maximum(-dbuy,0))*cfg.dt)
            br=raw[cfg.comparison_names[0]]; fr=focus
            bt=float( np.sum(br['plan_cost'][formal])+
                np.sum(br['adj_cost'][formal])+ np.sum(br['emg_cost'][formal]))
            ft=float( np.sum(fr['plan_cost'][formal])+
                np.sum(fr['adj_cost'][formal])+ np.sum(fr['emg_cost'][formal]))
            bemg=float(np.sum(br['emg_energy'][formal]))
            femg=float(np.sum(fr['emg_energy'][formal]))
            bsur=float(np.sum(br['surplus_energy'][formal]))
            fsur=float(np.sum(fr['surplus_energy'][formal]))
            with open(out/'Q4-3_0点与12点预报重点结论.txt','w',encoding='utf-8') as f:
                f.write('Q4-3：0:00 与 12:00 预报的边际价值分析\n'+'='*62+'\n')
                f.write('12:00同时更新附件3光伏预报与基于已实现实时价格的电价预测，再重解未来24小时LP。\n')
                f.write(f'corr(光伏预报修正量, 购电调整量) = {cpv:.4f}\n')
                f.write(f'corr(电价预测修正量, 购电调整量) = {cprice:.4f}（电价作用具有跨时段储能耦合，不要求简单单调符号）\n')
                f.write(f'corr(有效净负荷修正量, 购电调整量) = {cnet:.4f}\n')
                f.write(f'12:00累计向上调整 = {up_kwh:.2f} kWh\n12:00累计向下调整 = {down_kwh:.2f} kWh\n')
                f.write(f'仅0:00总费用 = {bt:.2f} 元\n0:00+12:00总费用 = {ft:.2f} 元，变化 {ft-bt:+.2f} 元\n')
                f.write( f'紧急购电变化 = {femg- bemg:+.2f} kWh\n富余电量变化 = {fsur-bsur:+.2f} kWh\n')
            logger.info('已输出预报修正与购电调整响应分析结果。')
    logger.info('完成：%s',csv_path); return results
def main():
    cfg=CFG; out=Path(cfg.out_dir); out.mkdir(parents=True,exist_ok=True)
    logger=setup_logger(out)
    logger.info('不同预报时点选择的结果影响')
    dates,price_prior,load_prior,load,pv,price_actual=load_inputs(cfg,logger)
    pv_hourly=load_attachment3(cfg,dates,logger); D=len(dates)
    n=cfg.slots_per_day; run_D=min(D,cfg.max_days) if cfg.max_days>0 else D
    cache_load,cache_price=try_load_q42_forecast_cache(cfg,D,logger)
    initial_buy=np.full((D,n),np.nan); effective_buy=np.full((D,n),np.nan)
    ch=np.zeros((D,n)); dis=np.zeros((D,n)); emg=np.zeros((D,n))
    sur=np.zeros((D,n)); soc_hist=np.full((D,n),np.nan)
    soc_start=np.full(D,np.nan); soc_end=np.full(D,np.nan)
    net_errors=np.full((D,n),np.nan); cplus=np.full((D,n),np.nan)
    cminus=np.full((D,n),np.nan)
    alpha_store=np.full((D,n),np.nan); fallback_reserve_store=np.zeros((D,n))
    load_fc_full_store=np.full((D,cfg.forecast_slots),np.nan)
    price_fc_full_store=np.full((D,cfg.forecast_slots),np.nan)
    stage_hist=np.full((D,4,cfg.rolling_horizon_slots),np.nan)
    price_stage_fc=np.full((D,4,cfg.rolling_horizon_slots),np.nan)
    plan_cost=np.zeros(D); adj_cost=np.zeros(D); emg_cost=np.zeros(D)
    up_energy=np.zeros(D); down_energy=np.zeros(D)
    load_rmse=np.zeros(D); pv_rmse0=np.zeros(D); mu0=np.zeros(D)
    lp=build_day_ahead_templates(cfg); soc=cfg.e0; stage_idx={0:0,6:1,12:2,18:3}
    hours=(0,6,12,18)
    bar=tqdm( range(run_D), desc='Q4-3全年计算', unit='day', dynamic_ncols=True,
        mininterval=.25)
    for d in bar:
        day=dates[d]; soc_start[d]=soc
        if d>=2: archive_stage_forecast_errors( d- 2, load_fc_full_store,
            pv_hourly, load, pv, stage_hist, cfg)
        if cache_load is not None: lf_full=np.asarray(cache_load[d],dtype=float)
        else: lf_full=np.asarray( load_forecast(d, load, load_prior, cfg,
            logger)[:cfg.forecast_slots], dtype=float)
        if cache_price is not None: pr_full=np.asarray( cache_price[d],
            dtype=float)
        else: pr_full=np.asarray( price_forecast(d, price_actual, price_prior,
            cfg, logger)[:cfg.forecast_slots], dtype=float)
        load_fc_full_store[d]=lf_full; price_fc_full_store[d]=pr_full
        lf0=lf_full[:n]
        pf0=pv_stage_forecast_horizon_10min(d,0,pv_hourly,pv,cfg)[:n]
        price_stage_fc[d,0]=pr_full[:n]
        update_risk_cost_cache( d, cplus, cminus, net_errors, price_actual,
            soc_hist, initial_buy, load, pv, cfg)
        alpha,_,fallback=dynamic_risk_weights(d,net_errors,cplus,cminus,cfg)
        alpha_store[d]=alpha; fallback_reserve_store[d]=fallback
        da=solve_day_ahead_price(pr_full,lf0,pf0,fallback,soc,cfg,lp)
        mu0[d]=da['mu']; initial_buy[d]=da['buy']; current_buy=da['buy'].copy()
        current_ch=da['ch'].copy(); current_dis=da['dis'].copy()
        effective_buy[d]=current_buy
        # 真实价格只用于事后结算，不进入0:00决策
        plan_cost[d]=float(np.sum(initial_buy[d]*price_actual[d])*cfg.dt)
        for si,h in enumerate(hours):
            start=h*6
            if h>0:
                lf_h=load_stage_forecast_horizon(lf_full,load[d],h,cfg)
                pf_h=pv_stage_forecast_horizon_10min(d,h,pv_hourly,pv,cfg)
                reserve_h=stage_reserve_from_history( d, stage_idx[h],
                    stage_hist, fallback, h, cfg)
                pr_h=price_stage_forecast_horizon(pr_full,price_actual[d],h,cfg)
                price_stage_fc[d,si]=pr_h
                adj=solve_rolling_mpc_24h_price( pr_h, lf_h, pf_h, reserve_h,
                    soc, current_buy, start, cfg)
                m=n-start; current_buy[start:]=adj['buy'][:m]
                current_ch[start:]=adj['ch'][:m]
                current_dis[start:]=adj['dis'][:m]
                # 真实结算
                real_today=price_actual[d,start:]
                upv=adj['up'][:m]; dnv=adj['down'][:m]
                adj_cost[ d]+=float( np.sum((1.5* real_today* upv- 0.5*
                    real_today* dnv)* cfg.dt))
                up_energy[d]+=adj['up_energy']
                down_energy[d]+=adj['down_energy']
            end=hours[si+1]*6 if si+1<len(hours) else n
            effective_buy[d,start:end]=current_buy[start:end]
            for t in range(start,end):
                soc_hist[d,t]=soc
                sol=operate_one_slot( soc, current_buy[t], load[d, t], pv[d, t],
                    current_ch[t], current_dis[t], cfg)
                ch[d,t]=sol['ch']; dis[d,t]=sol['dis']; emg[d,t]=sol['emg']
                sur[d,t]=sol['sur']; soc=sol['E']
        soc_end[d]=soc
        emg_cost[d]=float(np.sum(emg[d]*5*price_actual[d])*cfg.dt)
        net_errors[d]=(load[d]-pv[d])-(lf0-pf0)
        load_rmse[d]=float(np.sqrt(np.mean((load[d]-lf0)**2)))
        pv_rmse0[d]=float(np.sqrt(np.mean((pv[d]-pf0)**2)))
        if cfg.verbose_each_day:
            logger.info(
                '【%3d/%3d】%s | SOC %.0f→%.0f | 初始 %.2f | 调整 %.2f | 紧急 %.2f | 紧急 %.1f kWh',
                d+ 1, run_D, day, soc_start[d], soc_end[d], plan_cost[d],
                adj_cost[d], emg_cost[d], np.sum(emg[d])* cfg.dt)
        if (d+1)%30==0 or d==run_D-1:
            bar.set_postfix_str( f'{day} SOC={soc:.0f} 紧急={np.sum(emg[d])*cfg.dt:.0f}kWh')
        if cfg.checkpoint_every_days>0 and ( (d+ 1)%cfg.checkpoint_every_days==
            0 or d== run_D- 1):
            np.savez_compressed( out/ 'q4_3_checkpoint.npz', last_day=d,
                initial_buy=initial_buy[:d+ 1], effective_buy=effective_buy[:d+
                1], ch=ch[:d+ 1], dis=dis[:d+ 1], emg=emg[:d+ 1], sur=sur[:d+1],
                soc_start=soc_start[:d+ 1], soc_end=soc_end[:d+ 1],
                load_fc=load_fc_full_store[:d+ 1],
                price_fc=price_fc_full_store[:d+ 1],
                price_stage_fc=price_stage_fc[:d+ 1],
                fallback_reserve=fallback_reserve_store[:d+ 1],
                stage_net_error_hist=stage_hist[:d+ 1])
    buy_cost_after_adjust=plan_cost+adj_cost
    if run_D>=32: fill_result4_3( cfg.template, cfg.out_xlsx, dates,
        price_actual, initial_buy, effective_buy, buy_cost_after_adjust, ch,dis,
        soc_start, soc_end, emg, cfg, logger)
    save_outputs( dates, price_actual, price_fc_full_store, price_stage_fc,
        initial_buy, effective_buy, ch, dis, emg, sur, soc_start, soc_end,
        plan_cost, adj_cost, emg_cost, up_energy, down_energy, load_rmse,
        pv_rmse0, run_D, cfg, logger)
    # 表格
    if run_D>=32:
        export_paper_tables_q43( dates, price_actual, initial_buy,effective_buy,
            ch, dis, emg, soc_start, soc_end, plan_cost, adj_cost, emg_cost,
            run_D, cfg, logger)
    # 直接复用主模型已生成的全年预测/风险缓存
    if cfg.run_comparison and run_D>=32:
        run_update_time_comparison_q43( dates, price_actual, load, pv,
            load_fc_full_store, price_fc_full_store, pv_hourly,
            fallback_reserve_store, stage_hist, run_D, cfg, lp, logger )
    logger.info('='*72)
    logger.info('Q4-3完成。输出目录：%s',cfg.out_dir)
if __name__=='__main__': main()