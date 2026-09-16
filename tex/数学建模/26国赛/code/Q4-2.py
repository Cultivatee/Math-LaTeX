from __future__ import annotations
import sys
import logging
from dataclasses import dataclass
from pathlib import Path
from datetime import datetime, timedelta, date
from typing import Tuple, Dict
import numpy as np
from scipy.optimize import linprog
from scipy import sparse
from statsmodels.tsa.ar_model import AutoReg
try:
    from statsmodels.tsa.seasonal import MSTL
    HAS_MSTL = True
except Exception:
    HAS_MSTL = False
try:
    from openpyxl import load_workbook, Workbook
except ImportError as e:
    raise ImportError(
        "缺少 openpyxl。请在当前 Python 环境执行：pip install openpyxl"
    ) from e
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
plt.rcParams["font.family"] = "sans-serif"
plt.rcParams["font.sans-serif"] = ["Noto Sans CJK SC", "Microsoft YaHei", "SimHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["font.size"] = 12
plt.rcParams["axes.spines.top"] = False
plt.rcParams["axes.linewidth"] = 1.5
plt.rcParams["xtick.major.width"] = 1.5
plt.rcParams["ytick.major.width"] = 1.5
plt.rcParams["xtick.major.size"] = 6
plt.rcParams["ytick.major.size"] = 6
C_BUY, C_PV, C_DIS, C_LOAD, C_CH = "#2E74B5", "#27AE60", "#E67E22", "#34495E", "#E74C3C"
C_W, C_SOC, C_MIN, C_SOCB, C_RES = "#95A5A6", "#7F8C8D", "#2980B9", "#1F77B4", "#16A085"
C_ALPHA, C_MAX, C_PRICE = "#8E44AD", "#C0392B", "#E74C3C"
class Config:
    b_dir = str(Path(__file__).resolve().parent)
    att1 = str(Path(b_dir) / "附件1.xlsx")
    att2 = str(Path(b_dir) / "附件2.xlsx")
    att4 = str(Path(b_dir) / "附件4.xlsx")
    if not Path(att4).exists() and Path(b_dir, "附件4(2).xlsx").exists():
        att4 = str(Path(b_dir) / "附件4(2).xlsx")
    temp = str(Path(b_dir) / "附件5" / "result4-2.xlsx")
    if not Path(temp).exists() and Path(b_dir, "result4-2.xlsx").exists():
        temp = str(Path(b_dir) / "result4-2.xlsx")
    if not Path(temp).exists() and Path(b_dir, "result4-2(1).xlsx").exists():
        temp = str(Path(b_dir) / "result4-2(1).xlsx")
    o_dir = str(Path(b_dir) / "Q4-2结果")
    o_xlsx = str(Path(o_dir) / "result4-2_filled.xlsx")
    fig_dir = str(Path(o_dir) / "Q4-2图片输出")
    slots_pday = 144; dt = 1.0 / 6.0
    hori_slots = 144; fore_slots = 288
    eta_c = eta_d = 0.90
    e_min, e_max, p_max, e0 = 1200.0, 10800.0, 5000.0, 6000.0
    l_wdays = pri_wdays = 42; pv_wdays = 60
    l_mindays = pri_mindays = 21; mstl_days = 7
    ar_candi = pri_ar_candi = (1, 2, 3, 6, 12, 18, 36)
    pv_ar_candi = (1, 2, 3, 6, 12)
    kar_candi = (1, 2, 3, 5, 7)
    clear_quant = 0.93; pv_smo_win = 2
    rk_hisdays = 28; rk_decay = 0.90; rk_minsamp = 5
    alpha_cold, alpha_min, alpha_max = 4.0, 1.0, 8.0
    beta_fixed = 1.0; c_eps = 1e-4
    tau_min, tau_max = 0.55, 0.88; reserve_scale = 0.65
    eps_cyc, eps_surplus = 1e-5, 1e-8; lp_method = "highs"
    max_days = 365; checkpoint_every_days = 30
CFG = Config()
def setup_logger(out_dir: Path) -> logging.Logger:
    out_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("Q2")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    fmt = logging.Formatter("%(message)s")
    sh = logging.StreamHandler(sys.stdout)
    sh.setFormatter(fmt)
    logger.addHandler(sh)
    return logger
def excel_serial_to_date(x) -> datetime:
    return datetime(1899, 12, 30) + timedelta(days=float(x))
def cell_to_date(x):
    """兼容 openpyxl 返回的 date/datetime、Excel 序列号和常见日期字符串。"""
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    try:
        from datetime import date as _date
        if isinstance(x, _date):
            return x
    except Exception:
        pass
    if isinstance(x, (int, float, np.integer, np.floating)):
        return excel_serial_to_date(x).date()
    if isinstance(x, str):
        txt=x.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(txt, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"无法识别日期单元格：{x!r}")
def safe_float_array(x) -> np.ndarray:
    return np.asarray([[float(v) if v is not None else np.nan for v in row] for row in x], dtype=float)
def weighted_quantile(values: np.ndarray, weights: np.ndarray, q: float) -> float:
    values = np.asarray(values, dtype=float)
    weights = np.asarray(weights, dtype=float)
    m = np.isfinite(values) & np.isfinite(weights) & (weights > 0)
    if m.sum() == 0:
        return 0.0
    o=np.argsort(v:=values[m]);w=weights[m][o];v=v[o];cw=np.cumsum(w)/w.sum()
    return float(np.interp(q, cw, v))
def moving_average_circular(x: np.ndarray, half: int) -> np.ndarray:
    if half <= 0:
        return x.copy()
    n = len(x); out = np.zeros(n)
    for i in range(n):
        lo = max(0, i-half); hi = min(n, i+half+1)
        out[i] = np.mean(x[lo:hi])
    return out
def choose_ar_forecast(series: np.ndarray, horizon: int, candidates: Tuple[int, ...], fallback: float = 0.0) -> np.ndarray:
    """用AIC在少量AutoReg候选中选阶；失败时返回fallback/近期均值。"""
    y = np.asarray(series, dtype=float)
    y = y[np.isfinite(y)]
    if len(y) < 8 or np.std(y) < 1e-8:
        return np.full(horizon, float(np.mean(y) if len(y) else fallback))
    best = None
    for p in candidates:
        if len(y) <= p + 5:
            continue
        try:
            model = AutoReg(y, lags=p, trend="ct", old_names=False).fit()
            if np.isfinite(model.aic) and (best is None or model.aic < best[0]):
                best = (model.aic, model)
        except Exception:
            continue
    if best is None:
        return np.full(horizon, float(np.mean(y[-min(7, len(y)):])) )
    try:
        pred = np.asarray(best[1].predict(start=len(y), end=len(y)+horizon-1, dynamic=False), dtype=float)
        if np.any(~np.isfinite(pred)):
            raise ValueError("AR forecast nonfinite")
        return pred
    except Exception:
        return np.full(horizon, float(np.mean(y[-min(7, len(y)):])) )
def load_inputs(cfg: Config, logger):
    logger.info("【数据读取】正在读取附件1、附件2、附件4……")
    wb1 = load_workbook(cfg.att1, data_only=True, read_only=True)
    ws1 = wb1["Sheet1"] if "Sheet1" in wb1.sheetnames else wb1[wb1.sheetnames[0]]
    rows1 = list(ws1.iter_rows(min_row=2, max_row=145, min_col=2, max_col=4, values_only=True))
    wb1.close()
    price_prior = np.asarray([float(r[0]) for r in rows1], dtype=float)
    load_prior = np.asarray([float(r[1]) for r in rows1], dtype=float)
    pv_prior = np.asarray([float(r[2]) for r in rows1], dtype=float)
    wb2 = load_workbook(cfg.att2, data_only=True, read_only=True)
    sl = wb2["小区负载"]
    sp = wb2["光伏发电实际功率"]
    load_all = list(sl.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    pv_all = list(sp.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    wb2.close()
    dates = [cell_to_date(r[0]) for r in load_all]
    load = safe_float_array([r[1:] for r in load_all])
    pv_actual = safe_float_array([r[1:] for r in pv_all])
    wb4 = load_workbook(cfg.att4, data_only=True, read_only=True)
    ws4 = wb4["Sheet1"] if "Sheet1" in wb4.sheetnames else wb4[wb4.sheetnames[0]]
    price_rows = list(ws4.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    wb4.close()
    price_dates = [cell_to_date(r[0]) for r in price_rows]
    price_actual = safe_float_array([r[1:] for r in price_rows])
    if load.shape != (365,144) or pv_actual.shape != (365,144) or price_actual.shape != (365,144):
        raise ValueError(f"数据维度异常：load={load.shape}, pv={pv_actual.shape}, price={price_actual.shape}")
    if dates != price_dates:
        raise ValueError("附件2与附件4日期未完全对齐。")
    if np.any(~np.isfinite(price_actual)) or np.any(price_actual < 0):
        raise ValueError("附件4存在缺失或负电价；当前模型按非负实时电价处理。")
    logger.info("【数据读取】完成：365天×144个10分钟时段。")
    logger.info("电价范围 %.4f~%.4f 元/kWh，全年均价 %.4f 元/kWh。",
                float(np.min(price_actual)), float(np.max(price_actual)), float(np.mean(price_actual)))
    return dates, price_prior, load_prior, pv_prior, load, pv_actual, price_actual
_LOAD_MSTL_CACHE = {}
def load_forecast(day_idx: int, load: np.ndarray, prior: np.ndarray, cfg: Config, logger=None) -> np.ndarray:
    H = cfg.fore_slots
    if day_idx == 0:
        return np.tile(prior, 2)
    hist_days = min(day_idx, cfg.l_wdays)
    hist_mat = load[day_idx-hist_days:day_idx]
    # 冷启动：同一时刻近期日加权 + 附件1先验
    if hist_days < cfg.l_mindays or not HAS_MSTL:
        k = hist_mat.shape[0]
        weights = np.exp(-0.10*np.arange(k-1, -1, -1))
        weights /= weights.sum()
        shape = np.sum(hist_mat * weights[:, None], axis=0)
        if hist_days >= 7:
            shape = 0.55*shape + 0.45*load[day_idx-7]
        w_prior = max(0.0, 1.0 - hist_days / cfg.l_mindays)
        shape = (1-w_prior)*shape + w_prior*prior
        return np.tile(shape, 2)
    # 仅每隔若干天完整重估一次 MSTL；数据已确认无明显异常，因此关闭 robust 迭代
    need_refit = (
        not _LOAD_MSTL_CACHE
        or day_idx - _LOAD_MSTL_CACHE.get("refit_day", -10**9) >= cfg.mstl_days
    )
    if need_refit:
        y = hist_mat.reshape(-1)
        try:
            mstl = MSTL(y, periods=(144, 1008), stl_kwargs={"robust": False}).fit()
            trend = np.asarray(mstl.trend)
            seasonal = np.asarray(mstl.seasonal)
            if seasonal.ndim == 1:
                s_daily = seasonal; s_weekly = np.zeros_like(seasonal)
            else:
                s_daily, s_weekly = seasonal[:, 0], seasonal[:, 1]
            resid = np.asarray(mstl.resid)
            # 保存稳定的日周期与周周期结构，供未来几天复用
            daily_tail = s_daily[-min(len(s_daily), 28*144):]
            weekly_tail = s_weekly[-min(len(s_weekly), 6*1008):]
            daily_pattern = np.array([
                np.nanmean(daily_tail[i::144]) for i in range(144)
            ])
            weekly_pattern = np.array([
                np.nanmean(weekly_tail[i::1008]) for i in range(1008)
            ])
            _LOAD_MSTL_CACHE.clear()
            _LOAD_MSTL_CACHE.update({
                "refit_day": day_idx, "origin_day": day_idx - hist_days,
                "daily_pattern": daily_pattern, "weekly_pattern": weekly_pattern,
                "trend_last": float(trend[-1]),
                "trend_slope": float(np.clip(
                    np.polyfit(
                        np.arange(min(432, len(trend)), dtype=float),
                        trend[-min(432, len(trend)):],
                        1
                    )[0], -5.0, 5.0
                )),
                "resid_tail": resid[-min(len(resid), 14*144):].copy(),
            })
            if logger:
                logger.info("负荷预测：第 %d 天重新拟合 MSTL（日/周周期）。", day_idx + 1)
        except Exception as e:
            if logger:
                logger.warning("MSTL失败(day=%d)，回退季节加权预测：%s", day_idx, e)
            k = hist_mat.shape[0]
            weights = np.exp(-0.10*np.arange(k-1, -1, -1))
            weights /= weights.sum()
            shape = np.sum(hist_mat * weights[:, None], axis=0)
            if hist_days >= 7:
                shape = 0.5*shape + 0.5*load[day_idx-7]
            return np.tile(shape, 2)
    # 使用最近一次 MSTL 的季节结构，并用最新 3 天数据快速校正趋势和残差
    daily_pattern = _LOAD_MSTL_CACHE["daily_pattern"]
    weekly_pattern = _LOAD_MSTL_CACHE["weekly_pattern"]
    origin_day = _LOAD_MSTL_CACHE["origin_day"]
    recent_days = min(hist_days, 3)
    recent = load[day_idx-recent_days:day_idx].reshape(-1)
    global_start = (day_idx - recent_days) * 144
    gidx = global_start + np.arange(len(recent))
    sd_hist = daily_pattern[gidx % 144]
    sw_hist = weekly_pattern[(gidx - origin_day*144) % 1008]
    deseason = recent - sd_hist - sw_hist
    m = min(432, len(deseason))
    x = np.arange(m, dtype=float)
    coef = np.polyfit(x, deseason[-m:], 1)
    slope = float(np.clip(coef[0], -5.0, 5.0))
    level = float(deseason[-1])
    trend_future = level + slope*np.arange(1, H+1)
    future_gidx = day_idx*144 + np.arange(H)
    sd_future = daily_pattern[future_gidx % 144]
    sw_future = weekly_pattern[(future_gidx - origin_day*144) % 1008]
    # 残差 AR 仍每天轻量更新，不需要重新做 MSTL
    hist_axis = np.arange(len(deseason), dtype=float)
    fitted_trend_hist = level + slope*(hist_axis - (len(deseason)-1))
    resid_recent = deseason - fitted_trend_hist
    ar_future = choose_ar_forecast(
        resid_recent[-min(len(resid_recent), 3*144):],
        H,
        cfg.ar_candi,
        0.0,
    )
    pred = trend_future + sd_future + sw_future + ar_future
    return np.maximum(pred, 0.0)
_PRICE_MSTL_CACHE = {}
def price_forecast(day_idx: int, price_actual: np.ndarray, prior: np.ndarray, cfg: Config, logger=None) -> np.ndarray:
    H = cfg.fore_slots
    if day_idx == 0:
        return np.tile(prior, 2)
    hist_days = min(day_idx, cfg.pri_wdays)
    hist_mat = price_actual[day_idx-hist_days:day_idx]
    # 冷启动：同一时刻近期日加权 + 附件1先验
    if hist_days < cfg.pri_mindays or not HAS_MSTL:
        k = hist_mat.shape[0]
        weights = np.exp(-0.10*np.arange(k-1, -1, -1))
        weights /= weights.sum()
        shape = np.sum(hist_mat * weights[:, None], axis=0)
        if hist_days >= 7:
            shape = 0.55*shape + 0.45*price_actual[day_idx-7]
        w_prior = max(0.0, 1.0 - hist_days / cfg.pri_mindays)
        shape = (1-w_prior)*shape + w_prior*prior
        return np.tile(shape, 2)
    # 仅每隔若干天完整重估一次 MSTL；数据已确认无明显异常，因此关闭 robust 迭代
    need_refit = (
        not _PRICE_MSTL_CACHE
        or day_idx - _PRICE_MSTL_CACHE.get("refit_day", -10**9) >= cfg.mstl_days
    )
    if need_refit:
        y = hist_mat.reshape(-1)
        try:
            mstl = MSTL(y, periods=(144, 1008), stl_kwargs={"robust": False}).fit()
            trend = np.asarray(mstl.trend)
            seasonal = np.asarray(mstl.seasonal)
            if seasonal.ndim == 1:
                s_daily = seasonal; s_weekly = np.zeros_like(seasonal)
            else:
                s_daily, s_weekly = seasonal[:, 0], seasonal[:, 1]
            resid = np.asarray(mstl.resid)
            # 保存稳定的日周期与周周期结构，供未来几天复用
            daily_tail = s_daily[-min(len(s_daily), 28*144):]
            weekly_tail = s_weekly[-min(len(s_weekly), 6*1008):]
            daily_pattern = np.array([
                np.nanmean(daily_tail[i::144]) for i in range(144)
            ])
            weekly_pattern = np.array([
                np.nanmean(weekly_tail[i::1008]) for i in range(1008)
            ])
            _PRICE_MSTL_CACHE.clear()
            _PRICE_MSTL_CACHE.update({
                "refit_day": day_idx, "origin_day": day_idx - hist_days,
                "daily_pattern": daily_pattern, "weekly_pattern": weekly_pattern,
                "trend_last": float(trend[-1]),
                "trend_slope": float(np.clip(
                    np.polyfit(
                        np.arange(min(432, len(trend)), dtype=float),
                        trend[-min(432, len(trend)):],
                        1
                    )[0], -5.0, 5.0
                )),
                "resid_tail": resid[-min(len(resid), 14*144):].copy(),
            })
            if logger:
                logger.info("电价预测：第 %d 天重新拟合 MSTL（日/周周期）。", day_idx + 1)
        except Exception as e:
            if logger:
                logger.warning("MSTL失败(day=%d)，回退季节加权预测：%s", day_idx, e)
            k = hist_mat.shape[0]
            weights = np.exp(-0.10*np.arange(k-1, -1, -1))
            weights /= weights.sum()
            shape = np.sum(hist_mat * weights[:, None], axis=0)
            if hist_days >= 7:
                shape = 0.5*shape + 0.5*price_actual[day_idx-7]
            return np.tile(shape, 2)
    # 使用最近一次 MSTL 的季节结构，并用最新 3 天数据快速校正趋势和残差
    daily_pattern = _PRICE_MSTL_CACHE["daily_pattern"]
    weekly_pattern = _PRICE_MSTL_CACHE["weekly_pattern"]
    origin_day = _PRICE_MSTL_CACHE["origin_day"]
    recent_days = min(hist_days, 3)
    recent = price_actual[day_idx-recent_days:day_idx].reshape(-1)
    global_start = (day_idx - recent_days) * 144
    gidx = global_start + np.arange(len(recent))
    sd_hist = daily_pattern[gidx % 144]
    sw_hist = weekly_pattern[(gidx - origin_day*144) % 1008]
    deseason = recent - sd_hist - sw_hist
    m = min(432, len(deseason))
    x = np.arange(m, dtype=float)
    coef = np.polyfit(x, deseason[-m:], 1)
    slope = float(np.clip(coef[0], -5.0, 5.0))
    level = float(deseason[-1])
    trend_future = level + slope*np.arange(1, H+1)
    future_gidx = day_idx*144 + np.arange(H)
    sd_future = daily_pattern[future_gidx % 144]
    sw_future = weekly_pattern[(future_gidx - origin_day*144) % 1008]
    # 残差 AR 仍每天轻量更新，不需要重新做 MSTL
    hist_axis = np.arange(len(deseason), dtype=float)
    fitted_trend_hist = level + slope*(hist_axis - (len(deseason)-1))
    resid_recent = deseason - fitted_trend_hist
    ar_future = choose_ar_forecast(resid_recent[-min(len(resid_recent), 3*144):],
        H,cfg.pri_ar_candi,0.0,)
    pred = trend_future + sd_future + sw_future + ar_future
    return np.maximum(pred, 0.0)
def build_clear_envelope(day_idx: int, pv: np.ndarray, prior: np.ndarray, cfg: Config) -> np.ndarray:
    if day_idx == 0:
        base = prior.copy()
        return np.maximum(base, 0.0)
    w = min(day_idx, cfg.pv_wdays)
    hist = pv[day_idx-w:day_idx]
    q = np.quantile(hist, cfg.clear_quant, axis=0)
    q = moving_average_circular(q, cfg.pv_smo_win)
    # 样本短时与附件1先验混合
    prior_w = max(0.0, 1.0 - w/21.0)
    base = (1-prior_w)*q + prior_w*prior
    # 夜间稳定置0：历史90%以上为0即置0
    zero_ratio = np.mean(hist <= 1e-8, axis=0)
    base[zero_ratio > 0.90] = 0.0
    return np.maximum(base, 0.0)
def pv_forecast(day_idx: int, pv: np.ndarray, prior: np.ndarray, cfg: Config, logger=None) -> Tuple[np.ndarray, Dict]:
    H = cfg.fore_slots
    if day_idx == 0:
        return np.tile(prior, 2), {"kappa_d": 1.0, "kappa_next": 1.0, "envelope": prior.copy()}
    B_d = build_clear_envelope(day_idx, pv, prior, cfg)
    # 下一日包络：年周期变化慢，先用当前包络；用近期日能量变化率做轻微尺度外推
    recent_energy = np.sum(pv[max(0,day_idx-14):day_idx], axis=1)*cfg.dt
    scale_next = 1.0
    if len(recent_energy) >= 7:
        x = np.arange(len(recent_energy))
        slope = np.polyfit(x, recent_energy, 1)[0]
        denom = max(np.mean(recent_energy), 1.0)
        scale_next = float(np.clip(1.0 + slope/denom, 0.95, 1.05))
    B_next = B_d * scale_next
    # 用当前滚动包络近似反演近期日透射率；只使用过去日期
    w = min(day_idx, cfg.pv_wdays)
    hist = pv[day_idx-w:day_idx]
    denom = max(np.sum(B_d)*cfg.dt, 1e-6)
    kappas = np.clip(np.sum(hist, axis=1)*cfg.dt/denom, 0.0, 1.0)
    kpred = choose_ar_forecast(kappas, 2, cfg.kar_candi, fallback=float(np.mean(kappas)))
    kpred = np.clip(kpred, 0.0, 1.0)
    base_d = B_d*kpred[0]
    base_next = B_next*kpred[1]
    # 残差构造：每个历史日用其日能量对应kappa相对于当前包络重构
    residual_days = hist - kappas[:, None]*B_d[None, :]
    resid_series = residual_days.reshape(-1)
    rpred = choose_ar_forecast(resid_series[-min(len(resid_series), 21*144):], H, cfg.pv_ar_candi, 0.0)
    pred = np.concatenate([base_d, base_next]) + rpred
    pred = np.maximum(pred, 0.0)
    # 夜间稳定无发电时段置零
    zero_mask = B_d <= 1e-8
    for h in range(H):
        if zero_mask[h % 144]:
            pred[h] = 0.0
    return pred, {"kappa_d": float(kpred[0]), "kappa_next": float(kpred[1]), "envelope": B_d}
def compute_historical_marginal_cost_day(j: int,
                                         available_day_exclusive: int,
                                         net_errors: np.ndarray,
                                         price_actual: np.ndarray,
                                         soc_hist: np.ndarray,
                                         planned_hist: np.ndarray,
                                         load: np.ndarray,
                                         pv: np.ndarray,
                                         cfg: Config) -> Tuple[np.ndarray, np.ndarray]:
    nT = cfg.slots_pday
    cplus = np.full(nT, np.nan, dtype=float)
    cminus = np.full(nT, np.nan, dtype=float)
    if j < 0 or j >= available_day_exclusive:
        return cplus, cminus
    # 只把当前允许访问的历史展平；后续窗口均是连续切片
    n_avail = available_day_exclusive * nT
    plan_f = planned_hist[:available_day_exclusive].reshape(-1)
    load_f = load[:available_day_exclusive].reshape(-1)
    pv_f = pv[:available_day_exclusive].reshape(-1)
    soc_f = soc_hist[:available_day_exclusive].reshape(-1)
    err_f = net_errors[:available_day_exclusive].reshape(-1)
    slot_price = price_actual[:available_day_exclusive].reshape(-1)
    base = j * nT
    eff_rt = cfg.eta_c * cfg.eta_d
    for t in range(nT):
        e = net_errors[j, t]
        if not np.isfinite(e):
            continue
        lam_t = price_actual[j, t]
        s0 = soc_hist[j, t] if np.isfinite(soc_hist[j, t]) else cfg.e0
        g0 = base + t
        g1 = g0 + 1
        g2 = min(g0 + cfg.hori_slots, n_avail - 1)
        if e > 0:
            c_em = 5.0 * lam_t
            c_bat = np.inf
            if s0 > cfg.e_min + 1e-6 and g1 <= g2:
                sl = slice(g1, g2 + 1)
                margin = plan_f[sl] + pv_f[sl] - load_f[sl]
                sf = soc_f[sl]
                pf = slot_price[sl]
                soc_ok = (~np.isfinite(sf)) | (sf < cfg.e_max - 1e-6)
                feasible = soc_ok & ((margin > 0.0) | (pf < lam_t))
                if np.any(feasible):
                    c_bat = float(np.min(pf[feasible]) / eff_rt)
            cplus[t] = min(c_em, c_bat) if np.isfinite(c_bat) else c_em
        elif e < 0:
            recovery = 0.0
            if s0 < cfg.e_max - 1e-6 and g1 <= g2:
                sl = slice(g1, g2 + 1)
                ef = err_f[sl]
                pf = slot_price[sl]
                m = np.isfinite(ef) & (ef > 0.0)
                if np.any(m):
                    recovery = float(np.max(5.0 * pf[m] * eff_rt))
            cminus[t] = max(0.0, lam_t - recovery)
    return cplus, cminus
def update_risk_cost_cache(day_idx: int,
                           cplus_cache: np.ndarray,
                           cminus_cache: np.ndarray,
                           net_errors: np.ndarray,
                           price_actual: np.ndarray,
                           soc_hist: np.ndarray,
                           planned_hist: np.ndarray,
                           load: np.ndarray,
                           pv: np.ndarray,
                           cfg: Config) -> None:
    if day_idx <= 0:
        return
    # 昨天：只能使用到昨天24:00，保持严格无未来信息泄露
    j = day_idx - 1
    cp, cm = compute_historical_marginal_cost_day(
        j, day_idx, net_errors, price_actual, soc_hist, planned_hist, load, pv, cfg)
    cplus_cache[j] = cp
    cminus_cache[j] = cm
    if day_idx >= 2:
        j = day_idx - 2
        cp, cm = compute_historical_marginal_cost_day(
            j, day_idx, net_errors, price_actual, soc_hist, planned_hist, load, pv, cfg)
        cplus_cache[j] = cp
        cminus_cache[j] = cm
def dynamic_risk_weights(day_idx: int,
                         net_errors: np.ndarray,
                         cplus_cache: np.ndarray,
                         cminus_cache: np.ndarray,
                         cfg: Config) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    nT = cfg.slots_pday
    alpha = np.full(nT, cfg.alpha_cold, dtype=float)
    beta = np.full(nT, cfg.beta_fixed, dtype=float)
    reserve = np.zeros(nT, dtype=float)
    if day_idx <= 0:
        return alpha, beta, reserve
    j0 = max(0, day_idx - cfg.rk_hisdays)
    hist_days = np.arange(j0, day_idx)
    if hist_days.size == 0:
        return alpha, beta, reserve
    ages = day_idx - hist_days
    base_w = cfg.rk_decay ** (ages - 1)
    errs = net_errors[hist_days]          # (HIST, 144)
    cp_all = cplus_cache[hist_days]
    cm_all = cminus_cache[hist_days]
    for t in range(nT):
        et = errs[:, t]
        # 正向误差历史
        mp = (et > 0) & np.isfinite(et) & np.isfinite(cp_all[:, t])
        # 负向误差历史
        mn = (et < 0) & np.isfinite(et) & np.isfinite(cm_all[:, t])
        if np.count_nonzero(mp) >= cfg.rk_minsamp and np.count_nonzero(mn) >= cfg.rk_minsamp:
            cp = float(np.average(cp_all[mp, t], weights=base_w[mp]))
            cm = float(np.average(cm_all[mn, t], weights=base_w[mn]))
            raw = cp / max(cm, cfg.c_eps)
            alpha[t] = min(cfg.alpha_max, max(cfg.alpha_min, raw))
        tau = alpha[t] / (alpha[t] + beta[t])
        tau = min(cfg.tau_max, max(cfg.tau_min, tau))
        if np.count_nonzero(mp) >= cfg.rk_minsamp:
            reserve[t] = cfg.reserve_scale * weighted_quantile(et[mp], base_w[mp], tau)
    return alpha, beta, reserve
@dataclass
class LPStructure:
    Aeq: sparse.csr_matrix
    bounds: list
def build_lp_structure(cfg: Config) -> LPStructure:
    """全年只构造一次等式矩阵；每天仅更新价格目标与净负荷右端项。"""
    n = cfg.slots_pday
    N = 5*n
    ib, ic, idis, isu, ie = 0, n, 2*n, 3*n, 4*n
    rows, cols, data = [], [], []
    r = 0
    for t in range(n):
        for col,val in ((ib+t,1.0),(idis+t,1.0),(ic+t,-1.0),(isu+t,-1.0)):
            rows.append(r); cols.append(col); data.append(val)
        r += 1
    for t in range(n):
        rows.append(r); cols.append(ie+t); data.append(1.0)
        if t>0:
            rows.append(r); cols.append(ie+t-1); data.append(-1.0)
        rows.append(r); cols.append(ic+t); data.append(-cfg.eta_c*cfg.dt)
        rows.append(r); cols.append(idis+t); data.append(cfg.dt/cfg.eta_d)
        r += 1
    Aeq = sparse.csr_matrix((data,(rows,cols)), shape=(r,N))
    bounds = ([(0,None)]*n + [(0,cfg.p_max)]*n + [(0,cfg.p_max)]*n +
              [(0,None)]*n + [(cfg.e_min,cfg.e_max)]*n)
    return LPStructure(Aeq,bounds)
def solve_day_ahead(price_fc_48: np.ndarray, load_fc: np.ndarray,
                    pv_fc: np.ndarray, reserve: np.ndarray, soc0: float,
                    cfg: Config, lp: LPStructure) -> Dict[str,np.ndarray]:
    """每天0:00求解一次日前LP；当日价格用于目标，次日预测最低价定义日末SOC机会价值。"""
    n=cfg.slots_pday
    net_req = load_fc[:n] - pv_fc[:n] + reserve
    b=np.empty(2*n,dtype=float); b[:n]=net_req; b[n:]=0.0; b[n]=soc0
    price_today=np.asarray(price_fc_48[:n],dtype=float)
    if len(price_fc_48)>=2*n:
        next_day=np.asarray(price_fc_48[n:2*n],dtype=float)
    else:
        next_day=price_today
    mu_terminal=cfg.eta_d*float(np.min(next_day))
    c=np.zeros(5*n,dtype=float)
    c[:n]=price_today*cfg.dt; c[n:3*n]=cfg.eps_cyc*cfg.dt
    c[3*n:4*n]=cfg.eps_surplus*cfg.dt; c[5*n-1]=-mu_terminal
    res=linprog(c,A_eq=lp.Aeq,b_eq=b,bounds=lp.bounds,method=cfg.lp_method)
    if not res.success:
        raise RuntimeError(f"日前LP失败: {res.message}")
    x=res.x
    return {"buy":x[:n],"ch":x[n:2*n],"dis":x[2*n:3*n],"sur":x[3*n:4*n],
            "E":x[4*n:5*n],"obj":float(res.fun+mu_terminal*soc0),"mu_terminal":mu_terminal}
def operate_one_slot(soc0: float, buy: float, load_actual: float, pv_actual: float,
                     plan_ch: float, plan_dis: float, cfg: Config) -> Dict[str,float]:
    dt = cfg.dt
    # 先把计划动作投影到当前SOC可行范围
    ch = max(0.0, min(float(plan_ch), cfg.p_max))
    dis = max(0.0, min(float(plan_dis), cfg.p_max))
    max_ch_soc = max(0.0, (cfg.e_max - soc0) / (cfg.eta_c * dt))
    max_dis_soc = max(0.0, (soc0 - cfg.e_min) * cfg.eta_d / dt)
    ch = min(ch, max_ch_soc)
    dis = min(dis, max_dis_soc)
    # 基于真实负荷/光伏检查功率平衡
    gap = float(load_actual - (buy + pv_actual + dis - ch))
    emg = sur = 0.0
    if gap > 1e-10:
        cut_ch = min(ch, gap); ch -= cut_ch; gap -= cut_ch
        max_dis_total = min(cfg.p_max, max(0.0, (soc0 - cfg.e_min) * cfg.eta_d / dt))
        extra_dis = min(gap, max(0.0, max_dis_total - dis))
        dis += extra_dis; gap -= extra_dis
        emg = max(0.0, gap)
    elif gap < -1e-10:
        excess = -gap
        # 1. 先取消不必要的计划放电
        cut_dis = min(dis, excess); dis -= cut_dis; excess -= cut_dis
        # 2. 再增加充电
        max_ch_total = min(cfg.p_max, max(0.0, (cfg.e_max - soc0) / (cfg.eta_c * dt)))
        extra_ch = min(excess, max(0.0, max_ch_total - ch))
        ch += extra_ch; excess -= extra_ch
        # 3. 电池无法继续吸收的部分作为富余功率
        sur = max(0.0, excess)
    soc1 = soc0 + (cfg.eta_c*ch - dis/cfg.eta_d)*dt
    soc1 = min(cfg.e_max, max(cfg.e_min, soc1))
    return {"ch":ch, "dis":dis, "emg":emg, "sur":sur, "E":soc1}
def fill_result4_2(template_path:str, out_path:str, dates, price_actual,
                 plan_buy, ch, dis, soc_start, soc_end, emg, cfg:Config, logger):
    from copy import copy
    wb = load_workbook(template_path)
    sp, sc, se = wb["计划购电量"], wb["充放电量"], wb["紧急购电量"]
    date_to_idx={d:i for i,d in enumerate(dates)}
    first_submit = date(2025, 2, 1)
    computed_dates=[d for d in dates if d >= first_submit and d in date_to_idx
                    and np.isfinite(plan_buy[date_to_idx[d]]).all()]
    # 计划购电量：模板本身已经逐日展开
    for r in range(2, sp.max_row + 1):
        raw_date = sp.cell(r,1).value
        if raw_date is None:
            continue
        try:
            d=cell_to_date(raw_date)
        except Exception:
            continue
        if d not in date_to_idx:
            continue
        i=date_to_idx[d]
        if not np.isfinite(plan_buy[i]).all():
            continue
        vals=(plan_buy[i]*cfg.dt).tolist()  # kWh/10min
        for c, v in enumerate(vals, start=2):
            sp.cell(r,c).value = None if not np.isfinite(v) else float(v)
        sp.cell(r,146).value=float(np.nansum(plan_buy[i])*cfg.dt)            # EP 全天购电量
        sp.cell(r,147).value=float(np.nansum(plan_buy[i]*price_actual[i])*cfg.dt)      # EQ 全天购电费
    charge_proto=[]
    for rr in range(2,8):
        charge_proto.append({
            "height": sc.row_dimensions[rr].height,
            "cells": [sc.cell(rr,cc) for cc in range(1,7)]
        })
    emg_proto=[]
    for rr in range(2,5):
        emg_proto.append({"height": se.row_dimensions[rr].height,
            "cells": [se.cell(rr,cc) for cc in range(1,4)]})
    def copy_cell_style(src, dst):
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.font = copy(src.font); dst.fill = copy(src.fill)
        dst.border = copy(src.border); dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
    # 充放电量：彻底展开 2/1--已计算末日
    if sc.max_row >= 2:
        sc.delete_rows(2, sc.max_row-1)
    intervals=[(0,4,"0:00-4:00"),(4,8,"4:00-8:00"),(8,12,"8:00-12:00"),
               (12,16,"12:00-16:00"),(16,20,"16:00-20:00"),(20,24,"20:00-24:00")]
    r=2
    for d in computed_dates:
        i=date_to_idx[d]
        for k,(h0,h1,label) in enumerate(intervals):
            proto=charge_proto[k]
            sc.row_dimensions[r].height = proto["height"]
            for cc in range(1,7):
                copy_cell_style(proto["cells"][cc-1], sc.cell(r,cc))
            sc.cell(r,1).value = d if k==0 else None
            sc.cell(r,2).value = label
            s0=h0*6; s1=h1*6
            sc.cell(r,3).value=float(np.nansum(ch[i,s0:s1])*cfg.dt)
            sc.cell(r,4).value=float(np.nansum(dis[i,s0:s1])*cfg.dt)
            if k==0:
                sc.cell(r,5).value="0:00"
                sc.cell(r,6).value=float(soc_start[i])
            elif k==1:
                sc.cell(r,5).value="24:00"
                sc.cell(r,6).value=float(soc_end[i])
            else:
                sc.cell(r,5).value=None
                sc.cell(r,6).value=None
            r += 1
    # 紧急购电：逐日展开，不保留省略号
    def fmt_slot(k):
        mins=k*10
        if mins == 1440:
            return "24:00"
        hh=(mins//60)%24; mm=mins%60
        return f"{hh}:{mm:02d}"
    if se.max_row >= 2:
        se.delete_rows(2, se.max_row-1)
    r=2
    for d in computed_dates:
        i=date_to_idx[d]
        events=[]; t=0
        while t<144:
            if emg[i,t] > 1e-6:
                st=t; energy=0.0
                while t<144 and emg[i,t] > 1e-6:
                    energy += emg[i,t]*cfg.dt
                    t += 1
                events.append((f"{fmt_slot(st)}-{fmt_slot(t)}", float(energy)))
            else:
                t += 1
        if not events:
            events=[(None,None)]
        for k,(period,energy) in enumerate(events):
            proto=emg_proto[min(k, len(emg_proto)-1)]
            se.row_dimensions[r].height = proto["height"]
            for cc in range(1,4):
                copy_cell_style(proto["cells"][cc-1], se.cell(r,cc))
            se.cell(r,1).value = d if k==0 else None
            se.cell(r,2).value = period; se.cell(r,3).value = energy
            r += 1
    # 日期与数值格式统一
    sc.column_dimensions['A'].width = max(sc.column_dimensions['A'].width or 10, 12)
    sc.column_dimensions['B'].width = max(sc.column_dimensions['B'].width or 12, 15)
    se.column_dimensions['A'].width = max(se.column_dimensions['A'].width or 10, 12)
    se.column_dimensions['B'].width = max(se.column_dimensions['B'].width or 12, 18)
    for ws in (sc,se):
        for rr in range(2, ws.max_row+1):
            if ws.cell(rr,1).value is not None:
                ws.cell(rr,1).number_format='yyyy/m/d'
    for rr in range(2, sc.max_row+1):
        sc.cell(rr,3).number_format='0.00'
        sc.cell(rr,4).number_format='0.00'
        sc.cell(rr,6).number_format='0.00'
    for rr in range(2, se.max_row+1):
        se.cell(rr,3).number_format='0.00'
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    wb.close()
    logger.info("result4-2 已完整展开并保存：%s", out_path)
def write_specified_date_results(dates, price_actual, plan_buy, ch, dis, emg,
                                 soc_start, soc_end, cfg:Config, logger):
    targets=[date(2025,3,20), date(2025,6,21), date(2025,9,23), date(2025,12,21)]
    date_to_idx={d:i for i,d in enumerate(dates)}
    out_dir=Path(cfg.o_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    # 表1：指定六个10min时段 + 全天购电量/购电费
    slot_hours=[10,12,14,16,18,20]
    table1=[]
    for d in targets:
        if d not in date_to_idx:
            continue
        i=date_to_idx[d]
        if not np.isfinite(plan_buy[i]).all():
            continue
        vals=[]
        for h in slot_hours:
            t=h*6
            vals.append(float(plan_buy[i,t]*cfg.dt))
        total_energy=float(np.nansum(plan_buy[i])*cfg.dt)
        total_cost=float(np.nansum(plan_buy[i]*price_actual[i])*cfg.dt)
        table1.append([str(d),*vals,total_energy,total_cost])
    hdr1=['日期','10:00-10:10/kWh','12:00-12:10/kWh','14:00-14:10/kWh',
          '16:00-16:10/kWh','18:00-18:10/kWh','20:00-20:10/kWh',
          '全天购电量/kWh','全天购电费/元']
    # 表2：每个指定日期6个四小时区间 + 日初/日末SOC
    intervals=[(0,4,'0:00-4:00'),(4,8,'4:00-8:00'),(8,12,'8:00-12:00'),
               (12,16,'12:00-16:00'),(16,20,'16:00-20:00'),(20,24,'20:00-24:00')]
    table2=[]
    for d in targets:
        if d not in date_to_idx:
            continue
        i=date_to_idx[d]
        if not np.isfinite(plan_buy[i]).all():
            continue
        for k,(h0,h1,label) in enumerate(intervals):
            s0=h0*6; s1=h1*6
            table2.append([str(d) if k==0 else '', label,
                float(np.nansum(ch[i,s0:s1])*cfg.dt),
                float(np.nansum(dis[i,s0:s1])*cfg.dt),
                float(soc_start[i]) if k==0 else '',
                float(soc_end[i]) if k==0 else ''])
    hdr2=['日期','时间段','充电量/kWh','放电量/kWh','0:00储电量/kWh','24:00储电量/kWh']
    # 表3：紧急购电事件，连续10min时段自动合并
    def fmt_slot(k):
        mins=k*10
        if mins==1440:
            return '24:00'
        return f'{mins//60}:{mins%60:02d}'
    table3=[]
    for d in targets:
        if d not in date_to_idx:
            continue
        i=date_to_idx[d]
        if not np.isfinite(plan_buy[i]).all():
            continue
        events=[]; t=0
        while t<144:
            if emg[i,t] > 1e-6:
                st=t; energy=0.0
                while t<144 and emg[i,t] > 1e-6:
                    energy += emg[i,t]*cfg.dt
                    t += 1
                events.append((f'{fmt_slot(st)}-{fmt_slot(t)}', float(energy)))
            else:
                t += 1
        if not events:
            table3.append([str(d),'无紧急购电',0.0])
        else:
            for k,(period,energy) in enumerate(events):
                table3.append([str(d) if k==0 else '',period,energy])
    hdr3=['日期','紧急购电时间段','紧急购电量/kWh']
    # 三个表合并写入一个Excel（各占一个sheet）
    xlsx=out_dir/'Q4-2_指定日期数据.xlsx'
    book=Workbook(); book.remove(book.active)
    for name,hdr,rows in (('表1 计划购电',hdr1,table1),
                          ('表2 充放电',hdr2,table2),
                          ('表3 紧急购电',hdr3,table3)):
        ws=book.create_sheet(name)
        ws.append(hdr)
        for row in rows:
            ws.append(row)
    try:
        book.save(xlsx)
    except PermissionError:
        book.save(str(xlsx).replace('.xlsx','_new.xlsx'))  # 文件被占用时改存副本
    book.close()
    logger.info("指定日期表1/表2/表3 已合并保存：%s", xlsx)
def _style_axes(ax):
    """统一论文图风格：减少边框、轻虚线网格、字号一致。"""
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.grid(True, axis='y', linestyle='--', linewidth=0.7, alpha=0.3)
    ax.tick_params(labelsize=TICK_FS)
NO_TITLE, AXIS_FS, TICK_FS, LEGEND_FS = True, 15, 12, 17
def _save_line_fig(x, ys, labels, title, xlabel, ylabel, path, date_axis=False,
                   hlines=None, ylim=None, colors=None, line_styles=None,
                   xticks=None, xlim=None, linewidths=None, legend_loc=None):
    fig, ax = plt.subplots(figsize=(10.8, 5.4), dpi=300)
    for k, (y, lab) in enumerate(zip(ys, labels)):
        c = colors[k] if colors else None
        ls = line_styles[k] if line_styles else '-'
        lw = linewidths[k] if linewidths else 2.2
        ax.plot(x, y, linewidth=lw, label=lab, color=c, linestyle=ls)
    if hlines:
        for item in hlines:
            yv, lab = item[0], item[1]
            hc = item[2] if len(item) > 2 else C_SOC
            ax.axhline(yv, linestyle='--', linewidth=1.0, color=hc, alpha=0.55, label=lab)
    if not NO_TITLE:
        ax.set_title(title, fontsize=13, fontweight='bold', pad=28)
    ax.set_xlabel(xlabel, fontsize=AXIS_FS, color='black')
    ax.set_ylabel(ylabel, fontsize=AXIS_FS, color='black')
    _style_axes(ax)
    if labels or hlines:
        if legend_loc:  # 例如 (1.0, 0.90)：右上、限值线下方
            ax.legend(loc='upper right', frameon=False, ncol=1,
                      fontsize=LEGEND_FS-2, bbox_to_anchor=legend_loc)
        else:           # 默认：图外顶部居中
            ax.legend(loc='lower center', frameon=False,
                      ncol=max(1, len(labels)), fontsize=LEGEND_FS,
                      bbox_to_anchor=(0.5, 1.02), columnspacing=1.8)
    if date_axis:
        ax.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=5, maxticks=9))
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'))
        fig.autofmt_xdate(rotation=25)
    if xticks is not None:
        ax.set_xticks(xticks)
    if xlim is not None:
        ax.set_xlim(*xlim)
    if ylim is not None:
        ax.set_ylim(*ylim)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches='tight')
    plt.close(fig)
def _save_bar_fig(xlabels, values, title, xlabel, ylabel, path, color=C_BUY):
    fig, ax = plt.subplots(figsize=(10.8, 5.4), dpi=300)
    xx=np.arange(len(values))
    ax.bar(xx, values, width=0.55, color=color, alpha=0.85)
    ax.set_xticks(xx)
    ax.set_xticklabels(xlabels, fontsize=TICK_FS)
    if not NO_TITLE:
        ax.set_title(title, fontsize=13, fontweight='bold', pad=12)
    ax.set_xlabel(xlabel, fontsize=AXIS_FS, color='black')
    ax.set_ylabel(ylabel, fontsize=AXIS_FS, color='black')
    _style_axes(ax)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches='tight')
    plt.close(fig)
def generate_visualizations(dates, price_actual, price_fc_store, plan_buy, run_D, cfg:Config, logger):
    fig_dir=Path(cfg.fig_dir); fig_dir.mkdir(parents=True,exist_ok=True)
    valid=[i for i in range(run_D) if dates[i]>=date(2025,2,1) and np.isfinite(plan_buy[i]).all()]
    if not valid: return
    tt=np.arange(144)/6
    # 指定日波动电价预测对比（3-20 / 6-21 / 9-23 / 12-21）
    for td in [date(2025,3,20),date(2025,6,21),date(2025,9,23),date(2025,12,21)]:
        i=next((j for j in valid if dates[j]==td),None)
        if i is not None:
            _save_line_fig(tt,[price_actual[i],price_fc_store[i,:144]],
                           ['实际电价','0:00预测电价'],f'{td} 波动电价预测对比',
                           '时刻 / h','电价 / (元/kWh)',
                           fig_dir/f'第四问{td.month}-{td.day}波动电价预测对比.png',
                           colors=[C_BUY,C_DIS],line_styles=['-','--'],
                           linewidths=[2.0,1.8],
                           xticks=list(range(0,25,2)),xlim=(0,24))
    # 全年10分钟级电价预测一致性散点图
    a=price_actual[valid].reshape(-1); p=price_fc_store[valid,:144].reshape(-1)
    fig,ax=plt.subplots(figsize=(6.2,6.2),dpi=300)
    ax.scatter(a,p,s=6,alpha=.06,color=C_BUY)
    lo=float(min(a.min(),p.min())); hi=float(max(a.max(),p.max()))
    ax.plot([lo,hi],[lo,hi],'--',lw=1.4,color=C_CH,label='y=x')
    if not NO_TITLE:
        ax.set_title('全年10分钟级电价预测一致性', fontsize=13, fontweight='bold')
    ax.set_xlabel('实际电价 / (元/kWh)', fontsize=AXIS_FS, color='black')
    ax.set_ylabel('预测电价 / (元/kWh)', fontsize=AXIS_FS, color='black')
    ax.tick_params(labelsize=TICK_FS)
    _style_axes(ax)
    ax.legend(loc='upper left', frameon=False, fontsize=13)
    fig.tight_layout()
    fig.savefig(fig_dir/'第四问全年10分钟级电价预测一致性.png',dpi=300,bbox_inches='tight')
    plt.close(fig)
def write_result_summary(dates, price_actual, price_fc_store, plan_buy, emg, sur, ch, dis, soc_end,
                         alpha_store, run_D, cfg:Config, logger):
    valid=[i for i in range(run_D) if dates[i]>=date(2025,2,1) and np.isfinite(plan_buy[i]).all()]
    if not valid: return
    plan_cost=sum(float(np.sum(plan_buy[i]*price_actual[i])*cfg.dt) for i in valid)
    emg_cost=sum(float(np.sum(emg[i]*5*price_actual[i])*cfg.dt) for i in valid)
    emg_energy=sum(float(np.sum(emg[i])*cfg.dt) for i in valid)
    sur_energy=sum(float(np.sum(sur[i])*cfg.dt) for i in valid)
    charge_energy=sum(float(np.sum(ch[i])*cfg.dt) for i in valid)
    discharge_energy=sum(float(np.sum(dis[i])*cfg.dt) for i in valid)
    ae=np.concatenate([(price_actual[i]-price_fc_store[i,:144]) for i in valid])
    aa=np.concatenate([price_actual[i] for i in valid])
    price_mae=float(np.mean(np.abs(ae))); price_rmse=float(np.sqrt(np.mean(ae**2)))
    denom=np.maximum(np.abs(aa),1e-3); price_mape=float(np.mean(np.abs(ae)/denom)*100)
    zero=sum(1 for i in valid if np.sum(emg[i])*cfg.dt<=1e-8)
    lines=['问题4-2结果整合（电价预测驱动，1月仅预热）',
        f'正式统计日期范围：{dates[valid[0]]} 至 {dates[valid[-1]]}',
        f'累计计划购电费（真实电价结算）：{plan_cost:.2f} 元',
        f'累计紧急购电费（5倍真实电价）：{emg_cost:.2f} 元',
        f'累计总费用：{plan_cost+emg_cost:.2f} 元',
        f'累计紧急购电量：{emg_energy:.2f} kWh',
        f'累计富余电量：{sur_energy:.2f} kWh',
        f'累计充电量：{charge_energy:.2f} kWh',
        f'累计放电量：{discharge_energy:.2f} kWh',
        f'期末储电量：{soc_end[valid[-1]]:.2f} kWh',
        f'无紧急购电天数：{zero}/{len(valid)}',
        f'电价预测MAE：{price_mae:.6f} 元/kWh',
        f'电价预测RMSE：{price_rmse:.6f} 元/kWh',
        f'电价预测MAPE：{price_mape:.3f}%',
        f'动态alpha平均值：{float(np.nanmean(alpha_store[valid])):.4f}',
        f'正式结果文件：{cfg.o_xlsx}',f'可视化目录：{cfg.fig_dir}']
    p=Path(cfg.o_dir)/'Q4-2_结果整合.txt'; p.write_text('\n'.join(lines),encoding='utf-8')
def main():
    cfg=CFG; out_dir=Path(cfg.o_dir); out_dir.mkdir(parents=True,exist_ok=True); logger=setup_logger(out_dir)
    logger.info('Q4-2')
    logger.info('负荷与电价共享MSTL重估周期=%d天，但模型缓存完全独立。',cfg.mstl_days)
    dates, price_prior, load_prior, pv_prior, load, pv, price_actual = load_inputs(cfg,logger)
    D=len(dates); n=cfg.slots_pday; run_D=min(D,cfg.max_days) if cfg.max_days>0 else D
    plan_buy=np.full((D,n),np.nan); plan_ch=np.zeros((D,n)); plan_dis=np.zeros((D,n))
    ch=np.zeros((D,n)); dis=np.zeros((D,n)); emg=np.zeros((D,n)); sur=np.zeros((D,n))
    soc_hist=np.full((D,n),np.nan); soc_start=np.full(D,np.nan); soc_end=np.full(D,np.nan)
    load_fc_store=np.full((D,cfg.fore_slots),np.nan); pv_fc_store=np.full((D,cfg.fore_slots),np.nan) 
    price_fc_store=np.full((D,cfg.fore_slots),np.nan)
    alpha_store=np.full((D,n),np.nan); reserve_store=np.zeros((D,n)); net_errors=np.full((D,n),np.nan)
    cplus_cache=np.full((D,n),np.nan); cminus_cache=np.full((D,n),np.nan); mu_store=np.full(D,np.nan)
    lp=build_lp_structure(cfg); soc=cfg.e0
    for d in range(run_D):
        soc_start[d]=soc
        lf=load_forecast(d,load,load_prior,cfg,logger)
        pf,_=pv_forecast(d,pv,pv_prior,cfg,logger)
        prf=price_forecast(d,price_actual,price_prior,cfg,logger)
        load_fc_store[d]=lf; pv_fc_store[d]=pf; price_fc_store[d]=prf
        update_risk_cost_cache(d,cplus_cache,cminus_cache,net_errors,price_actual,soc_hist,plan_buy,load,pv,cfg)
        alpha,_,reserve=dynamic_risk_weights(d,net_errors,cplus_cache,cminus_cache,cfg)
        alpha_store[d]=alpha; reserve_store[d]=reserve
        da=solve_day_ahead(prf,lf,pf,reserve,soc,cfg,lp); plan_buy[d]=da['buy']; plan_ch[d]=da['ch']; plan_dis[d]=da['dis']; mu_store[d]=da['mu_terminal']
        for t in range(n):
            soc_hist[d,t]=soc
            sol=operate_one_slot(soc,plan_buy[d,t],load[d,t],pv[d,t],plan_ch[d,t],plan_dis[d,t],cfg)
            ch[d,t]=sol['ch']; dis[d,t]=sol['dis']; emg[d,t]=sol['emg']; sur[d,t]=sol['sur']; soc=sol['E']
        soc_end[d]=soc
        net_errors[d]=(load[d]-pv[d])-(lf[:n]-pf[:n])
        if cfg.checkpoint_every_days>0 and ((d+1)%cfg.checkpoint_every_days==0 or d==run_D-1):
            np.savez_compressed(out_dir/'q4_2_checkpoint.npz',last_day=d,plan_buy=plan_buy[:d+1],ch=ch[:d+1],dis=dis[:d+1],emg=emg[:d+1],sur=sur[:d+1],soc_start=soc_start[:d+1],soc_end=soc_end[:d+1],load_fc=load_fc_store[:d+1],pv_fc=pv_fc_store[:d+1],price_fc=price_fc_store[:d+1],alpha=alpha_store[:d+1],reserve=reserve_store[:d+1],mu=mu_store[:d+1])
    if run_D>=32 and Path(cfg.temp).exists():
        fill_result4_2(cfg.temp,cfg.o_xlsx,dates,price_actual,plan_buy,ch,dis,soc_start,soc_end,emg,cfg,logger)
    elif run_D>=32:
        logger.warning('未找到 %s；图和其余结果仍正常输出。',cfg.temp)
    write_specified_date_results(dates,price_actual,plan_buy,ch,dis,emg,soc_start,soc_end,cfg,logger)
    generate_visualizations(dates,price_actual,price_fc_store,plan_buy,run_D,cfg,logger)
    write_result_summary(dates,price_actual,price_fc_store,plan_buy,emg,sur,ch,dis,soc_end,alpha_store,run_D,cfg,logger)
    logger.info('Q4-2全部完成，输出目录：%s',out_dir)
if __name__ == '__main__':
    main()