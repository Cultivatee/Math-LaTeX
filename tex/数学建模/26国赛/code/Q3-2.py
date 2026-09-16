from __future__ import annotations
import csv
import math
import importlib.util
import sys
from pathlib import Path
from datetime import datetime, date, timedelta
from typing import Dict, List, Tuple
import numpy as np
from openpyxl import load_workbook, Workbook
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False; plt.rcParams["font.size"] = 12
plt.rcParams["axes.spines.top"] = False; plt.rcParams["axes.linewidth"] = 1.5
plt.rcParams["xtick.major.width"] = 1.5; plt.rcParams["ytick.major.width"] = 1.5
plt.rcParams["xtick.major.size"] = 6; plt.rcParams["ytick.major.size"] = 6
C_BUY, C_PV, C_DIS, C_LOAD = "#2E74B5", "#27AE60", "#E67E22", "#34495E"
C_CH, C_SOC = "#E74C3C", "#7F8C8D"
class RunConfig:
    BASE = Path(__file__).resolve().parent if "__file__" in globals() else Path(".")
    Q3_MAIN_CANDIDATES = [BASE / "Q3.py"]
    ATT1_CANDIDATES = [BASE / "附件1.xlsx", BASE / "附件1(3).xlsx", BASE / "附件1(1).xlsx"]
    ATT2_CANDIDATES = [BASE / "附件2.xlsx", BASE / "附件2(3).xlsx", BASE / "附件2(2).xlsx"]
    ATT3_CANDIDATES = [BASE / "附件3.xlsx", BASE / "附件3(1).xlsx"]
    CHECKPOINT_CANDIDATES = [
        BASE / "Q3结果" / "q3_checkpoint_forecast_driven.npz",
        BASE / "q3_outputs_forecast_driven" / "q3_checkpoint_forecast_driven.npz",
        BASE / "q3_checkpoint_forecast_driven.npz", BASE / "q3_checkpoint_forecast_driven(1).npz",
    ]
    OUT = BASE / "Q3_2结果"
    FIG = OUT / "Q3_2图片输出"
    CACHE = OUT / "Q3_02_preprocess_cache.npz"
    CANDIDATE_INTERVALS = ((0, 6), (6, 12), (12, 18))
    MAX_EXTRA_HOURS, MIN_INTERVAL_RISK_SHARE = 2, 0.05
    MANUAL_EXTRA_HOURS = None     # 如需固定，可改为 (10, 14)
    RECENT_ERROR_HOURS, BIAS_DECAY_HOURS, HISTORY_DAYS_FOR_TAIL = 3, 4.0, 30
    MAX_BIAS_ABS_KW, MAX_DAYS = 1500.0, 365
    REPRESENTATIVE_DATES = (date(2025, 3, 20), date(2025, 6, 21),
                            date(2025, 9, 23), date(2025, 12, 21))
def first_existing(paths: List[Path]) -> Path:
    for p in paths:
        if p.exists():
            return p
    raise FileNotFoundError("未找到文件：\n" + "\n".join(str(p) for p in paths))
def ensure_dir(p: Path):
    p.mkdir(parents=True, exist_ok=True)
def import_py(path: Path, module_name: str):
    spec = importlib.util.spec_from_file_location(module_name, str(path))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载Python脚本：{path}")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = mod
    spec.loader.exec_module(mod)
    return mod
def write_csv(path: Path, rows: List[dict]):
    if not rows:
        return
    ensure_dir(path.parent)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)
def style(ax):
    ax.grid(True, axis="y", linestyle="--", linewidth=0.7, alpha=0.3)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(labelsize=TICK_FS)
NO_TITLE, AXIS_FS, TICK_FS, LEGEND_FS = True, 15, 12, 17
SOFT_RED, SOFT_GREEN, SOFT_BLUE, SOFT_GRAY = '#E9A6A1', '#8BCF8B', '#7097CA', '#CFCECE'
BAR_EDGE  = dict(edgecolor='black', linewidth=1.2)
def _set_labels(ax, x, y):
    ax.set_xlabel(x, fontsize=AXIS_FS, color='black')
    ax.set_ylabel(y, fontsize=AXIS_FS, color='black')
def _legend_top(ax, ncol=None):
    h, l = ax.get_legend_handles_labels()
    if h:
        ax.legend(h, l, loc='lower center', frameon=False, ncol=ncol or len(h),
                  fontsize=LEGEND_FS, bbox_to_anchor=(0.5, 1.02), columnspacing=1.8)
def load_core_inputs(att1: Path, att2: Path, q3):
    wb1 = load_workbook(att1, data_only=True, read_only=True)
    ws1 = wb1[wb1.sheetnames[0]]
    rows1 = list(ws1.iter_rows(min_row=2, max_row=145, min_col=2, max_col=4, values_only=True))
    wb1.close()
    price = np.asarray([float(r[0]) for r in rows1], dtype=float)
    load_prior = np.asarray([float(r[1]) for r in rows1], dtype=float)
    pv_prior = np.asarray([float(r[2]) for r in rows1], dtype=float)
    wb2 = load_workbook(att2, data_only=True, read_only=True)
    sl, sp = wb2["小区负载"], wb2["光伏发电实际功率"]
    load_rows = list(sl.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    pv_rows = list(sp.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    wb2.close()
    dates = [q3.cell_to_date(r[0]) for r in load_rows]
    load = np.asarray([[float(v) for v in r[1:145]] for r in load_rows], dtype=float)
    pv = np.asarray([[float(v) for v in r[1:145]] for r in pv_rows], dtype=float)
    if price.shape != (144,) or load.shape != (365, 144) or pv.shape != (365, 144):
        raise ValueError(f"输入尺寸异常：price={price.shape}, load={load.shape}, pv={pv.shape}")
    return dates, price, load_prior, pv_prior, load, pv
def load_checkpoint(cp_path: Path, run_D: int):
    d = np.load(cp_path, allow_pickle=True)
    required = ["soc_start", "ch", "dis", "fallback_reserve", "stage_net_error_hist"]
    missing = [k for k in required if k not in d.files]
    if missing:
        raise ValueError(
            "Q3_01 checkpoint缺少字段：" + ", ".join(missing) +
            "。请先用当前正式Q3主程序完整运行一次。"
        )
    out = {k: np.asarray(d[k]) for k in d.files}
    if len(out["soc_start"]) < run_D:
        raise ValueError(
            f"checkpoint只运行到{len(out['soc_start'])}天，但本程序要求{run_D}天。"
            "请先把Q3_01完整运行到相同天数。"
        )
    return out
def reconstruct_soc_hourly(checkpoint: dict, run_D: int, cfg) -> np.ndarray:
    """用Q3_01真实10分钟充放电轨迹重构每小时起点SOC，供风险带使用。"""
    soc_start = np.asarray(checkpoint["soc_start"][:run_D], dtype=float)
    ch = np.asarray(checkpoint["ch"][:run_D], dtype=float)
    dis = np.asarray(checkpoint["dis"][:run_D], dtype=float)
    out = np.zeros((run_D, 24), dtype=float)
    for d in range(run_D):
        soc = float(soc_start[d])
        for h in range(24):
            out[d, h] = soc
            for k in range(h * 6, h * 6 + 6):
                soc += cfg.eta_c * ch[d, k] * cfg.dt - dis[d, k] * cfg.dt / cfg.eta_d
                soc = min(max(soc, cfg.e_min), cfg.e_max)
    return out
class RiskBandConfig:
    DAYLIGHT_THRESHOLD, BAND_QUANTILE, BAND_MIN_KW, BAND_FLOOR_RATIO = 100.0, 0.85, 120.0, 0.08
    USE_HOUR_SMOOTH, C_ES, C_SUR, K_SHORT = True, 0.03, 0.06, 5.0
    K_DIS, K_CH, EPS = 1.5, 0.5, 1e-9
def load_forecast_records_xlsx(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws, current_date = wb[wb.sheetnames[0]], None
    recs = []
    for r in ws.iter_rows(min_row=2, max_col=26, values_only=True):
        if r[0] not in (None, ""):
            if isinstance(r[0], datetime):
                current_date = r[0].date()
            elif isinstance(r[0], date):
                current_date = r[0]
            else:
                current_date = datetime.strptime(str(r[0]).replace('/', '-')[:10], '%Y-%m-%d').date()
        if current_date is None:
            continue
        txt = str(r[1]).strip()
        issue_hour = int(txt.split(':')[0]) if ':' in txt else int(float(txt))
        fc = np.asarray([0.0 if v is None else float(v) for v in r[2:26]], dtype=float)
        issue_dt = datetime.combine(current_date, datetime.min.time()) + timedelta(hours=issue_hour)
        recs.append({
            'issue_dt': issue_dt,
            'issue_hour': issue_hour,
            'fc': np.maximum(fc, 0.0),
        })
    wb.close()
    recs.sort(key=lambda z: z['issue_dt'])
    return recs
def build_target_times(dates: List[date]):
    return [datetime.combine(dd, datetime.min.time()) + timedelta(hours=h)
            for dd in dates for h in range(1, 25)]
def effective_forecast_records(records: List[dict], target_times: List[datetime]):
    values = np.full(len(target_times), np.nan)
    for i, t in enumerate(target_times):
        candidates = []
        for r in records:
            if r['issue_dt'] >= t:
                break
            val, _ = record_value_at_target(r, t)
            if val is not None:
                candidates.append((r['issue_dt'], val))
        if candidates:
            values[i] = max(candidates, key=lambda z: z[0])[1]
    if np.isnan(values).any():
        bad = np.where(np.isnan(values))[0]
        raise ValueError(f'有{len(bad)}个目标时刻缺少生效预报，首个索引={bad[0]}')
    return values
def rb_qsafe(arr, q):
    x = np.asarray(arr, dtype=float); x = x[np.isfinite(x)]
    return float(np.quantile(x, q)) if x.size else np.nan
def rb_mean(arr):
    x = np.asarray(arr, dtype=float); x = x[np.isfinite(x)]
    return float(np.mean(x)) if x.size else np.nan
def rb_rmse(arr):
    x = np.asarray(arr, dtype=float); x = x[np.isfinite(x)]
    return float(np.sqrt(np.mean(x*x))) if x.size else np.nan
def rb_mae(arr):
    x = np.asarray(arr, dtype=float); x = x[np.isfinite(x)]
    return float(np.mean(np.abs(x))) if x.size else np.nan
def risk_economic_costs(price: float, load: float, fc_pv: float, soc: float, cfg):
    a_dis = max(0.0, min(cfg.p_max, cfg.eta_d * max(soc - cfg.e_min, 0.0)))
    a_ch = max(0.0, min(cfg.p_max, max(cfg.e_max - soc, 0.0) / cfg.eta_c))
    net = load - fc_pv; shortage = max(net, 0.0); surplus = max(-net, 0.0)
    pi_plus = a_dis / (a_dis + shortage + RiskBandConfig.EPS)
    pi_minus = a_ch / (a_ch + surplus + RiskBandConfig.EPS)
    c_plus = RiskBandConfig.C_ES + price * (
        RiskBandConfig.K_DIS * pi_plus + RiskBandConfig.K_SHORT * (1.0 - pi_plus)
    )
    c_minus = RiskBandConfig.C_ES + RiskBandConfig.K_CH * price * pi_minus + RiskBandConfig.C_SUR * (1.0 - pi_minus)
    return c_plus, c_minus
def build_hour_scale(actual_flat, fc_flat, clock):
    err = fc_flat - actual_flat
    daylight = (actual_flat >= RiskBandConfig.DAYLIGHT_THRESHOLD) | (fc_flat >= RiskBandConfig.DAYLIGHT_THRESHOLD)
    base = {}
    for h in range(1, 25):
        m = (clock == h) & daylight
        if np.any(m):
            qv = rb_qsafe(np.abs(err[m]), RiskBandConfig.BAND_QUANTILE)
            med = float(np.median(actual_flat[m]))
            floor = max(RiskBandConfig.BAND_FLOOR_RATIO * med, RiskBandConfig.BAND_MIN_KW)
            base[h] = max(0.0 if np.isnan(qv) else float(qv), floor)
        else:
            base[h] = 0.0
    if RiskBandConfig.USE_HOUR_SMOOTH:
        vals = np.array([base[h] for h in range(1, 25)], dtype=float); sm = vals.copy()
        for i in range(24):
            sm[i] = 0.25*vals[max(i-1,0)] + 0.50*vals[i] + 0.25*vals[min(i+1,23)]
        return {h: float(sm[h-1]) for h in range(1,25)}
    return dict(base)
def analyze_risk_scheme(name: str, dates, load_hourly, pv_actual_hourly,
                        price_hourly, soc_hourly, fc_flat, cfg):
    target_times = build_target_times(dates)
    actual = pv_actual_hourly.reshape(-1)
    load = load_hourly.reshape(-1)
    price = np.tile(price_hourly, len(dates))
    soc = soc_hourly.reshape(-1)
    clock = np.asarray([24 if t.hour == 0 else t.hour for t in target_times], dtype=int)
    error = fc_flat - actual
    smooth_scale = build_hour_scale(actual, fc_flat, clock)
    n = len(target_times)
    lower, upper, bp, bm, danger, safe = (np.zeros(n) for _ in range(6))
    for i in range(n):
        cp, cm = risk_economic_costs(price[i], load[i], fc_flat[i], soc[i], cfg)
        w = smooth_scale[int(clock[i])]
        den = max(cp + cm, RiskBandConfig.EPS)
        bplus, bminus = 2.0 * cm / den * w, 2.0 * cp / den * w
        lo, hi = actual[i] - bminus, actual[i] + bplus
        d_ex = max(fc_flat[i] - hi, 0.0)
        s_ex = max(lo - fc_flat[i], 0.0)
        lower[i]=lo; upper[i]=hi; bp[i]=bplus; bm[i]=bminus
        danger[i]=d_ex; safe[i]=s_ex
    hourly=[]
    for h in range(1,25):
        m=(clock==h)
        hourly.append({
            '目标小时':h, '样本数':int(np.sum(m)),
            '危险高估次数':int(np.sum(danger[m]>0)),
            '危险高估总量kW':float(np.sum(danger[m])),
            '保守低估次数':int(np.sum(safe[m]>0)),
            '保守低估总量kW':float(np.sum(safe[m])),
            '总带外次数':int(np.sum((danger[m]+safe[m])>0)),
            '总带外量kW':float(np.sum(danger[m]+safe[m])),
            'RMSE':rb_rmse(error[m]), 'MAE':rb_mae(error[m]),
            '平均b_plus':rb_mean(bp[m]), '平均b_minus':rb_mean(bm[m]),
        })
    overall={
        '方案':name,
        '危险高估次数':int(np.sum(danger>0)),
        '危险高估总量kW':float(np.sum(danger)),
        '保守低估次数':int(np.sum(safe>0)),
        '保守低估总量kW':float(np.sum(safe)),
        '总带外次数':int(np.sum((danger+safe)>0)),
        '总带外量kW':float(np.sum(danger+safe)),
        'RMSE':rb_rmse(error), 'MAE':rb_mae(error),
        '平均b_plus':rb_mean(bp), '平均b_minus':rb_mean(bm),
    }
    return {
        'target_times':target_times,'actual':actual,'forecast':fc_flat,
        'lower':lower,'upper':upper,'danger':danger,'safe':safe,
        'hourly':hourly,'overall':overall,
    }
def risk_compare_table(base, enhanced):
    keys=['危险高估次数','危险高估总量kW','保守低估次数','保守低估总量kW',
          '总带外次数','总带外量kW','RMSE','MAE','平均b_plus','平均b_minus']
    out=[]
    for k in keys:
        b=float(base['overall'][k]); e=float(enhanced['overall'][k])
        out.append({'指标':k,'原方案':b,'增强方案':e,
                    '变化率%':((e-b)/b*100.0 if abs(b)>1e-12 else np.nan)})
    return out
def risk_day_slice(analysis, day_str):
    dd=datetime.strptime(day_str,'%Y-%m-%d').date()
    idx=[]
    for i,t in enumerate(analysis['target_times']):
        td=t.date() if t.hour!=0 else t.date()-timedelta(days=1)
        if td==dd: idx.append(i)
    if len(idx)!=24: return None
    return slice(idx[0],idx[-1]+1)
def plot_risk_day_compare(base, enhanced, day_str, path):
    sl1=risk_day_slice(base,day_str); sl2=risk_day_slice(enhanced,day_str)
    if sl1 is None or sl2 is None: return
    fig,axes=plt.subplots(2,1,figsize=(12.5,9.0),sharex=True,dpi=300)
    for ax,ana,sl,ptitle in [(axes[0],base,sl1,'原四时点预报'),(axes[1],enhanced,sl2,'增加预报后')]:
        x=np.arange(1,25); ac=ana['actual'][sl]; fc=ana['forecast'][sl]
        lo=ana['lower'][sl]; hi=ana['upper'][sl]
        md=ana['danger'][sl]>0; ms=ana['safe'][sl]>0
        ax.fill_between(x,lo,hi,alpha=.12,color=C_BUY,label='动态非对称风险带')
        ax.plot(x,ac,lw=2.2,color=C_PV,label='实际光伏')
        ax.plot(x,fc,lw=2.2,color=C_BUY,label='生效光伏预报')
        ax.plot(x,lo,lw=1.0,ls='--',alpha=.75,color=C_SOC,label='下界')
        ax.plot(x,hi,lw=1.0,ls='--',alpha=.75,color=C_SOC,label='上界')
        if np.any(md): ax.scatter(x[md],fc[md],s=46,marker='^',color=C_CH,zorder=6,label='危险高估点')
        if np.any(ms): ax.scatter(x[ms],fc[ms],s=46,marker='v',color=C_DIS,zorder=6,label='保守低估点')
        ax.text(0.012,0.97,ptitle,transform=ax.transAxes,fontsize=13,
                color='#444444',va='top',ha='left')
        _set_labels(ax,'','光伏功率 / kW')
        style(ax)
        ax.legend(frameon=False,ncol=4, fontsize=14, loc='lower center',
                  bbox_to_anchor=(0.5,1.02), columnspacing=1.4)
    axes[-1].set_xlim(1,24); axes[-1].set_xticks(range(1,25))
    _set_labels(axes[-1],'目标小时 / h','')
    fig.tight_layout(); fig.savefig(path,dpi=300,bbox_inches='tight'); plt.close(fig)
def plot_risk_overall(base, enhanced, path):
    labels=['危险高估总量kW','保守低估总量kW','总带外量kW','RMSE','MAE']
    b=[base['overall'][k] for k in labels]; e=[enhanced['overall'][k] for k in labels]
    rel=[100*ee/bb if abs(bb)>1e-12 else np.nan for ee,bb in zip(e,b)]
    x=np.arange(len(labels)); width=.36
    fig,ax=plt.subplots(figsize=(10.8,5.4),dpi=300)
    ax.bar(x-width/2,[100]*len(labels),width,color=SOFT_GRAY,alpha=0.95,label='原方案',zorder=3,**BAR_EDGE)
    ax.bar(x+width/2,rel,width,color=SOFT_GREEN,alpha=0.95,label='增强方案',zorder=3,**BAR_EDGE)
    ax.axhline(100,ls='--',lw=1.0,alpha=0.55,color=C_LOAD)
    ax.set_xticks(x,labels)
    if not NO_TITLE:
        ax.set_title('增强方案相对原方案的风险指标变化', fontsize=13, fontweight='bold')
    _set_labels(ax,'风险指标','相对原方案 / %')
    _legend_top(ax,ncol=2)
    style(ax)
    fig.tight_layout(); fig.savefig(path,dpi=300,bbox_inches='tight'); plt.close(fig)
def interval_candidate(hourly_rows: List[dict], start: int, end: int):
    stats = {int(r["目标小时"]): r for r in hourly_rows}
    hours = list(range(start + 1, end + 1))
    total_energy = sum(float(stats[h]["危险高估总量kW"]) for h in hours)
    total_count = sum(int(stats[h]["危险高估次数"]) for h in hours)
    if total_energy <= 0:
        return None
    weights = []
    for h in hours:
        # 以危险高估总量为主，次数为辅。
        w = float(stats[h]["危险高估总量kW"]) + 5.0 * int(stats[h]["危险高估次数"])
        weights.append(max(w, 0.0))
    centroid = sum(h * w for h, w in zip(hours, weights)) / max(sum(weights), 1e-12)
    candidate = int(round(centroid - 1.0))
    candidate = max(start + 1, min(candidate, end - 1 if end - start > 1 else end))
    return {
        "区间": f"{start}:00-{end}:00", "风险重心": centroid, "候选时点": candidate,
        "危险高估次数": total_count, "危险高估总量kW": total_energy,
    }
def select_extra_hours(base_analysis: dict):
    if RunConfig.MANUAL_EXTRA_HOURS is not None:
        return sorted(set(int(x) for x in RunConfig.MANUAL_EXTRA_HOURS)), []
    candidates = []
    for a, b in RunConfig.CANDIDATE_INTERVALS:
        r = interval_candidate(base_analysis["hourly"], a, b)
        if r is not None:
            candidates.append(r)
    total = sum(r["危险高估总量kW"] for r in candidates)
    for r in candidates:
        r["风险占比"] = r["危险高估总量kW"] / max(total, 1e-12)
    kept = [r for r in candidates if r["风险占比"] >= RunConfig.MIN_INTERVAL_RISK_SHARE]
    kept.sort(key=lambda z: z["危险高估总量kW"], reverse=True)
    kept = kept[:RunConfig.MAX_EXTRA_HOURS]
    hours = sorted({int(r["候选时点"]) for r in kept})
    return hours, candidates
def actual_hour_value(target_dt: datetime, dates: List[date], actual_hourly: np.ndarray):
    if target_dt.hour == 0:
        dd = target_dt.date() - timedelta(days=1)
        h = 24
    else:
        dd, h = target_dt.date(), target_dt.hour
    try:
        di = dates.index(dd)
    except ValueError:
        return None
    return float(actual_hourly[di, h - 1])
def past_climatology(target_dt: datetime, dates: List[date], actual_hourly: np.ndarray,
                     issue_dt: datetime):
    if target_dt.hour == 0:
        target_date = target_dt.date() - timedelta(days=1)
        h = 24
    else:
        target_date, h = target_dt.date(), target_dt.hour
    vals = []
    for back in range(1, RunConfig.HISTORY_DAYS_FOR_TAIL + 1):
        dd = target_date - timedelta(days=back)
        if dd >= issue_dt.date():
            continue
        try:
            di = dates.index(dd)
        except ValueError:
            continue
        vals.append(float(actual_hourly[di, h - 1]))
    return max(float(np.median(vals)), 0.0) if vals else 0.0
def latest_record_before(records: List[dict], issue_dt: datetime):
    c = [r for r in records if r["issue_dt"] < issue_dt]
    return max(c, key=lambda z: z["issue_dt"]) if c else None
def record_value_at_target(rec: dict, target_dt: datetime):
    lead = int(round((target_dt - rec["issue_dt"]).total_seconds() / 3600.0))
    if 1 <= lead <= 24:
        return float(rec["fc"][lead - 1]), lead
    return None, None
def build_extra_record(extra_issue_dt: datetime, official_records: List[dict],
                       dates: List[date], actual_hourly: np.ndarray):
    old = latest_record_before(official_records, extra_issue_dt)
    if old is None:
        return None
    residuals, weights = [], []
    for back in range(RunConfig.RECENT_ERROR_HOURS - 1, -1, -1):
        target_dt = extra_issue_dt - timedelta(hours=back)
        old_val, _ = record_value_at_target(old, target_dt)
        actual_val = actual_hour_value(target_dt, dates, actual_hourly)
        if old_val is None or actual_val is None:
            continue
        residuals.append(actual_val - old_val)
        weights.append(math.exp(-back / 1.5))
    if residuals:
        w = np.asarray(weights, dtype=float); w /= w.sum()
        bias = float(np.dot(w, np.asarray(residuals, dtype=float)))
    else:
        bias = 0.0
    bias = float(np.clip(bias, -RunConfig.MAX_BIAS_ABS_KW, RunConfig.MAX_BIAS_ABS_KW))
    fc_new = np.zeros(24, dtype=float)
    for lead in range(1, 25):
        target_dt = extra_issue_dt + timedelta(hours=lead)
        base_val, _ = record_value_at_target(old, target_dt)
        if base_val is None:
            base_val = past_climatology(target_dt, dates, actual_hourly, extra_issue_dt)
        correction = bias * math.exp(-(lead - 1) / RunConfig.BIAS_DECAY_HOURS)
        if base_val < 100.0:
            correction *= 0.25
        fc_new[lead - 1] = max(base_val + correction, 0.0)
    return {"issue_dt": extra_issue_dt, "issue_hour": extra_issue_dt.hour, "fc": fc_new}
def build_augmented_records(official_records: List[dict], extra_hours: List[int],
                            dates: List[date], actual_hourly: np.ndarray):
    out = [dict(r) for r in official_records]
    official_dates = sorted({r["issue_dt"].date() for r in official_records})
    for dd in official_dates:
        for h in extra_hours:
            issue_dt = datetime.combine(dd, datetime.min.time()) + timedelta(hours=h)
            rec = build_extra_record(issue_dt, official_records, dates, actual_hourly)
            if rec is not None:
                out.append(rec)
    out.sort(key=lambda z: z["issue_dt"])
    return out
def export_records_xlsx(path: Path, records: List[dict]):
    wb = Workbook(); ws = wb.active
    ws.title = "Sheet1"
    ws.append(["日期", "预报时刻"] + [f"预报{i}小时" for i in range(1, 25)])
    by_date: Dict[date, List[dict]] = {}
    for r in records:
        by_date.setdefault(r["issue_dt"].date(), []).append(r)
    for dd in sorted(by_date):
        rr = sorted(by_date[dd], key=lambda z: z["issue_dt"])
        for j, r in enumerate(rr):
            date_cell = dd.isoformat() if j == 0 else None
            ws.append([date_cell, f"{r['issue_hour']}:00"] + [float(v) for v in r["fc"]])
    ensure_dir(path.parent)
    wb.save(path)
    wb.close()
def build_or_load_preprocess_cache(q3, cfg, dates, load, load_prior, run_D: int):
    if RunConfig.CACHE.exists():
        z = np.load(RunConfig.CACHE, allow_pickle=True)
        if "load_fc_full_store" in z.files:
            arr = np.asarray(z["load_fc_full_store"], dtype=float)
            if arr.shape[0] >= run_D and arr.shape[1] == cfg.fore_slots:
                return arr[:run_D].copy()
    if hasattr(q3, "_LOAD_MSTL_CACHE"):
        q3._LOAD_MSTL_CACHE.clear()
    store = np.full((run_D, cfg.fore_slots), np.nan, dtype=float)
    for d in range(run_D):
        fc = q3.load_forecast(d, load, load_prior, cfg, logger=None)
        store[d] = np.asarray(fc[:cfg.fore_slots], dtype=float)
    np.savez_compressed(RunConfig.CACHE, load_fc_full_store=store)
    return store
def records_by_key(records: List[dict]):
    return {(r["issue_dt"].date(), int(r["issue_hour"])): r for r in records}
def pv_record_horizon_10min(day_idx: int, stage_hour: int, rec_map: dict,
                            dates: List[date], pv_actual: np.ndarray, cfg):
    key = (dates[day_idx], int(stage_hour))
    if key not in rec_map:
        raise KeyError(f"缺少{dates[day_idx]} {stage_hour}:00的光伏预报记录")
    hourly = np.asarray(rec_map[key]["fc"], dtype=float)
    start_slot = int(stage_hour) * 6
    anchor = 0.0 if stage_hour == 0 else float(pv_actual[day_idx, start_slot - 1])
    xp = np.arange(0, 25, dtype=float)
    fp = np.concatenate([[anchor], hourly])
    target_rel = (np.arange(cfg.hori_slots) + 1) / 6.0
    pred = np.interp(target_rel, xp, fp)
    return np.maximum(pred, 0.0)
def actual_horizon(day_idx: int, stage_hour: int, arr: np.ndarray, cfg):
    start = stage_hour * 6
    H = cfg.hori_slots
    flat = arr.reshape(-1)
    g0 = day_idx * cfg.slots_pday + start
    if g0 + H > len(flat):
        return None
    return np.asarray(flat[g0:g0 + H], dtype=float)
def build_extra_stage_error_history(extra_hours: List[int], rec_map: dict,
                                    dates: List[date], load_fc_full_store: np.ndarray,
                                    load: np.ndarray, pv: np.ndarray, cfg, run_D: int):
    """事后预计算历史误差数组；仿真调用时仍严格只截取d-2及更早历史。"""
    out = {h: np.full((run_D, cfg.hori_slots), np.nan, dtype=float)
           for h in extra_hours}
    for d in range(run_D):
        for h in extra_hours:
            la = actual_horizon(d, h, load, cfg)
            pa = actual_horizon(d, h, pv, cfg)
            if la is None or pa is None:
                continue
            lf = q3_global.load_stage_forecast_horizon(load_fc_full_store[d], load[d], h, cfg)
            pf = pv_record_horizon_10min(d, h, rec_map, dates, pv, cfg)
            out[h][d] = (la - pa) - (lf - pf)
    return out
def extra_stage_reserve(day_idx: int, stage_hour: int, hist_by_hour: dict,
                        fallback_reserve_day: np.ndarray, cfg):
    H = cfg.hori_slots
    fb = np.asarray(fallback_reserve_day, dtype=float)
    clock = (stage_hour * 6 + np.arange(H)) % cfg.slots_pday
    fallback = fb[clock]
    # 与Q3_01官方阶段备用相同：仅使用d-2及更早已完整实现的24h预报误差。
    end = max(0, day_idx - 1)
    start = max(0, end - cfg.stage_rhisdays)
    hist = hist_by_hour[stage_hour][start:end]
    if hist.shape[0] < cfg.stage_mindays:
        return fallback.copy()
    out = np.zeros(H, dtype=float)
    age = np.arange(hist.shape[0] - 1, -1, -1)
    day_w = cfg.stage_decay ** age
    for k in range(H):
        vals = hist[:, k]
        mask = np.isfinite(vals)
        if int(np.sum(mask)) < cfg.stage_mindays:
            out[k] = fallback[k]
            continue
        q = q3_global.weighted_quantile(vals[mask], day_w[mask], cfg.stage_quant)
        out[k] = cfg.stage_scale * max(0.0, float(q))
    return np.clip(out, 0.0, 5000.0)
def simulate_scheme(name: str,
                    update_hours: Tuple[int, ...],
                    records: List[dict],
                    dates, price, load, pv,
                    load_fc_full_store,
                    fallback_reserve_store,
                    official_stage_error_hist,
                    extra_stage_error_hist,
                    run_D: int, cfg, lp):
    n = cfg.slots_pday
    official_idx = {0: 0, 6: 1, 12: 2, 18: 3}
    update_hours = tuple(sorted(set(int(h) for h in update_hours)))
    rec_map = records_by_key(records)
    plan_cost = np.zeros(run_D); adj_cost = np.zeros(run_D); emg_cost = np.zeros(run_D)
    emg_energy = np.zeros(run_D); surplus_energy = np.zeros(run_D); up_energy = np.zeros(run_D)
    down_energy = np.zeros(run_D); soc_start = np.zeros(run_D); soc_end = np.zeros(run_D)
    soc = cfg.e0
    for d in range(run_D):
        soc_start[d] = soc
        lf_full = load_fc_full_store[d]
        lf0 = lf_full[:n]
        pf0 = pv_record_horizon_10min(d, 0, rec_map, dates, pv, cfg)[:n]
        fallback_reserve = fallback_reserve_store[d]
        da = q3_global.solve_day_ahead(price, lf0, pf0, fallback_reserve, soc, cfg, lp)
        initial_buy = da["buy"].copy(); current_buy = da["buy"].copy()
        current_ch = da["ch"].copy(); current_dis = da["dis"].copy()
        plan_cost[d] = float(np.sum(initial_buy * price) * cfg.dt)
        stages = (0,) + update_hours
        for si, h in enumerate(stages):
            start = h * 6
            if h > 0:
                lf_h = q3_global.load_stage_forecast_horizon(lf_full, load[d], h, cfg)
                pf_h = pv_record_horizon_10min(d, h, rec_map, dates, pv, cfg)
                if h in official_idx:
                    reserve_h = q3_global.stage_reserve_from_history(
                        d, official_idx[h], official_stage_error_hist,
                        fallback_reserve, h, cfg
                    )
                else:
                    reserve_h = extra_stage_reserve(
                        d, h, extra_stage_error_hist,
                        fallback_reserve, cfg
                    )
                adj = q3_global.solve_rolling_mpc_24h(
                    price, lf_h, pf_h, reserve_h,
                    soc, current_buy, start, cfg
                )
                m = n - start
                current_buy[start:] = adj["buy"][:m]; current_ch[start:] = adj["ch"][:m]
                current_dis[start:] = adj["dis"][:m]; adj_cost[d] += adj["adj_net_cost"]
                up_energy[d] += adj["up_energy"]; down_energy[d] += adj["down_energy"]
            end = stages[si + 1] * 6 if si + 1 < len(stages) else n
            for t in range(start, end):
                sol = q3_global.operate_one_slot(
                    soc, current_buy[t], load[d, t], pv[d, t],
                    current_ch[t], current_dis[t], cfg
                )
                soc = sol["E"]; emg_energy[d] += sol["emg"] * cfg.dt
                surplus_energy[d] += sol["sur"] * cfg.dt
                emg_cost[d] += sol["emg"] * 5.0 * price[t] * cfg.dt
        soc_end[d] = soc
    return {
        "plan_cost": plan_cost, "adj_cost": adj_cost, "emg_cost": emg_cost,
        "emg_energy": emg_energy, "surplus_energy": surplus_energy, "up_energy": up_energy,
        "down_energy": down_energy, "soc_end": soc_end,
    }
def summarize_scheme(name: str, hours: Tuple[int, ...], result: dict, run_D: int):
    start = 31 if run_D >= 32 else 0
    sl = slice(start, run_D)
    daily = np.asarray(result["emg_energy"][sl], dtype=float)
    plan = float(np.sum(result["plan_cost"][sl]))
    adj = float(np.sum(result["adj_cost"][sl]))
    emgc = float(np.sum(result["emg_cost"][sl]))
    return {
        "方案": name, "全部预报时点": ",".join(["0"] + [str(h) for h in hours]),
        "初始计划购电费": plan, "调整净费用": adj, "紧急购电费": emgc,
        "总费用": plan + adj + emgc, "紧急购电量kWh": float(np.sum(daily)),
        "日紧急购电量P95kWh": float(np.quantile(daily, 0.95)) if len(daily) else 0.0,
        "单日最大紧急购电量kWh": float(np.max(daily)) if len(daily) else 0.0,
        "紧急购电天数": int(np.sum(daily > 1e-6)),
        "富余电量kWh": float(np.sum(result["surplus_energy"][sl])),
        "向上调整电量kWh": float(np.sum(result["up_energy"][sl])),
        "向下调整电量kWh": float(np.sum(result["down_energy"][sl])),
        "期末SOCkWh": float(result["soc_end"][run_D - 1]),
    }
def plot_lp_comparison(rows: List[dict], out_dir: Path):
    ensure_dir(out_dir)
    names = [r["方案"] for r in rows]; x = np.arange(len(rows))
    # 01/02：单指标柱状图（01为相对基线变化），柔色+黑描边、无标题、大字号
    for key, scale, color, ylabel, title, fname, fmt, diff in (
            ("总费用", 1e4, SOFT_BLUE, "相对基线总费用变化 / 万元",
             "新增预报时点对总费用的影响", "第三问新增预报时点对总费用的影响.png", "{:+.2f}", True),
            ("紧急购电量kWh", 1e3, SOFT_RED, "紧急购电量 / MWh",
             "新增预报时点对紧急购电量的影响", "第三问新增预报时点对紧急购电量的影响.png", "{:.2f}", False)):
        vals = np.asarray([r[key] - rows[0][key] if diff else r[key] for r in rows]) / scale
        fig, ax = plt.subplots(figsize=(10.8, 5.4), dpi=300)
        ax.bar(x, vals, width=0.40, color=color, alpha=0.95, zorder=3, **BAR_EDGE)
        if diff:
            ax.axhline(0, lw=1, color='black')
        ax.set_xticks(x, names)
        if not NO_TITLE:
            ax.set_title(title, fontsize=13, fontweight="bold")
        _set_labels(ax, '预报时点方案', ylabel)
        style(ax)
        for i, v in enumerate(vals):
            dy = (0.15 if v >= 0 else -0.15) if diff else 0.0
            ax.text(i, v + dy, fmt.format(v), ha="center",
                    va="bottom" if (not diff or v >= 0) else "top", fontsize=11)
        fig.tight_layout()
        fig.savefig(out_dir / fname, dpi=300, bbox_inches="tight")
        plt.close(fig)
    # 03 P95与最大单日紧急购电图已按需删除（数据仍随 CSV 输出）
q3_global = None
def main():
    global q3_global
    ensure_dir(RunConfig.OUT); ensure_dir(RunConfig.FIG)
    q3_path = first_existing(RunConfig.Q3_MAIN_CANDIDATES)
    att1 = first_existing(RunConfig.ATT1_CANDIDATES)
    att2 = first_existing(RunConfig.ATT2_CANDIDATES)
    att3 = first_existing(RunConfig.ATT3_CANDIDATES)
    cp_path = first_existing(RunConfig.CHECKPOINT_CANDIDATES)
    q3 = import_py(q3_path, "q3_main_model"); q3_global = q3; cfg = q3.CFG
    cfg.base_dir = str(RunConfig.BASE); cfg.att1 = str(att1); cfg.att2 = str(att2); cfg.att3 = str(att3)
    cfg.max_days = RunConfig.MAX_DAYS; cfg.use_bias = False
    dates, price, load_prior, pv_prior, load, pv = load_core_inputs(att1, att2, q3)
    run_D = min(len(dates), RunConfig.MAX_DAYS)
    dates_run = dates[:run_D]; load_run = load[:run_D]; pv_run = pv[:run_D]
    checkpoint = load_checkpoint(cp_path, run_D)
    fallback_reserve_store = np.asarray(checkpoint["fallback_reserve"][:run_D], dtype=float)
    official_stage_error_hist = np.asarray(checkpoint["stage_net_error_hist"][:run_D], dtype=float)
    price_hourly = price.reshape(24, 6).mean(axis=1)
    hour_idx = [6 * h - 1 for h in range(1, 25)]
    load_hourly = load_run[:, hour_idx]; pv_hourly_actual = pv_run[:, hour_idx]
    soc_hourly = reconstruct_soc_hourly(checkpoint, run_D, cfg)
    official_records_all = load_forecast_records_xlsx(att3)
    official_records = [r for r in official_records_all if r["issue_dt"].date() in set(dates_run)]
    target_times = build_target_times(dates_run)
    fc0 = effective_forecast_records(official_records, target_times)
    base_risk = analyze_risk_scheme("原四时点预报", dates_run, load_hourly, pv_hourly_actual,
                                    price_hourly, soc_hourly, fc0, cfg)
    write_csv(RunConfig.OUT / "风险带_原方案按目标小时汇总.csv", base_risk["hourly"])
    extra_hours, _ = select_extra_hours(base_risk)
    print("\n风险带筛选")
    if not extra_hours:
        raise RuntimeError("未筛选出新增预报时点，无法继续增强方案验证。")
    print("最终新增时点：", ", ".join(f"{h}:00" for h in extra_hours))
    augmented_records = build_augmented_records(
        official_records, extra_hours, dates_run, pv_hourly_actual
    )
    extra_tag = "_".join(str(h) for h in extra_hours)
    enhanced_xlsx = RunConfig.OUT / f"附件3新增预报_{extra_tag}_结果.xlsx"
    export_records_xlsx(enhanced_xlsx, augmented_records)
    fc1 = effective_forecast_records(augmented_records, target_times)
    enhanced_risk = analyze_risk_scheme("增加预报后", dates_run, load_hourly, pv_hourly_actual,
                                        price_hourly, soc_hourly, fc1, cfg)
    write_csv(RunConfig.OUT / "风险带_增强方案按目标小时汇总.csv", enhanced_risk["hourly"])
    risk_cmp = risk_compare_table(base_risk, enhanced_risk)
    write_csv(RunConfig.OUT / "风险带_原方案vs增强方案.csv", risk_cmp)
    for dd in RunConfig.REPRESENTATIVE_DATES:
        if dd in dates_run:
            plot_risk_day_compare(
                base_risk, enhanced_risk, dd.isoformat(),
                RunConfig.FIG / f"第三问{dd.month}-{dd.day}动态非对称风险带前后对比.png"
            )
    # 各目标小时带外量对比图已按需删除
    plot_risk_overall(
        base_risk, enhanced_risk,
        RunConfig.FIG / "第三问增强方案相对原方案的风险指标变化.png"
    )
    load_fc_full_store = build_or_load_preprocess_cache(
        q3, cfg, dates_run, load_run, load_prior, run_D)
    aug_map = records_by_key(augmented_records)
    extra_stage_hist = build_extra_stage_error_history(
        extra_hours, aug_map, dates_run, load_fc_full_store, load_run, pv_run, cfg, run_D)
    lp = q3.build_lp_templates(price, cfg)
    base_updates = (6, 12, 18)
    schemes = [("原0/6/12/18", base_updates)]
    for h in extra_hours:
        hrs = tuple(sorted(set(base_updates + (h,))))
        schemes.append((f"原方案+{h}:00", hrs))
    if len(extra_hours) >= 2:
        all_hrs = tuple(sorted(set(base_updates + tuple(extra_hours))))
        schemes.append(("原方案+全部新增时点", all_hrs))
    schemes = list(dict.fromkeys(schemes))
    summaries = []
    print("\n完整24h滚动LP验证")
    for name, hrs in schemes:
        records = official_records if hrs == base_updates else augmented_records
        r = simulate_scheme(
            name, hrs, records, dates_run, price, load_run, pv_run,
            load_fc_full_store, fallback_reserve_store, official_stage_error_hist,
            extra_stage_hist, run_D, cfg, lp)
        summaries.append(summarize_scheme(name, hrs, r, run_D))
    write_csv(RunConfig.OUT / "新增预报时点_完整滚动LP方案对比.csv", summaries)
    plot_lp_comparison(summaries, RunConfig.FIG)
    base_row = summaries[0]
    delta_rows = []
    metrics = ["总费用", "紧急购电量kWh", "日紧急购电量P95kWh",
               "单日最大紧急购电量kWh", "富余电量kWh", "向上调整电量kWh", "向下调整电量kWh"]
    for row in summaries[1:]:
        item = {"方案": row["方案"]}
        for m in metrics:
            b = float(base_row[m])
            v = float(row[m])
            item[f"{m}_变化量"] = v - b
            item[f"{m}_变化率%"] = (v - b) / b * 100.0 if abs(b) > 1e-12 else np.nan
        delta_rows.append(item)
    write_csv(RunConfig.OUT / "新增预报方案_相对基线改善.csv", delta_rows)
    print("\n最终方案比较")
    for r in summaries:
        print(f"{r['方案']}: 总费用={r['总费用']:.2f} 元 | "
            f"紧急购电={r['紧急购电量kWh']:.2f} kWh | "
            f"P95={r['日紧急购电量P95kWh']:.2f} kWh | "
            f"最大单日={r['单日最大紧急购电量kWh']:.2f} kWh | "
            f"富余={r['富余电量kWh']:.2f} kWh")
    min_cost = min(summaries, key=lambda r: r["总费用"])
    min_emg = min(summaries, key=lambda r: r["紧急购电量kWh"])
    conclusion = ["问题3新增预报时点结果整合", "=" * 70,
                  "风险带筛选新增时点：" + ", ".join(f"{h}:00" for h in extra_hours),
                  f"最低总费用方案：{min_cost['方案']}，{min_cost['总费用']:.2f} 元",
                  f"最低紧急购电方案：{min_emg['方案']}，{min_emg['紧急购电量kWh']:.2f} kWh" ]
    (RunConfig.OUT / "结果整合.txt").write_text("\n".join(conclusion), encoding="utf-8")
    print("\n输出目录：", RunConfig.OUT)
    print("方案对比：", RunConfig.OUT / "新增预报时点_完整滚动LP方案对比.csv")
if __name__ == "__main__":
    main()