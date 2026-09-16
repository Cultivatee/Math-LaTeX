from __future__ import annotations
import csv
import math
import os
import sys
import importlib.util
from pathlib import Path
from datetime import datetime, date, timedelta
from typing import List, Dict, Tuple, Optional
import numpy as np
from openpyxl import load_workbook, Workbook
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei",
    "Noto Sans CJK SC", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
# 全局绘图样式
plt.rcParams["font.size"] = 12
plt.rcParams["axes.spines.top"] = False
plt.rcParams["axes.linewidth"] = 1.5
plt.rcParams["xtick.major.width"] = 1.5
plt.rcParams["ytick.major.width"] = 1.5
plt.rcParams["xtick.major.size"] = 6
plt.rcParams["ytick.major.size"] = 6
# 配色
C_BUY="#2E74B5";C_PV="#27AE60";C_DIS="#E67E22";C_LOAD="#34495E"
C_CH="#E74C3C";C_W="#95A5A6";C_SOC="#7F8C8D";C_MIN="#2980B9"
C_RES="#16A085";C_ALPHA="#8E44AD"
class RunConfig:
    BASE=Path(__file__).resolve().parent if "__file__" in globals() else Path(".")
    Q43_CANDIDATES = [BASE / "Q4-3.py"]
    ATT3_CANDIDATES = [BASE / "附件3.xlsx", BASE / "附件3(1).xlsx"]
    CHECKPOINT_CANDIDATES = [
        BASE / "Q4-3结果" / "q4_3_checkpoint.npz",
        BASE / "q4_3_outputs_price_forecast" / "q4_3_checkpoint.npz",
        BASE / "q4_3_checkpoint.npz",
        BASE / "_q43_test_outputs" / "q4_3_checkpoint.npz",
    ]
    OUT = BASE / "Q4-3-2结果"
    FIG = OUT / "Q4-3-2图片输出"
    EXISTING_HOURS = (0, 6, 12, 18)
    POSSIBLE_EXTRA_HOURS=tuple(h for h in range(1,24) if h not in (0,6,12,18))
    LOOKAHEAD_HOURS = 4;LOOKAHEAD_DECAY_TAU = 1.5
    MAX_AUTO_CANDIDATES = 7;MIN_AUTO_CANDIDATES = 4
    MIN_CANDIDATE_GAP_HOURS = 2;MANUAL_EXTRA_HOURS = None
    RUN_ALL_CANDIDATES_TOGETHER = True
    RECENT_ERROR_HOURS = 3;BIAS_DECAY_HOURS = 4.0
    HISTORY_DAYS_FOR_TAIL = 30;MAX_BIAS_ABS_KW = 1500.0
    MAX_DAYS = int(os.environ.get("Q43_EXTRA_MAX_DAYS", "365"))
    REPRESENTATIVE_DATES = (
        date(2025, 3, 20), date(2025, 6, 21),
        date(2025, 9, 23), date(2025, 12, 21),
    )
class RiskBandConfig:
    DAYLIGHT_THRESHOLD = 100.0; BAND_QUANTILE = 0.85
    BAND_MIN_KW = 120.0; BAND_FLOOR_RATIO = 0.08; USE_HOUR_SMOOTH = True
    C_ES = 0.03; C_SUR = 0.06; K_SHORT = 5.0
    K_DIS = 1.5; K_CH = 0.5; EPS = 1e-9
def first_existing(paths: List[Path]) -> Path:
    for p in paths:
        if p.exists():
            return p
    raise FileNotFoundError("未找到文件：\n" + "\n".join(str(p) for p in paths))
def ensure_dir(p: Path):
    p.mkdir(parents=True, exist_ok=True)
def import_py(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, str(path))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载：{path}")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod
def write_csv(path: Path, rows: List[dict]):
    if not rows:
        return
    ensure_dir(path.parent)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)
def style(ax):
    ax.grid(True, axis="y", linestyle="--", linewidth=0.7, alpha=0.3)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(labelsize=TICK_FS)
NO_TITLE  = True
AXIS_FS   = 15
TICK_FS   = 12
LEGEND_FS = 17
SOFT_RED,SOFT_GREEN,SOFT_BLUE,SOFT_GRAY='#E9A6A1','#8BCF8B','#7097CA','#CFCECF'
BAR_EDGE  = dict(edgecolor='black', linewidth=1.2)
def _set_labels(ax, x, y):
    ax.set_xlabel(x, fontsize=AXIS_FS, color='black')
    ax.set_ylabel(y, fontsize=AXIS_FS, color='black')
def _legend_top(ax, ncol=None, fs=LEGEND_FS):
    h, l = ax.get_legend_handles_labels()
    if h:
        ax.legend(h, l, loc='lower center', frameon=False, ncol=ncol or len(h),
                  fontsize=fs, bbox_to_anchor=(0.5, 1.02), columnspacing=1.6)
def load_checkpoint(path: Path, run_D: int):
    z = np.load(path, allow_pickle=True)
    need = ["soc_start", "ch", "dis", "effective_buy", "load_fc", "price_fc",
            "fallback_reserve", "stage_net_error_hist"]
    miss = [k for k in need if k not in z.files]
    if miss:
        raise ValueError("Q4-3 checkpoint缺少字段：" + ", ".join(miss))
    out = {k: np.asarray(z[k]) for k in z.files}
    if len(out["soc_start"]) < run_D:
        raise ValueError(f"checkpoint只有{len(out['soc_start'])}天，"
            f"要求{run_D}天。请先完整运行Q4-3。")
    return out
def reconstruct_soc_hourly(cp: dict, run_D: int, cfg):
    soc_start = np.asarray(cp["soc_start"][:run_D], float)
    ch = np.asarray(cp["ch"][:run_D], float)
    dis = np.asarray(cp["dis"][:run_D], float)
    out = np.zeros((run_D, 24), float)
    for d in range(run_D):
        soc = float(soc_start[d])
        for h in range(24):
            out[d, h] = soc
            for k in range(h*6, h*6+6):
                soc += cfg.eta_c*ch[d,k]*cfg.dt - dis[d,k]*cfg.dt/cfg.eta_d
                soc = min(cfg.e_max, max(cfg.e_min, soc))
    return out
def load_forecast_records_xlsx(path: Path, kind="official"):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    cur = None; recs=[]
    for r in ws.iter_rows(min_row=2, max_col=26, values_only=True):
        if r[0] not in (None, ""):
            x=r[0]
            if isinstance(x, datetime): cur=x.date()
            elif isinstance(x, date): cur=x
            else: cur=datetime.strptime(str(x).replace('/','-')[:10], '%Y-%m-%d').date()
        if cur is None: continue
        txt=str(r[1]).strip(); h=int(txt.split(':')[0]) if ':' in txt else int(float(txt))
        fc=np.asarray([0.0 if v is None else float(v) for v in r[2:26]],float)
        recs.append({"issue_dt":datetime.combine(cur,datetime.min.time())+timedelta(hours=h),
                     "issue_hour":h,"fc":np.maximum(fc,0.0),"kind":kind})
    wb.close(); recs.sort(key=lambda z:z["issue_dt"])
    return recs
def records_by_key(records):
    return {(r["issue_dt"].date(), int(r["issue_hour"])):r for r in records}
def latest_record_before(records, issue_dt):
    c=[r for r in records if r["issue_dt"] < issue_dt]
    return max(c,key=lambda z:z["issue_dt"]) if c else None
def record_value_at_target(rec, target_dt):
    lead=int(round((target_dt-rec["issue_dt"]).total_seconds()/3600.0))
    if 1 <= lead <= 24:
        return float(rec["fc"][lead-1]), lead
    return None, None
def actual_hour_value(target_dt, dates, actual_hourly):
    dmap={d:i for i,d in enumerate(dates)}
    dd=target_dt.date(); h=target_dt.hour
    if h==0:
        dd=dd-timedelta(days=1); hh=23
    else: hh=h-1
    if dd not in dmap: return None
    return float(actual_hourly[dmap[dd],hh])
def past_climatology(target_dt, dates, actual_hourly, issue_dt):
    dmap={d:i for i,d in enumerate(dates)}; vals=[]
    dd=target_dt.date(); h=target_dt.hour
    if h==0: dd=dd-timedelta(days=1); hh=23
    else: hh=h-1
    if dd not in dmap: return 0.0
    i=dmap[dd]
    for j in range(max(0,i-RunConfig.HISTORY_DAYS_FOR_TAIL),i):
        vals.append(actual_hourly[j,hh])
    return float(np.median(vals)) if vals else 0.0
def build_extra_record(extra_issue_dt, official_records, dates, actual_hourly):
    old=latest_record_before(official_records,extra_issue_dt)
    if old is None: return None
    residuals=[]; weights=[]
    for back in range(RunConfig.RECENT_ERROR_HOURS-1,-1,-1):
        t=extra_issue_dt-timedelta(hours=back)
        ov,_=record_value_at_target(old,t); av=actual_hour_value(t,dates,actual_hourly)
        if ov is None or av is None: continue
        residuals.append(av-ov); weights.append(math.exp(-back/1.5))
    if residuals:
        w=np.asarray(weights,float); w/=w.sum()
        bias=float(np.dot(w,np.asarray(residuals,float)))
    else: bias=0.0
    bias=float(np.clip(bias,-RunConfig.MAX_BIAS_ABS_KW,RunConfig.MAX_BIAS_ABS_KW))
    fc=np.zeros(24,float)
    for lead in range(1,25):
        t=extra_issue_dt+timedelta(hours=lead)
        base,_=record_value_at_target(old,t)
        if base is None: base=past_climatology(t,dates,actual_hourly,extra_issue_dt)
        corr=bias*math.exp(-(lead-1)/RunConfig.BIAS_DECAY_HOURS)
        if base < 100.0: corr*=0.25
        fc[lead-1]=max(base+corr,0.0)
    return {"issue_dt":extra_issue_dt,"issue_hour":extra_issue_dt.hour,"fc":fc,
            "kind":"extra_proxy","bias":bias,"source_issue_hour":old["issue_hour"]}
def build_augmented_records(official_records, extra_hours, dates, actual_hourly):
    out=[dict(r) for r in official_records]; diag=[]
    days=sorted({r["issue_dt"].date() for r in official_records})
    for dd in days:
        for h in extra_hours:
            issue=datetime.combine(dd,datetime.min.time())+timedelta(hours=h)
            rec=build_extra_record(issue,official_records,dates,actual_hourly)
            if rec is None: continue
            out.append(rec); diag.append({"日期":dd.isoformat(),"新增发布时间":h,
                "上一官方发布时间":rec["source_issue_hour"],
                "截至发布时间估计偏差bias_kW":round(rec["bias"],4)})
    out.sort(key=lambda z:z["issue_dt"])
    return out,diag
def export_records_xlsx(path, records):
    wb=Workbook(); ws=wb.active; ws.title="Sheet1"
    ws.append(["日期","预报时刻"]+[f"预报{i}小时" for i in range(1,25)])
    by={}
    for r in records: by.setdefault(r["issue_dt"].date(),[]).append(r)
    for dd in sorted(by):
        rr=sorted(by[dd],key=lambda z:z["issue_dt"])
        for j,r in enumerate(rr):
            ws.append([dd.isoformat() if j==0 else None,f"{r['issue_hour']}:00"]
                +[float(v) for v in r["fc"]])
    ensure_dir(path.parent); wb.save(path); wb.close()
def pv_record_horizon_10min(day_idx, stage_hour, rec_map, dates, pv_actual, cfg):
    rec=rec_map[(dates[day_idx],int(stage_hour))]
    hourly=np.asarray(rec["fc"],float); start=stage_hour*6
    anchor=0.0 if stage_hour==0 else float(pv_actual[day_idx,start-1])
    xp=np.arange(25,dtype=float); fp=np.concatenate([[anchor],hourly])
    target=(np.arange(cfg.rolling_horizon_slots)+1)/6.0
    return np.maximum(np.interp(target,xp,fp),0.0)
def effective_pv_hourly(records, dates):
    """得到每天1..24点目标时刻的最后可用PV预报及其发布时间。"""
    vals=np.full((len(dates),24),np.nan); issue=np.full((len(dates),24),-1,int)
    for di,dd in enumerate(dates):
        d0=datetime.combine(dd,datetime.min.time())
        for hh in range(1,25):
            t=d0+timedelta(hours=hh); best=None
            for r in records:
                if r["issue_dt"] >= t: break
                v,lead=record_value_at_target(r,t)
                if v is not None and (best is None or r["issue_dt"]>best[0]):
                    best=(r["issue_dt"],v,r["issue_hour"])
            if best is not None:
                vals[di,hh-1]=best[1]; issue[di,hh-1]=best[2]
    if np.isnan(vals).any(): raise ValueError("存在目标小时缺少PV预报")
    return vals,issue
def risk_economic_costs(price, load, fc_pv, soc, cfg):
    # 与Q3风险带一致，但直接使用Q4实际波动电价。
    a_dis=max(0.0,min(cfg.p_max,cfg.eta_d*max(soc-cfg.e_min,0.0)))
    a_ch=max(0.0,min(cfg.p_max,max(cfg.e_max-soc,0.0)/cfg.eta_c))
    net=load-fc_pv; shortage=max(net,0.0); surplus=max(-net,0.0)
    pi_plus=a_dis/(a_dis+shortage+RiskBandConfig.EPS)
    pi_minus=a_ch/(a_ch+surplus+RiskBandConfig.EPS)
    cp=RiskBandConfig.C_ES+price*(RiskBandConfig.K_DIS*pi_plus+RiskBandConfig.K_SHORT*(1-pi_plus))
    cm=RiskBandConfig.C_ES+RiskBandConfig.K_CH*price*pi_minus+RiskBandConfig.C_SUR*(1-pi_minus)
    return cp,cm
def build_hour_scale(actual,fc):
    err=fc-actual; raw=np.zeros(24); base=np.zeros(24)
    for h in range(24):
        m=((actual[:,h]>=RiskBandConfig.DAYLIGHT_THRESHOLD)
           |(fc[:,h]>=RiskBandConfig.DAYLIGHT_THRESHOLD))
        if np.any(m):
            q=float(np.quantile(np.abs(err[m,h]),RiskBandConfig.BAND_QUANTILE))
            med=float(np.median(actual[m,h]))
            floor=max(RiskBandConfig.BAND_FLOOR_RATIO*med,
                RiskBandConfig.BAND_MIN_KW)
            raw[h]=q; base[h]=max(q,floor)
    if RiskBandConfig.USE_HOUR_SMOOTH:
        sm=base.copy()
        for h in range(24): sm[h]=.25*base[max(h-1,0)]+.5*base[h]+.25*base[min(h+1,23)]
        base=sm
    return raw,base
def build_effective_price_hourly(q43, price_fc_full, price_actual, cfg, run_D):
    """按原0/6/12/18方案，给每个目标整点配置当时最后一次可用的价格预测。"""
    pred=np.full((run_D,24),np.nan); issue=np.full((run_D,24),-1,int)
    stages=(0,6,12,18)
    for d in range(run_D):
        for h0 in stages:
            ph=q43.price_stage_forecast_horizon(price_fc_full[d],price_actual[d],h0,cfg)
            # ph[0]对应发布后第一个10分钟；整点目标 h= h0+1...，取每小时末点位置 5,11,...
            next_stage={0:6,6:12,12:18,18:24}[h0]
            for target_h in range(h0+1,next_stage+1):
                j=(target_h-h0)*6-1
                pred[d,target_h-1]=ph[j]; issue[d,target_h-1]=h0
    return pred,issue
def build_combined_hourly_risk(q43, dates, load, pv, price_actual, cp, cfg,
    run_D, official_records):
    # 小时样本取每小时末10分钟点，与附件3整点目标保持一致。
    idx=np.array([6*h-1 for h in range(1,25)])
    load_h=load[:run_D][:,idx]; pv_h=pv[:run_D][:,idx]; price_h=price_actual[:run_D][:,idx]
    soc_h=reconstruct_soc_hourly(cp,run_D,cfg)
    buy_h=np.zeros((run_D,24));
    for h in range(24): buy_h[:,h]=np.sum(cp["effective_buy"][:run_D,h*6:(h+1)*6],axis=1)*cfg.dt
    pv_fc,pv_issue=effective_pv_hourly(official_records,dates[:run_D])
    _,scale=build_hour_scale(pv_h,pv_fc)
    pfc_h,p_issue=build_effective_price_hourly(q43,cp["price_fc"][:run_D],
        price_actual[:run_D],cfg,run_D)
    rows=[]; agg=[]
    for h in range(24):
        pv_econ=[]; price_econ=[]; danger=[]
        for d in range(run_D):
            cpv,cmv=risk_economic_costs(price_h[d,h],load_h[d,h],pv_fc[d,h],soc_h[d,h],cfg)
            den=max(cpv+cmv,RiskBandConfig.EPS)
            bplus=2*cmv/den*scale[h]
            d_ex=max(pv_fc[d,h]-(pv_h[d,h]+bplus),0.0)
            # 危险高估近似造成的1小时缺口经济暴露（元）
            pv_loss=d_ex*cpv
            # 价格预测误差 × 该小时基线购电电量（元）
            price_loss=abs(pfc_h[d,h]-price_h[d,h])*buy_h[d,h]
            pv_econ.append(pv_loss); price_econ.append(price_loss); danger.append(d_ex)
        sp=float(np.sum(pv_econ)); sl=float(np.sum(price_econ)); sc=sp+sl
        agg.append({"目标小时":h+1,"PV危险高估经济风险元":sp,"价格预测经济暴露元":sl,
                    "综合信息风险元":sc,"PV危险高估总量kW":float(np.sum(danger)),
                    "平均实际电价":float(np.mean(price_h[:,h])),
                    "平均绝对价格误差":float(np.mean(np.abs(pfc_h[:,h]-price_h[:,h]))),
                    "平均基线购电量kWh":float(np.mean(buy_h[:,h]))})
    return agg,pv_fc,pfc_h
def _future_coverage_score(hourly_rows, issue_hour):
    stats = {int(r["目标小时"]): float(r["综合信息风险元"]) for r in hourly_rows}
    score = 0.0
    covered = []
    for k in range(1, RunConfig.LOOKAHEAD_HOURS + 1):
        target = issue_hour + k
        if target > 24:
            break
        w = math.exp(-(k - 1) / RunConfig.LOOKAHEAD_DECAY_TAU)
        val = stats.get(target, 0.0); score += w * val
        covered.append((target, w, val))
    return float(score), covered
def build_candidate_score_rows(hourly_rows):
    raw_risk = {int(r["目标小时"]): float(r["综合信息风险元"]) for r in hourly_rows}
    rows = []
    for h in RunConfig.POSSIBLE_EXTRA_HOURS:
        score, covered = _future_coverage_score(hourly_rows, h)
        rows.append({
            "候选发布时间": int(h),
            "未来风险覆盖得分元": float(score),
            "覆盖目标小时": ",".join(str(x[0]) for x in covered),
            "该时点自身综合风险元": float(raw_risk.get(h, 0.0)),
            "是否局部峰值": 0,
            "是否最终候选": 0,
            "候选排名": "",
        })
    smap = {int(r["候选发布时间"]): float(r["未来风险覆盖得分元"]) for r in rows}
    possible = list(RunConfig.POSSIBLE_EXTRA_HOURS)
    for r in rows:
        h = int(r["候选发布时间"])
        s = smap[h]
        # 在“可选发布时间序列”上判断局部峰值，避免 6/12/18 的空缺导致误判。
        pos = possible.index(h)
        left = smap[possible[pos-1]] if pos > 0 else -np.inf
        right = smap[possible[pos+1]] if pos < len(possible)-1 else -np.inf
        if s >= left and s >= right and s > 0:
            r["是否局部峰值"] = 1
    return rows
def select_extra_hours(hourly_rows):
    if RunConfig.MANUAL_EXTRA_HOURS is not None:
        manual = sorted(set(int(x) for x in RunConfig.MANUAL_EXTRA_HOURS))
        rows = build_candidate_score_rows(hourly_rows)
        for r in rows:
            if int(r["候选发布时间"]) in manual:
                r["是否最终候选"] = 1
        return manual, rows
    rows = build_candidate_score_rows(hourly_rows)
    local = [r for r in rows if int(r["是否局部峰值"]) == 1]
    local.sort(key=lambda z: float(z["未来风险覆盖得分元"]), reverse=True)
    selected = []
    def far_enough(h):
        return all(abs(h - x) >= RunConfig.MIN_CANDIDATE_GAP_HOURS for x in selected)
    # 第一轮：优先选择局部峰值。
    for r in local:
        h = int(r["候选发布时间"])
        if far_enough(h):
            selected.append(h)
        if len(selected) >= RunConfig.MAX_AUTO_CANDIDATES:
            break
    # 第二轮：若局部峰值太少，则按全局评分补足，但仍保持最小时距。
    if len(selected) < RunConfig.MIN_AUTO_CANDIDATES:
        all_rows = sorted(rows, key=lambda z: float(z["未来风险覆盖得分元"]), reverse=True)
        for r in all_rows:
            h = int(r["候选发布时间"])
            if h in selected:
                continue
            if far_enough(h):
                selected.append(h)
            if len(selected) >= RunConfig.MIN_AUTO_CANDIDATES:
                break
    selected = sorted(selected)
    rank_order = sorted(selected, key=lambda h: next(
        float(r["未来风险覆盖得分元"]) for r in rows if int(r["候选发布时间"]) == h
    ), reverse=True)
    rank_map = {h: i+1 for i, h in enumerate(rank_order)}
    for r in rows:
        h = int(r["候选发布时间"])
        if h in selected:
            r["是否最终候选"] = 1
            r["候选排名"] = rank_map[h]
    return selected, rows
def actual_horizon(day_idx,stage_hour,arr,cfg):
    H=cfg.rolling_horizon_slots; flat=arr.reshape(-1); g0=day_idx*cfg.slots_per_day+stage_hour*6
    if g0+H>len(flat): return None
    return np.asarray(flat[g0:g0+H],float)
def build_extra_stage_error_history(q43,extra_hours,rec_map,dates,load_fc,load,pv,cfg,run_D):
    out={h:np.full((run_D,cfg.rolling_horizon_slots),np.nan) for h in extra_hours}
    for d in range(run_D):
        for h in extra_hours:
            la=actual_horizon(d,h,load,cfg); pa=actual_horizon(d,h,pv,cfg)
            if la is None or pa is None:continue
            lf=q43.load_stage_forecast_horizon(load_fc[d],load[d],h,cfg)
            pf=pv_record_horizon_10min(d,h,rec_map,dates,pv,cfg)
            out[h][d]=(la-pa)-(lf-pf)
    return out
def extra_stage_reserve(q43,day_idx,h,hist_by_hour,fallback_day,cfg):
    H=cfg.rolling_horizon_slots; fb=np.asarray(fallback_day,float)
    clock=(h*6+np.arange(H))%cfg.slots_per_day
    fallback=fb[clock]; end=max(0,day_idx-1)
    start=max(0,end-cfg.stage_reserve_history_days)
    hist=hist_by_hour[h][start:end]
    if hist.shape[0] < cfg.stage_reserve_min_days:return fallback.copy()
    out=np.zeros(H); age=np.arange(hist.shape[0]-1,-1,-1); w=cfg.stage_reserve_decay**age
    for k in range(H):
        vals=hist[:,k]; m=np.isfinite(vals)
        if np.sum(m)<cfg.stage_reserve_min_days:out[k]=fallback[k]
        else:out[k]=cfg.stage_reserve_scale*max(0.0,float(
            q43.weighted_quantile(vals[m],w[m],cfg.stage_reserve_quantile)))
    return np.clip(out,0,5000)
def simulate_scheme(q43,name,update_hours,records,dates,price_actual,load,pv,load_fc,price_fc,
                    fallback_store,official_hist,extra_hist,run_D,cfg,lp):
    n=cfg.slots_per_day; official_idx={0:0,6:1,12:2,18:3}; rec_map=records_by_key(records)
    update_hours=tuple(sorted(set(int(h) for h in update_hours)))
    plan=np.zeros(run_D); adjc=np.zeros(run_D); emgc=np.zeros(run_D)
    emge=np.zeros(run_D); sure=np.zeros(run_D)
    upe=np.zeros(run_D); dne=np.zeros(run_D); s0a=np.zeros(run_D)
    s1a=np.zeros(run_D); soc=cfg.e0
    for d in range(run_D):
        s0a[d]=soc; lf_full=load_fc[d]; pr_full=price_fc[d]; fallback=fallback_store[d]
        pf0=pv_record_horizon_10min(d,0,rec_map,dates,pv,cfg)[:n]
        da=q43.solve_day_ahead_price(pr_full,lf_full[:n],pf0,fallback,soc,cfg,lp)
        init=da["buy"].copy(); cb=da["buy"].copy(); cc=da["ch"].copy(); cd=da["dis"].copy()
        plan[d]=float(np.sum(init*price_actual[d])*cfg.dt)
        stages=(0,)+update_hours
        for si,h in enumerate(stages):
            start=h*6
            if h>0:
                lf=q43.load_stage_forecast_horizon(lf_full,load[d],h,cfg)
                pf=pv_record_horizon_10min(d,h,rec_map,dates,pv,cfg)
                if h in official_idx:
                    reserve=q43.stage_reserve_from_history(d,official_idx[h],
                        official_hist,fallback,h,cfg)
                else:
                    reserve=extra_stage_reserve(q43,d,h,extra_hist,fallback,cfg)
                pr=q43.price_stage_forecast_horizon(pr_full,price_actual[d],h,cfg)
                sol=q43.solve_rolling_mpc_24h_price(pr,lf,pf,reserve,soc,cb,start,cfg)
                m=n-start; cb[start:]=sol["buy"][:m]
                cc[start:]=sol["ch"][:m]; cd[start:]=sol["dis"][:m]
                up=sol["up"][:m]; dn=sol["down"][:m]; real=price_actual[d,start:]
                adjc[d]+=float(np.sum((1.5*real*up-0.5*real*dn)*cfg.dt))
                upe[d]+=sol["up_energy"]; dne[d]+=sol["down_energy"]
            end=stages[si+1]*6 if si+1<len(stages) else n
            for t in range(start,end):
                o=q43.operate_one_slot(soc,cb[t],load[d,t],pv[d,t],cc[t],cd[t],cfg)
                soc=o["E"]; emge[d]+=o["emg"]*cfg.dt; sure[d]+=o["sur"]*cfg.dt
                emgc[d]+=o["emg"]*5*price_actual[d,t]*cfg.dt
        s1a[d]=soc
    return {"plan_cost":plan,"adj_cost":adjc,"emg_cost":emgc,"emg_energy":emge,"surplus_energy":sure,
            "up_energy":upe,"down_energy":dne,"soc_start":s0a,"soc_end":s1a}
def summarize(name,hours,r,run_D):
    s=31 if run_D>=32 else 0; sl=slice(s,run_D); daily=np.asarray(r["emg_energy"][sl],float)
    plan=float(np.sum(r["plan_cost"][sl])); adj=float(np.sum(r["adj_cost"][sl]))
    emg=float(np.sum(r["emg_cost"][sl]))
    return {"方案":name,"全部预报时点":",".join(["0"]+[str(h) for h in hours]),
            "初始计划购电费":plan,"调整净费用":adj,
            "紧急购电费":emg,"总费用":plan+adj+emg,"紧急购电量kWh":float(np.sum(daily)),
            "日紧急购电量P95kWh":float(np.quantile(daily,.95)) if len(daily) else 0.0,
            "单日最大紧急购电量kWh":float(np.max(daily)) if len(daily) else 0.0,
            "紧急购电天数":int(np.sum(daily>1e-6)),
            "富余电量kWh":float(np.sum(r["surplus_energy"][sl])),
            "向上调整电量kWh":float(np.sum(r["up_energy"][sl])),
            "向下调整电量kWh":float(np.sum(r["down_energy"][sl])),
            "期末SOCkWh":float(r["soc_end"][run_D-1])}
def build_risk_band_diagnostics(q43, dates, load, pv, price_actual, cp, cfg, run_D, pv_fc_h):
    idx=np.array([6*h-1 for h in range(1,25)])
    load_h=load[:run_D][:,idx]
    pv_h=pv[:run_D][:,idx]
    price_h=price_actual[:run_D][:,idx]
    soc_h=reconstruct_soc_hourly(cp,run_D,cfg)
    _,scale=build_hour_scale(pv_h,pv_fc_h)
    rows=[]
    by_day={}
    for d in range(run_D):
        lo=np.zeros(24); hi=np.zeros(24)
        bp=np.zeros(24); bm=np.zeros(24)
        danger=np.zeros(24,dtype=bool); safe=np.zeros(24,dtype=bool)
        for h in range(24):
            cplus,cminus=risk_economic_costs(
                price_h[d,h],load_h[d,h],pv_fc_h[d,h],soc_h[d,h],cfg
            )
            den=max(cplus+cminus,RiskBandConfig.EPS)
            bplus=2*cminus/den*scale[h]
            bminus=2*cplus/den*scale[h]
            lower=max(0.0,pv_h[d,h]-bminus)
            upper=pv_h[d,h]+bplus
            bp[h]=bplus; bm[h]=bminus
            lo[h]=lower; hi[h]=upper
            danger[h]=pv_fc_h[d,h] > upper
            safe[h]=pv_fc_h[d,h] < lower
            rows.append({
                "日期": str(dates[d]), "目标小时": h+1,
                "实际光伏kW": float(pv_h[d,h]), "生效光伏预报kW": float(pv_fc_h[d,h]),
                "风险带下界kW": float(lower), "风险带上界kW": float(upper),
                "b_minus_kW": float(bminus), "b_plus_kW": float(bplus),
                "危险高估": int(danger[h]), "保守低估": int(safe[h]),
                "电价元每kWh": float(price_h[d,h]), "SOC近似kWh": float(soc_h[d,h]),
                "c_plus_元每kWh": float(cplus), "c_minus_元每kWh": float(cminus),
            })
        by_day[dates[d]]={"actual":pv_h[d].copy(),"forecast":pv_fc_h[d].copy(),
            "lower":lo,"upper":hi,"danger":danger,
            "safe":safe,"b_plus":bp,"b_minus":bm}
    return rows,by_day
def plot_risk_band_single(day, ana, path):
    x=np.arange(1,25)
    ac=ana["actual"]; fc=ana["forecast"]
    lo=ana["lower"]; hi=ana["upper"]
    danger=ana["danger"]; safe=ana["safe"]
    fig,ax=plt.subplots(figsize=(12.6,5.8),dpi=200)
    ax.fill_between(x,lo,hi,alpha=.12,color=C_BUY,label="动态非对称风险带")
    ax.plot(x,ac,lw=2.2,color=C_PV,label="实际光伏")
    ax.plot(x,fc,lw=2.0,color=C_BUY,label="生效光伏预报")
    ax.plot(x,lo,lw=1.0,ls="--",alpha=.75,color=C_SOC,label="风险带下界")
    ax.plot(x,hi,lw=1.0,ls="--",alpha=.75,color=C_SOC,label="风险带上界")
    if np.any(danger):
        ax.scatter(x[danger],fc[danger],s=56,marker="^",color=C_CH,zorder=6,label="危险高估")
    if np.any(safe):
        ax.scatter(x[safe],fc[safe],s=56,marker="v",color=C_DIS,zorder=6,label="保守低估")
    ax.set_xlim(1,24)
    ax.set_xticks(range(1,25))
    ax.set_xlabel("目标小时 / h", fontsize=12)
    ax.set_ylabel("光伏功率 / kW", fontsize=12)
    ax.set_title(f"{day}：动态非对称光伏经济风险带", fontsize=13, fontweight="bold", pad=12)
    ax.legend(frameon=False,ncol=3, fontsize=10)
    style(ax)
    fig.tight_layout()
    fig.savefig(path,bbox_inches="tight")
    plt.close(fig)
def plot_risk_band_summary(rows, path):
    danger=np.zeros(24); safe=np.zeros(24); total=np.zeros(24)
    for r in rows:
        h=int(r["目标小时"])-1
        danger[h]+=int(r["危险高估"])
        safe[h]+=int(r["保守低估"])
        total[h]+=1
    total=np.maximum(total,1); danger_rate=100*danger/total
    safe_rate=100*safe/total; x=np.arange(1,25)
    fig,ax=plt.subplots(figsize=(12.6,5.5),dpi=200)
    ax.bar(x-.18,danger_rate,.36,color=C_CH,alpha=0.85,label="危险高估越界率")
    ax.bar(x+.18,safe_rate,.36,color=C_DIS,alpha=0.85,label="保守低估越界率")
    ax.set_xticks(range(1,25))
    ax.set_xlabel("目标小时 / h", fontsize=12)
    ax.set_ylabel("全年越界率 / %", fontsize=12)
    ax.set_title("动态非对称光伏风险带：各目标小时全年越界率", fontsize=13, fontweight="bold", pad=12)
    ax.legend(frameon=False, fontsize=10)
    style(ax)
    fig.tight_layout()
    fig.savefig(path,bbox_inches="tight")
    plt.close(fig)
def plot_representative_risk_bands(q43, dates, load, pv, price_actual, cp,
    cfg, run_D, pv_fc_h, out_dir):
    rows,by_day=build_risk_band_diagnostics(
        q43,dates,load,pv,price_actual,cp,cfg,run_D,pv_fc_h
    )
    write_csv(RunConfig.OUT/"动态非对称风险带逐时诊断.csv",rows)
    # 09/10 风险带图已按需删除（逐时诊断数据仍随 CSV 输出）
    return rows
def plot_candidate_risk(hourly, candidate_rows, extra_hours, path):
    x=np.arange(1,25)
    pv=np.array([r["PV危险高估经济风险元"] for r in hourly])/1e4
    pr=np.array([r["价格预测经济暴露元"] for r in hourly])/1e4
    # 候选时点：光伏风险与电价信息价值联合分布（柔色+黑描边，无标题）
    fig,ax=plt.subplots(figsize=(12.5,5.8),dpi=300)
    ax.bar(x,pv,width=0.55,color=SOFT_RED,alpha=0.95,label="光伏危险高估经济风险",zorder=3,**BAR_EDGE)
    ax.bar(x,pr,bottom=pv,width=0.55,color=SOFT_BLUE,alpha=0.95,
        label="电价预测经济暴露",zorder=3,**BAR_EDGE)
    for i,h in enumerate(extra_hours):
        ax.axvline(h,ls="--",lw=1.1,color='#888888',alpha=0.8,
                   label=f"自动候选 {h}:00" if i==0 else None)
    ax.set_xticks(range(1,25))
    if not NO_TITLE:
        ax.set_title("Q4-3新增预报时点：光伏风险与电价信息价值联合分布",
                     fontsize=13,fontweight="bold",pad=12)
    _set_labels(ax,"目标小时 / h","累计风险尺度 / 万元")
    style(ax); _legend_top(ax, ncol=3)
    fig.tight_layout()
    fig.savefig(path,bbox_inches="tight"); plt.close(fig)
    # 信息价值评分曲线（无标题、大字号、图例图外上方）
    xx=np.array([int(r["候选发布时间"]) for r in candidate_rows])
    ss=np.array([float(r["未来风险覆盖得分元"]) for r in candidate_rows])/1e4
    local=np.array([int(r["是否局部峰值"]) for r in candidate_rows],dtype=bool)
    chosen=np.array([int(r["是否最终候选"]) for r in candidate_rows],dtype=bool)
    fig,ax=plt.subplots(figsize=(12.5,5.5),dpi=300)
    # 参考 figures4papers：LineCollection 渐变透明度折线（左浅右深）+ 曲线下浅填充
    LINE_BLUE = '#3775BA'
    from matplotlib.collections import LineCollection
    pts = np.array([xx, ss]).T.reshape(-1, 1, 2)
    segs = np.concatenate([pts[:-1], pts[1:]], axis=1)
    alphas = np.linspace(0.30, 1.0, len(segs))
    lc = LineCollection(segs, colors=[(0.216, 0.459, 0.729, a) for a in alphas],
                        linewidth=3.2, capstyle='round', zorder=3)
    ax.add_collection(lc)
    ax.fill_between(xx, ss, 0, color=LINE_BLUE, alpha=0.05, linewidth=0, zorder=1)
    ax.margins(x=0.02)
    ax.plot(xx, ss, linestyle='none', marker="o", markersize=7.5,
            markerfacecolor=LINE_BLUE, markeredgecolor='white',
            markeredgewidth=1.4, label="未来风险覆盖得分", zorder=4)
    if np.any(local):
        ax.scatter(xx[local],ss[local],s=110,marker="^",color='#E9A6A1',
                   edgecolor='#3775BA',linewidth=1.0,label="局部峰值",zorder=5)
    if np.any(chosen):
        # 星形沿纵轴上移半格，避免覆盖下方的局部峰值三角
        ax.scatter(xx[chosen],ss[chosen]+0.55,s=260,marker="*",color='#E67E22',
                   edgecolor='white',linewidth=1.2,label="最终自动候选",zorder=6)
    ax.set_xticks(range(1,24))
    if not NO_TITLE:
        ax.set_title("新增预报发布时间的未来风险覆盖得分", fontsize=13, fontweight="bold", pad=12)
    _set_labels(ax,"新增预报发布时间 / h","可覆盖信息价值 / 万元")
    style(ax); _legend_top(ax, ncol=3)
    fig.tight_layout()
    fig.savefig(path.parent/"第四问预报时间的信息价值.png",bbox_inches="tight"); plt.close(fig)
def plot_scheme(rows,out):
    names=[r["方案"] for r in rows]; x=np.arange(len(rows)); base=rows[0]
    # 相对基线总费用变化（figures4papers 暖色柔粉红 + 黑描边，无标题、大字号）
    vals=np.array([r["总费用"]-base["总费用"] for r in rows])/1e4
    f,a=plt.subplots(figsize=(max(10.5,1.15*len(rows)),5.2),dpi=300)
    a.bar(x,vals,width=0.40,color='#E9A6A1',alpha=0.95,zorder=3,**BAR_EDGE)
    a.axhline(0,lw=.9,color='black')
    a.set_xticks(x,names,rotation=18,ha="right")
    if not NO_TITLE:
        a.set_title("新增预报时点对总费用的影响", fontsize=13, fontweight="bold", pad=12)
    _set_labels(a,"预报时点方案","相对基线总费用变化 / 万元")
    a.tick_params(labelsize=TICK_FS)
    for i,v in enumerate(vals):
        a.text(i,v+(0.12 if v>=0 else -0.12),f"{v:+.2f}",ha="center",
            va="bottom" if v>=0 else "top",fontsize=11)
    style(a); f.tight_layout()
    f.savefig(out/"第四问新增预报时点对总费用的影响.png",bbox_inches="tight"); plt.close(f)
    dc=np.array([r["总费用"]-base["总费用"] for r in rows])/1e4
    de=np.array([r["紧急购电量kWh"]-base["紧急购电量kWh"] for r in rows])/1000
    f,a=plt.subplots(figsize=(8.6,6.2),dpi=300)
    a.axhline(0,lw=.8,color='black'); a.axvline(0,lw=.8,color='black')
    cmap=plt.get_cmap('OrRd')
    import matplotlib.colors as mcolors
    norm=mcolors.Normalize(vmin=float(de.min()), vmax=float(de.max())+1e-9)
    core=cmap(norm(de))
    a.scatter(dc,de,s=430,color=core,alpha=0.10,zorder=4,linewidths=0)
    a.scatter(dc,de,s=230,color=core,alpha=0.22,zorder=5,linewidths=0)
    a.scatter(dc,de,s=130,color=core,edgecolor='black',linewidth=1.3,zorder=6)
    offsets={
        "原0/6/12/18":        (10,-16,'left','center'),  
        "原方案+21:00":       (12,0,'left','center'),     
        "原方案+16:00":       (0,13,'center','bottom'),   
        "原方案+11:00":       (12,10,'left','center'),    
        "原方案+1:00":        (0,13,'center','bottom'),   
        "原方案+全部自动候选": (-14,0,'right','center'),  
    }
    for i,n in enumerate(names):
        dx,dy,ha,va=offsets.get(n,(0,13,'center','bottom'))
        a.annotate(n,(dc[i],de[i]),xytext=(dx,dy),textcoords="offset points",
                   fontsize=13,ha=ha,va=va)
    sm=plt.cm.ScalarMappable(cmap=cmap,norm=norm); sm.set_array([])
    cbar=f.colorbar(sm,ax=a,pad=0.015,aspect=28)
    cbar.set_label('相对基线紧急购电变化 / MWh',fontsize=12)
    cbar.ax.tick_params(labelsize=10)
    a.set_xlim(-0.8,19.6); a.set_ylim(-0.5,5.6)
    if not NO_TITLE:
        a.set_title("新增预报方案的经济性—可靠性 Pareto 分布", fontsize=13, fontweight="bold", pad=12)
    _set_labels(a,"相对基线总费用变化 / 万元","相对基线紧急购电变化 / MWh")
    a.tick_params(labelsize=TICK_FS)
    a.grid(True,axis="y",linestyle="--",linewidth=0.7,alpha=0.3)
    a.spines["top"].set_visible(False); a.spines["right"].set_visible(False)
    f.tight_layout()
    f.savefig(out/"第四问可靠性Pareto分布.png",bbox_inches="tight"); plt.close(f)
def main():
    ensure_dir(RunConfig.OUT); ensure_dir(RunConfig.FIG)
    q43_path=first_existing(RunConfig.Q43_CANDIDATES)
    att3=first_existing(RunConfig.ATT3_CANDIDATES)
    cp_path=first_existing(RunConfig.CHECKPOINT_CANDIDATES)
    q43=import_py(q43_path,"q43_main_for_extra"); cfg=q43.CFG; cfg.base_dir=str(RunConfig.BASE)
    cfg.max_days=RunConfig.MAX_DAYS;cfg.verbose_each_day=False;cfg.use_intraday_load_bias=False
    # 修正附件路径到当前目录
    for attr,name in [("att1","附件1.xlsx"),("att2","附件2.xlsx"),
        ("att3","附件3.xlsx"),("att4","附件4.xlsx")]:
        p=RunConfig.BASE/name
        if p.exists(): setattr(cfg,attr,str(p))
    class Silent:
        def info(self,*a,**k): pass
        def warning(self,*a,**k): pass
    dates,price_prior,load_prior,load,pv,price_actual=q43.load_inputs(cfg,Silent())
    run_D=min(len(dates),RunConfig.MAX_DAYS); dates=dates[:run_D]
    load=load[:run_D]; pv=pv[:run_D]; price_actual=price_actual[:run_D]
    cp=load_checkpoint(cp_path,run_D)
    load_fc=np.asarray(cp["load_fc"][:run_D],float)
    price_fc=np.asarray(cp["price_fc"][:run_D],float)
    fallback=np.asarray(cp["fallback_reserve"][:run_D],float)
    official_hist=np.asarray(cp["stage_net_error_hist"][:run_D],float)
    official_all=load_forecast_records_xlsx(att3,"official"); dayset=set(dates)
    official=[r for r in official_all if r["issue_dt"].date() in dayset]
    hourly,pv_fc_h,price_fc_h=build_combined_hourly_risk(q43,dates,load,pv,
        price_actual,cp,cfg,run_D,official)
    write_csv(RunConfig.OUT/"各小时_光伏风险与电价经济暴露.csv",hourly)
    extra,crows=select_extra_hours(hourly)
    write_csv(RunConfig.OUT/"自动候选时点筛选.csv",crows)
    # 兼容旧文件名，便于已有论文/脚本继续读取
    write_csv(RunConfig.OUT/"综合风险候选时点筛选.csv",crows)
    if not extra:
        raise RuntimeError("未筛选出新增预报时点")
    plot_candidate_risk(hourly,crows,extra,RunConfig.FIG/"第四问候选时点.png")
    plot_representative_risk_bands(
        q43,dates,load,pv,price_actual,cp,cfg,run_D,pv_fc_h,RunConfig.FIG
    )
    # 生成新增PV代理；
    augmented,diag=build_augmented_records(official,extra,dates,pv[:,[6*h-1 for h in range(1,25)]])
    write_csv(RunConfig.OUT/"新增光伏预报偏差修正明细.csv",diag)
    tag="_".join(map(str,extra))
    export_records_xlsx(RunConfig.OUT/f"附件3_新增预报_{tag}_在线代理.xlsx",augmented)
    rec_map=records_by_key(augmented)
    extra_hist=build_extra_stage_error_history(q43,extra,rec_map,dates,load_fc,
        load,pv,cfg,run_D)
    lp=q43.build_day_ahead_templates(cfg)
    base=(6,12,18); schemes=[("原0/6/12/18",base)]
    for h in extra:
        schemes.append((f"原方案+{h}:00",tuple(sorted(set(base+(h,))))))
    if RunConfig.RUN_ALL_CANDIDATES_TOGETHER and len(extra)>=2:
        schemes.append(("原方案+全部自动候选",tuple(sorted(set(base+tuple(extra))))))
    uniq=[]; seen=set()
    for x in schemes:
        if x[1] not in seen: uniq.append(x); seen.add(x[1])
    rows=[]
    for name,hrs in uniq:
        records=official if hrs==base else augmented
        r=simulate_scheme(q43,name,hrs,records,dates,price_actual,load,
            pv,load_fc,price_fc,fallback,official_hist,extra_hist,run_D,cfg,lp)
        rows.append(summarize(name,hrs,r,run_D))
    write_csv(RunConfig.OUT/"新增预报时点_完整滚动LP方案对比.csv",rows)
    plot_scheme(rows,RunConfig.FIG)
    base_row=rows[0]; metrics=["总费用","紧急购电量kWh","日紧急购电量P95kWh",
        "单日最大紧急购电量kWh","富余电量kWh","向上调整电量kWh","向下调整电量kWh"]
    delta=[]
    for r in rows[1:]:
        z={"方案":r["方案"]}
        for m in metrics:
            b=float(base_row[m]); v=float(r[m]); z[m+"_变化量"]=v-b
            z[m+"_变化率%"]=(v-b)/b*100 if abs(b)>1e-12 else np.nan
        delta.append(z)
    write_csv(RunConfig.OUT/"新增预报方案_相对基线改善.csv",delta)
    minc=min(rows,key=lambda r:r["总费用"]); mine=min(rows,key=lambda r:r["紧急购电量kWh"])
    lines=["Q4-3新增其他预报时点综合验证摘要","="*72,
           "自动筛选候选时点："+", ".join(f"{h}:00" for h in extra),
           "筛选原则：不预设具体时刻。先将PV危险高估经济风险与电价预测经济暴露统一为元，",
           "再计算每个可能发布时间对未来若干小时风险的指数衰减覆盖得分，由局部峰值和最小时距自动形成候选集。",
           f"最低总费用方案：{minc['方案']}，{minc['总费用']:.2f} 元",
           f"最低紧急购电方案：{mine['方案']}，{mine['紧急购电量kWh']:.2f} kWh","",
           "最终是否增加时点，应以完整LP的总费用、紧急购电、P95、最大单日风险和富余电量共同判断；",
           "风险评分只负责生成候选，不直接替代滚动LP的最终验证。",
           "论文图形证据链：动态非对称风险带->光伏/电价联合风险->自动候选评分->完整LP经济性与可靠性比较。"]
    (RunConfig.OUT/"综合验证摘要.txt").write_text("\n".join(lines),encoding="utf-8")
    print("[完成] Q4-3-2 运行结束，输出目录：",RunConfig.OUT)
if __name__ == "__main__":
    main()
